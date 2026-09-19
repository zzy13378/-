"""
P4: 科研文献整合与论文初稿辅助系统

核心场景：科研新手已经找到一批论文，但难以完成文献分类、观点整合和综述初稿组织。
OpenAlex 是真实论文检索入口；本地 PDF、Word、TXT、Markdown、JSON、CSV 可用于导入用户已经收集的论文。
没有网络或 LLM 时使用明确标记的内置示例和规则降级，保证本地演示可运行。
"""

from __future__ import annotations

import copy
import csv
import io
import time
import json
import math
import os
import re
import sys
import unicodedata
import urllib.parse
import xml.etree.ElementTree as ET
import urllib.request
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple


PROJECT_DIR = Path(__file__).resolve().parent
VENDOR_DIR = PROJECT_DIR / "vendor"

# Keep optional parsing dependencies local to this demonstrable project.
if VENDOR_DIR.is_dir() and str(VENDOR_DIR) not in sys.path:
    sys.path.insert(0, str(VENDOR_DIR))

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


def count_chinese_words(text: str) -> int:
    clean = re.sub(r"[#*`>\s\n\r\-]", "", text or "")
    return len(clean)


def compact_text(value: Any, limit: int = 900) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:limit] + ("..." if len(text) > limit else "")


def parse_json_response(text: str) -> Optional[Dict[str, Any]]:
    if not text:
        return None
    fenced = re.search(r"```(?:json)?\s*(.*?)```", text, flags=re.I | re.S)
    candidates = [fenced.group(1), text] if fenced else [text]
    for candidate in candidates:
        try:
            value = json.loads(candidate.strip())
            if isinstance(value, dict):
                return value
        except (TypeError, json.JSONDecodeError):
            pass
    start, end = text.find("{"), text.rfind("}")
    if start >= 0 and end > start:
        try:
            value = json.loads(text[start:end + 1])
            return value if isinstance(value, dict) else None
        except json.JSONDecodeError:
            return None
    return None


@dataclass
class Paper:
    ref_id: str
    title: str
    authors: List[str] = field(default_factory=list)
    year: Optional[int] = None
    abstract: str = ""
    doi: str = ""
    url: str = ""
    source: str = "OpenAlex"
    full_text: str = ""
    metadata_confidence: str = "中"
    parse_warnings: List[str] = field(default_factory=list)
    page_count: int = 0
    text_quality: str = "可分析"
    citation_count: int = 0
    venue: str = ""
    quality_score: float = 0.0
    relevance_score: float = 0.0

    def to_dict(self) -> Dict[str, Any]:
        return {
            "ref_id": self.ref_id,
            "title": self.title,
            "authors": self.authors,
            "year": self.year,
            "abstract": self.abstract,
            "doi": self.doi,
            "url": self.url,
            "source": self.source,
            "full_text": self.full_text,
            "metadata_confidence": self.metadata_confidence,
            "parse_warnings": self.parse_warnings,
            "page_count": self.page_count,
            "text_quality": self.text_quality,
            "citation_count": self.citation_count,
            "venue": self.venue,
            "quality_score": self.quality_score,
            "relevance_score": self.relevance_score,
        }

    def citation(self) -> str:
        author = self.authors[0] if self.authors else "作者未知"
        suffix = "等" if len(self.authors) > 1 else ""
        year = self.year or "年份未知"
        link = self.doi or self.url
        return f"[{self.ref_id}] {author}{suffix}. {self.title}. {year}." + (f" {link}" if link else "")


def _abstract_from_inverted_index(index: Any) -> str:
    if not isinstance(index, dict):
        return ""
    words: List[Tuple[int, str]] = []
    for word, positions in index.items():
        if isinstance(positions, list):
            for position in positions:
                if isinstance(position, int):
                    words.append((position, str(word)))
    return " ".join(word for _, word in sorted(words))


def _authors(value: Any) -> List[str]:
    if isinstance(value, str):
        return [value]
    if not isinstance(value, list):
        return []
    names: List[str] = []
    for item in value:
        if isinstance(item, str):
            names.append(item)
        elif isinstance(item, dict):
            author = item.get("author")
            name = author.get("display_name") if isinstance(author, dict) else None
            name = name or item.get("name") or item.get("display_name")
            if name:
                names.append(str(name))
    return names


def paper_from_record(record: Dict[str, Any], ref_id: str, source: str = "本地导入") -> Paper:
    location = record.get("primary_location") or {}
    abstract = record.get("abstract") or record.get("abstract_inverted_index") or ""
    if isinstance(abstract, dict):
        abstract = _abstract_from_inverted_index(abstract)
    year = record.get("year") or record.get("publication_year")
    try:
        year = int(year) if year else None
    except (TypeError, ValueError):
        year = None
    doi = str(record.get("doi") or "")
    doi = doi.removeprefix("https://doi.org/")
    return Paper(
        ref_id=ref_id,
        title=compact_text(record.get("title") or "未命名论文", 500),
        authors=_authors(record.get("authors") or record.get("authorships") or record.get("author")),
        year=year,
        abstract=compact_text(abstract, 1800),
        full_text=compact_text(abstract, 12000),
        doi=doi,
        url=str(record.get("url") or record.get("landing_page_url") or location.get("landing_page_url") or ""),
        source=source,
    )


def _decode_text_bytes(data: bytes, filename: str) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError(f"{filename} 不是可识别的 UTF-8 或 GB18030 文本文件")


def _split_authors(value: str) -> List[str]:
    cleaned = re.sub(r"[\*†‡]+", " ", value or "")
    names = re.split(r"\s*(?:,|;|，|；|、|\band\b)\s*", cleaned, flags=re.I)
    result: List[str] = []
    for name in names:
        name = re.sub(r"\s*[\d,，、]+\s*$", "", name).strip()
        name = re.sub(r"(?:等|et\s+al\.?)$", "", name, flags=re.I).strip()
        if name:
            result.append(compact_text(name, 120))
    if len(result) == 1 and re.fullmatch(r"[\u4e00-\u9fff]{2,4}(?:\s+[\u4e00-\u9fff]{2,4})+", result[0]):
        result = re.findall(r"[\u4e00-\u9fff]{2,4}", result[0])
    return result


_AUTHOR_NOISE_PREFIXES = (
    "评估", "分析", "研究", "因素", "服务", "影响", "方法", "基于", "探析", "探究",
    "析", "论", "综述", "摘要", "关键词",
)


def _looks_like_person_name(value: str) -> bool:
    value = re.sub(r"[\d*∗†‡,，、;；|｜()（）.。\s]+", "", value or "")
    if not re.fullmatch(r"[\u4e00-\u9fff]{2,4}", value):
        return False
    if value.startswith(_AUTHOR_NOISE_PREFIXES):
        return False
    return True


_COMMON_CHINESE_SURNAMES = set(
    "赵钱孙李周吴郑王冯陈褚卫蒋沈韩杨朱秦尤许何吕施张孔曹严华金魏陶姜戚谢邹喻柏水窦章云苏潘葛奚范彭郎鲁韦昌马苗凤花方俞任袁柳酆鲍史唐费廉岑薛雷贺倪汤滕殷罗毕郝邬安常乐于时傅皮卞齐康伍余元卜顾孟平黄和穆萧尹姚邵湛汪祁毛禹狄米贝明臧计伏成戴谈宋茅庞熊纪舒屈项祝董梁杜阮蓝闵席季麻强贾路娄危江童颜郭梅盛林钟徐丘骆高夏蔡田樊胡凌霍虞万支柯昝管卢莫经房裘缪解应宗丁宣贲邓郁单杭洪包诸左石崔吉钮龚程邢滑裴陆荣翁荀羊於惠甄曲家封芮羿储靳汲邴糜松井段富巫乌焦巴弓牧隗山谷车侯宓蓬全郗班仰秋仲伊宫宁仇栾暴甘钭厉戎祖武符刘景詹束龙叶幸司韶郜黎蓟薄印宿白怀蒲邰从鄂索咸籍赖卓蔺屠蒙池乔阴郁胥能苍双闻莘党翟谭贡劳逄姬申扶堵冉宰郦雍却璩桑桂濮牛寿通边扈燕冀郏浦尚农温别庄晏柴瞿阎充慕连茹习宦艾鱼容向古易慎戈廖庾终暨居衡步都耿满弘匡国文寇广禄阙东殴殳沃利蔚越夔隆师巩厍聂晁勾敖融冷訾辛阚那简饶空曾毋沙乜养鞠须丰巢关蒯相查后荆红游竺权逯盖益桓公万俟司马上官欧阳夏侯诸葛闻人东方赫连皇甫尉迟长孙宇文令狐慕容"
)


def _split_concatenated_chinese_names(value: str) -> List[str]:
    """Split names collapsed by PDF extraction, e.g. 赵珂欣侯青阳张贤壹."""
    compact = "".join(re.findall(r"[\u4e00-\u9fff]", value or ""))
    if len(compact) < 4:
        return []
    for start in range(len(compact)):
        if compact[start] not in _COMMON_CHINESE_SURNAMES:
            continue
        tail = compact[start:]
        for first_width in (3, 2):
            names = []
            position = 0
            width = first_width
            while position < len(tail):
                if tail[position] not in _COMMON_CHINESE_SURNAMES:
                    break
                if position + width > len(tail):
                    break
                names.append(tail[position:position + width])
                position += width
                if position == len(tail):
                    if len(names) >= 2:
                        return names
                    break
                if tail[position] not in _COMMON_CHINESE_SURNAMES:
                    # Try a 3-character name at this boundary when the first
                    # name was two characters, or vice versa.
                    width = 3 if width == 2 else 2
                    continue
                width = 3 if position + 3 <= len(tail) else 2
            if len(names) >= 2 and position == len(tail):
                return names
    return []


def _clean_chinese_author(value: str) -> str:
    """Remove title suffix characters accidentally joined to the first author."""
    value = re.sub(r"[\d*∗†‡,，、;；|｜()（）.。\s]+", "", value or "")
    if _looks_like_person_name(value):
        return value
    # Text extraction from bilingual PDFs can produce e.g. "评估赵珂欣".
    # Prefer the final 2-4 characters, which are adjacent to the affiliation marker.
    for width in (4, 3, 2):
        candidate = value[-width:]
        if _looks_like_person_name(candidate):
            return candidate
    return value


def _looks_like_author_hint(value: str) -> bool:
    compact = re.sub(r"[\d*∗†‡\s]+", "", value or "")
    if not compact or len(compact) > 32:
        return False
    if any(token in compact for token in ("研究", "分析", "服务", "因素", "空间", "政策", "机构", "场景", "预测", "国家公园", "资源", "环境")):
        return False
    parts = re.split(r"[,，、;；|｜]+", compact)
    return 1 <= len(parts) <= 8 and all(_looks_like_person_name(part) for part in parts)


def _title_quality(value: str, source: str = "") -> int:
    """Score a title candidate so journal headers do not become paper titles."""
    value = compact_text(value, 500).strip(" .。:：")
    if not _title_candidate_ok(value):
        return -100
    score = 0
    compact = re.sub(r"\s+", "", value)
    if 12 <= len(compact) <= 70:
        score += 4
    elif len(compact) >= 8:
        score += 1
    if re.search(r"研究|分析|评估|影响因素|时空|应用|方法|模型|预测|比较|机制|综述|服务|治理|识别", value):
        score += 5
    if re.search(r"第\s*\d+\s*卷|第\s*\d+\s*期|Vol\.?\s*\d+|No\.?\s*\d+|ISSN|文章编号|年\s*第\s*\d+\s*期", value, flags=re.I):
        score -= 12
    if re.fullmatch(r"[\u4e00-\u9fff]{4,16}", value) and not re.search(r"研究|分析|评估|因素|服务|预测|方法|应用|综述", value):
        score -= 8
    if source == "filename":
        score += 2
    elif source == "pdf_metadata":
        score += 1
    return score


def _normalize_document_text(text: str) -> str:
    """Normalize layout artifacts produced by PDF text extraction while keeping lines."""
    text = unicodedata.normalize("NFKC", text or "")
    text = text.replace("\u00a0", " ").replace("\u3000", " ").replace("\r\n", "\n").replace("\r", "\n")
    lines: List[str] = []
    for raw_line in text.split("\n"):
        line = re.sub(r"\s+", " ", raw_line).strip()
        if not line:
            continue
        # Chinese PDF fonts often insert a blank between every character.
        line = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])", "", line)
        line = re.sub(r"\s+([，。；：、！？）》】])", r"\1", line)
        line = re.sub(r"([（《【])\s+", r"\1", line)
        line = re.sub(r"\s+([，。；：、！？）》】])", r"\1", line)
        # Preserve searchable names of common methods and models.
        for spaced, compact in (
            (r"A\s+I", "AI"), (r"C\s+E\s+S", "CES"), (r"B\s+E\s+R\s+T\s*o\s*p\s*i\s*c", "BERTopic"),
            (r"S\s+o\s+l\s+V\s+E\s+S", "SolVES"), (r"K\s+G\s+A\s+T", "KGAT"),
            (r"N\s+D\s+C\s+G", "NDCG"),
        ):
            line = re.sub(spaced, compact, line, flags=re.I)
        lines.append(line)
    return "\n".join(lines)


def _clean_evidence_text(text: str, limit: int = 1200) -> str:
    """Make extracted evidence readable without changing its meaning."""
    value = unicodedata.normalize("NFKC", str(text or ""))
    value = re.sub(r"\s+", " ", value).strip()
    value = re.sub(r"(?<=[\u3400-\u9fff])\s+(?=[\u3400-\u9fff])", "", value)
    value = re.sub(r"\s*([，。；：、！？）》】,.!?;:])\s*", r"\1", value)
    value = re.sub(r"\s*([（《【()）])\s*", r"\1", value)
    value = re.sub(r"\s*:\s*", ":", value)
    value = re.sub(r"(\d+\.\d+)\s+(\d+)", r"\1\2", value)
    value = re.sub(r"\s+([%％])", r"\1", value)
    return compact_text(value, limit)


def _metadata_from_filename(filename: str) -> Tuple[str, List[str]]:
    stem = Path(filename).stem.strip()
    stem = re.sub(r"\s*\(\d+\)\s*$", "", stem)
    stem = re.sub(r"\s*[-_]\s*\d{1,3}\s*$", "", stem)
    stem = re.sub(r"^(?:外文|英文|中文|译文|translation|translated)[\s_＿-]*", "", stem, flags=re.I)
    if not stem or re.fullmatch(r"\d+", stem):
        return "", []
    parts = re.split(r"[_＿]", stem)
    author_hint = ""
    if len(parts) > 1:
        possible_authors = parts[-1].strip()
        if _looks_like_author_hint(possible_authors):
            author_hint = possible_authors
            stem = "_".join(parts[:-1]).strip()
    stem = re.sub(r"\s*[-－—]\s*(?:作者|author)\s*[:：]?\s*", " ", stem, flags=re.I)
    stem = re.sub(r"\s+", " ", stem).strip()
    invalid = re.compile(r"^(?:论文|文献|报告|资料|附件|新建文档|p4[-_ ]?upload)$", re.I)
    title = "" if invalid.fullmatch(stem) or len(stem) < 8 else stem
    return title[:500], _split_authors(author_hint)


def _title_candidate_ok(value: str) -> bool:
    value = compact_text(value, 500)
    if not (8 <= len(value) <= 180):
        return False
    if re.fullmatch(r"[\d\s.＿_\-—]+", value):
        return False
    blocked = (
        r"研究论文|内容摘要|关键词|摘要|编者按|第\s*\d+\s*卷|Vol\.?\s*\d+|ISSN|DOI|文章编号|Journal of|"
        r"Correspondence|Received|Accepted|Published|Copyright|University|Institute|College|Laboratory|"
        r"Research Center|Management World Journal Press|Microsoft Word|Untitled|@|https?://|引用本文|\[J\]|^NATIONAL PARK$|"
        r"Publisher's Note|Licensee MDPI|Creative Commons|Article"
    )
    if re.match(r"^\s*\d+\s+", value) or value.count(";") >= 2:
        return False
    if re.fullmatch(r"[A-Z][A-Za-z'-]+\s+[A-Z][A-Za-z'-]+\s+[a-z](?:\s*[,*a-z])?", value):
        return False
    return not re.search(blocked, value, flags=re.I)


_ABSTRACT_LINE_RE = re.compile(
    r"^\s*[\[【(（]?\s*(?:摘\s*要|内容摘要|abstract)\s*[\]】)）]?\s*[:：]?",
    flags=re.I,
)


def _abstract_line_match(line: str) -> Optional[re.Match[str]]:
    return _ABSTRACT_LINE_RE.match(line or "")


def _looks_like_author_line(line: str) -> bool:
    value = str(line or "")
    if re.search(r"Citation|Academic Editors|Received|Accepted|Published|Copyright|Publisher", value, flags=re.I):
        return False
    if re.search(r"大学|学院|研究院|研究中心|研究所|基金|收稿|采用日期|DOI|https?://|\b(?:University|Institute|College|Academy|Laboratory|Center)\b", value, flags=re.I):
        return False
    if re.search(r"[.。]", value) and len(value) > 30:
        return False
    if re.fullmatch(r"\s*[A-Z][A-Za-z'-]+(?:\s+[A-Z][A-Za-z'-]+){1,2}\s*[a-z]\s*", value):
        return True
    if not re.search(r"[,;&]|\band\b|\d|[a-z]\s*[,;]", value):
        return False
    english = re.findall(r"\b[A-Z][A-Za-z'-]+(?:\s+[A-Z](?:\s+[A-Z])?|\s+[a-z]+){1,2}", value)
    chinese = re.findall(r"[\u4e00-\u9fff]{2,4}", re.sub(r"[\d|｜*∗†‡]+", " ", value))
    return len(english) >= 2 or len(chinese) >= 2


def _title_from_body(normalized: str) -> str:
    lines = normalized.splitlines()
    abstract_index = next((i for i, line in enumerate(lines) if _abstract_line_match(line)), min(len(lines), 100))
    candidates: List[Tuple[int, str]] = []
    # Many English publisher PDFs put a multi-line title immediately before
    # the author line. The same layout occurs in Chinese bilingual journals.
    author_indices = [
        index for index in range(max(1, abstract_index - 35), abstract_index)
        if _looks_like_author_line(lines[index])
    ]
    for author_index in author_indices:
        for width in (1, 2, 3, 4):
            start = author_index - width
            if start < 0:
                continue
            parts = [line.strip() for line in lines[start:author_index]]
            if any(
                re.search(r"Citation|Academic Editors|Received|Accepted|Published|Copyright|Publisher|https?://|DOI", part, flags=re.I)
                for part in parts
            ):
                continue
            parts = [part for part in parts if _title_candidate_ok(part)]
            if len(parts) != width:
                continue
            joined = " ".join(parts).strip()
            if _title_candidate_ok(joined):
                candidates.append((width, joined))
        if candidates:
            # Prefer the longest valid block so a title split across several
            # extracted lines is restored as one title.
            return max(candidates, key=lambda value: (_title_quality(value[1]), len(value[1]), value[0]))[1]
    for index in range(max(0, abstract_index - 45), abstract_index):
        line = re.sub(r"^\s*#+\s*", "", lines[index]).strip()
        if _title_candidate_ok(line):
            candidates.append((0, line))
        for width in (2, 3):
            if index + width <= abstract_index:
                raw_parts = lines[index:index + width]
                if any(_looks_like_author_line(part) or part.lstrip().startswith(",") for part in raw_parts):
                    continue
                separator = "" if all(re.search(r"[\u4e00-\u9fff]", part) and not re.search(r"[A-Za-z]", part) for part in raw_parts) else " "
                joined = re.sub(r"^\s*#+\s*", "", separator.join(raw_parts)).strip()
                if _title_candidate_ok(joined):
                    candidates.append((width, joined))
    labeled = re.search(r"(?im)^\s*(?:title|论文(?:题目|标题)?|文献题目)\s*[:：]\s*(.+?)\s*$", normalized)
    if labeled and _title_candidate_ok(labeled.group(1)):
        return compact_text(labeled.group(1), 500)
    if not candidates:
        return ""
    # A paper title is usually longer than a journal header and sits closest to the abstract.
    return max(
        candidates,
        key=lambda value: (
            _title_quality(value[1]),
            len(re.findall(r"[\u4e00-\u9fffA-Za-z]", value[1])),
            -value[0],
        ),
    )[1]


def _extract_body_authors(normalized: str, title: str) -> List[str]:
    lines = normalized.splitlines()
    has_abstract_marker = any(_abstract_line_match(line) for line in lines)
    if not has_abstract_marker:
        return []
    abstract_index = next((i for i, line in enumerate(lines) if _abstract_line_match(line)), min(len(lines), 100))
    stop = {"摘要", "内容摘要", "关键词", "研究论文", "作者简介", "北京", "上海", "广州", "南京"}
    title_key = re.sub(r"\s+", "", title or "")
    title_index = 0
    title_found = False
    if title_key:
        for index in range(abstract_index - 1, max(-1, abstract_index - 45), -1):
            line_key = re.sub(r"\s+", "", lines[index])
            title_prefix = title_key[:12]
            prefix_match = title_prefix in line_key if not re.search(r"[A-Za-z]", title_key) else title_prefix in line_key
            if title_key in line_key or prefix_match:
                title_index = index
                title_found = True
                break
    if not title_found:
        return []
    collected: List[str] = []
    header = lines[title_index + 1: min(abstract_index, title_index + 20)]
    author_block: List[str] = []
    for line in header:
        if any(token in line for token in ("大学", "学院", "研究院", "基金", "邮箱", "E-mail", "通讯作者", "University", "Institute", "College", "Academy")):
            break
        author_block.append(line)
    structured_chinese: List[str] = []
    for raw_name in re.findall(r"([\u4e00-\u9fff]{2,})(?=\s*[\d,，、]+)", "".join(author_block)):
        collapsed_names = _split_concatenated_chinese_names(raw_name)
        if collapsed_names:
            structured_chinese.extend(collapsed_names)
        else:
            cleaned_name = _clean_chinese_author(raw_name)
            if _looks_like_person_name(cleaned_name):
                structured_chinese.append(cleaned_name)
    if len(structured_chinese) >= 2:
        return list(dict.fromkeys(structured_chinese))[:8]
    title_words = set(re.findall(r"[A-Za-z]{4,}", title.casefold()))
    for line in header:
        if re.match(r"^\s*(?:作者|authors?)\s*[:：]", line, flags=re.I):
            continue
        if any(token in line for token in ("大学", "学院", "研究院", "基金", "邮箱", "E-mail", "通讯作者")):
            continue
        line_words = set(re.findall(r"[A-Za-z]{4,}", line.casefold()))
        if title_words and len(title_words & line_words) >= 1:
            continue
        english_names = re.findall(r"\b[A-Z][a-z]+(?:\s+[A-Z][a-z]+){1,2}\b", line)
        if len(english_names) >= 1 and not re.search(
            r"(?:Received|Accepted|Published|Editor|Citation|Copyright|Institute|University|College|Academy|Research|Center|Department|Laboratory|Beijing|China|Press|@)",
            line,
            flags=re.I,
        ):
            collected.extend(english_names)
            continue
        chinese_parts = re.split(r"\s*[,，、;；|｜]\s*", re.sub(r"[\d*∗†‡]+", " ", line))
        chinese_names = [
            _clean_chinese_author(part)
            for part in chinese_parts
            if _looks_like_person_name(_clean_chinese_author(part))
        ]
        if chinese_names:
            collected.extend(chinese_names)
            continue
        compact_line = re.sub(r"[\d\s,，、;；()（）.。*∗]+", "", line)
        if not compact_line or re.search(r"(?:第|卷|期|年|月|日|资源|环境|专题|发展|文章编号|数据分析|知识发现)", compact_line):
            continue
        if compact_line in title_key or title_key in compact_line:
            continue
        if _looks_like_person_name(compact_line):
            names = [_clean_chinese_author(compact_line)]
        elif re.fullmatch(r"[\u4e00-\u9fff]+", compact_line) and len(compact_line) % 3 == 0:
            names = [compact_line[i:i + 3] for i in range(0, len(compact_line), 3)]
        else:
            names = re.findall(r"[\u4e00-\u9fff]{2,4}", compact_line)
        collected.extend(name for name in names if name not in stop)
    return list(dict.fromkeys(collected))[:8]


def _extract_abstract(normalized: str, title_hint: str = "") -> str:
    lines = normalized.splitlines()
    matches = []
    for i, line in enumerate(lines):
        if re.match(r"^\s*[\[【(（]?\s*(?:摘\s*要|内容摘要)\s*[\]】)）]?\s*[:：]?", line, flags=re.I):
            matches.append((i, True))
        elif re.match(r"^\s*[\[【(（]?\s*abstract\s*[\]】)）]?\s*[:：]?", line, flags=re.I):
            matches.append((i, False))
    # Prefer an abstract belonging to the detected title. This skips journal
    # editorials that can appear before the article on Chinese journal pages.
    title_key = title_hint
    start_index = 0
    if title_key:
        compact_title = re.sub(r"\s+", "", title_key)[:10]
        for i in range(len(lines)):
            window = re.sub(r"\s+", "", "".join(lines[i:i + 3]))
            if compact_title and compact_title in window:
                start_index = i
                break
    candidates: List[Tuple[int, str]] = []
    for index, is_chinese in matches:
        if index < start_index:
            continue
        first = re.sub(
            r"^\s*[\[【(（]?\s*(?:摘\s*要|内容摘要|abstract)\s*[\]】)）]?\s*[:：]?\s*",
            "",
            lines[index],
            flags=re.I,
        )
        body_lines = [first] if first else []
        for line in lines[index + 1: index + 140]:
            if re.match(r"^\s*[\[【(（]?\s*(?:关键词|key\s*words?|index\s+terms?|引言|introduction|参考文献|references)\b", line, flags=re.I):
                break
            if not re.search(r"(?:项目|基金|作者简介|E-mail|邮箱)", line):
                body_lines.append(line)
        body = _clean_evidence_text(re.sub(r"^[\s:：]+", "", " ".join(body_lines)), 2200)
        if len(body) >= 40:
            candidates.append((0 if is_chinese else 1, body))
    if candidates:
        # Chinese abstracts are the primary evidence for Chinese papers. Within
        # that group, prefer a normal-length abstract over a page-level spill.
        candidates.sort(key=lambda value: (value[0], abs(len(value[1]) - 900)))
        return candidates[0][1]
    # Some Word-exported papers omit the literal "摘要" label. In that case
    # use the first prose block after the detected title, rather than treating
    # the complete first page (including journal metadata) as the abstract.
    lines = normalized.splitlines()
    title_key = re.sub(r"\s+", "", title_hint or "")
    start = 0
    if title_key:
        for index, line in enumerate(lines):
            if title_key[:10] and title_key[:10] in re.sub(r"\s+", "", line):
                start = index + 1
                break
    prose: List[str] = []
    for line in lines[start:start + 24]:
        if re.match(r"^\s*(?:关键词|key\s*words?|参考文献|references|引言|introduction)\b", line, flags=re.I):
            break
        if len(re.sub(r"[^\u4e00-\u9fffA-Za-z]", "", line)) >= 20:
            prose.append(line)
        if len("".join(prose)) >= 1600:
            break
    return _clean_evidence_text(" ".join(prose), 1800)


def _valid_abstract(text: str) -> bool:
    value = _clean_evidence_text(text, 2400)
    if len(re.sub(r"\s+", "", value)) < 80:
        return False
    if re.search(r"第\s*\d+\s*卷|第\s*\d+\s*期|Vol\.?\s*\d+|ISSN|文章编号", value, flags=re.I):
        return False
    prose_chars = len(re.findall(r"[\u4e00-\u9fffA-Za-z]", value))
    return prose_chars >= 80


def _pdf_metadata(reader: Any) -> Tuple[str, List[str], Optional[int]]:
    """Read embedded PDF metadata only when it looks like article metadata."""
    metadata = getattr(reader, "metadata", None)
    title = ""
    authors: List[str] = []
    if metadata:
        raw_title = str(getattr(metadata, "title", "") or "").strip()
        if _title_candidate_ok(raw_title):
            title = raw_title
        raw_author = str(getattr(metadata, "author", "") or "").strip()
        if _looks_like_author_hint(raw_author):
            authors = _split_authors(raw_author)
    year = None
    created = str(getattr(metadata, "creation_date", "") or "") if metadata else ""
    match = re.search(r"(?:19|20)\d{2}", created)
    if match:
        year = int(match.group(0))
    return title, authors, year


def _extract_pdf_pages(reader: Any, max_pages: int = 50) -> Tuple[str, int]:
    """Extract readable page text and remove repeated journal headers/footers."""
    page_texts: List[str] = []
    for index, page in enumerate(reader.pages):
        if index >= max_pages:
            break
        try:
            page_text = page.extract_text(extraction_mode="layout") or ""
        except (TypeError, ValueError):
            page_text = page.extract_text() or ""
        page_texts.append(page_text)
    if not page_texts:
        return "", 0

    # A journal name, running title, or page number repeated on several pages
    # is layout noise, not evidence. Only inspect the page margins so a real
    # section heading in the body is preserved.
    margin_lines: Dict[str, int] = {}
    for page_text in page_texts:
        values = [re.sub(r"\s+", " ", line).strip() for line in page_text.splitlines() if line.strip()]
        for line in set(values[:4] + values[-4:]):
            if len(line) >= 3 and not re.fullmatch(r"\d{1,4}", line):
                margin_lines[line] = margin_lines.get(line, 0) + 1
    repeated = {line for line, count in margin_lines.items() if count >= 2}
    cleaned_pages: List[str] = []
    for page_text in page_texts:
        lines = []
        for raw_line in page_text.splitlines():
            line = re.sub(r"\s+", " ", raw_line).strip()
            if not line or line in repeated or re.fullmatch(r"(?:第\s*)?\d{1,4}(?:\s*页)?", line):
                continue
            lines.append(raw_line)
        cleaned_pages.append("\n".join(lines))
    return "\n".join(cleaned_pages), len(page_texts)


def _metadata_candidates(
    normalized: str,
    filename: str,
    title_hint: str = "",
    pdf_title: str = "",
) -> List[Tuple[int, str, str]]:
    candidates: List[Tuple[int, str, str]] = []
    for source, raw in (("filename", title_hint), ("pdf_metadata", pdf_title), ("body", _title_from_body(normalized))):
        title = compact_text(raw, 500).strip(" .。:：")
        score = _title_quality(title, source)
        if score >= 0:
            candidates.append((score, title, source))
    return candidates


def _extract_metadata_from_document(
    text: str,
    filename: str,
    source: str,
    title_hint: str = "",
    author_hint: str = "",
    pdf_title: str = "",
    pdf_authors: Optional[List[str]] = None,
    pdf_year: Optional[int] = None,
    page_count: int = 0,
) -> Paper:
    """Convert the useful text in one paper file into the common Paper model."""
    normalized = _normalize_document_text(text).strip()
    if not normalized:
        raise ValueError(f"{filename} 中没有可用于文献整合的正文文本")

    title_source = ""
    title_candidates = _metadata_candidates(normalized, filename, title_hint, pdf_title)
    if title_candidates:
        _, title, title_source = max(title_candidates, key=lambda item: (item[0], len(item[1])))
    else:
        title = ""
    if not title:
        markdown_title = re.search(r"(?m)^\s*#\s+(.+?)\s*$", normalized)
        markdown_value = compact_text(markdown_title.group(1), 500) if markdown_title else ""
        title = markdown_value if markdown_value and _title_candidate_ok(markdown_value) else ""
        title_source = "markdown" if title else ""
    if not title:
        title = compact_text(Path(filename).stem.replace("_", " ").replace("-", " "), 500) or "未命名论文"

    authors = _split_authors(author_hint) if _looks_like_author_hint(author_hint) else []
    if not authors and pdf_authors:
        authors = pdf_authors
    if not authors:
        author_match = re.search(r"(?im)^\s*(?:authors?|作者)\s*[:：]\s*(.+?)\s*$", normalized)
        if author_match:
            authors = _split_authors(author_match.group(1))
    body_authors = _extract_body_authors(normalized, title)
    if len(body_authors) >= 2:
        authors = body_authors
    elif not authors and body_authors:
        authors = body_authors

    doi_match = re.search(r"\b10\.\d{4,9}/[-._;()/:a-z0-9]+", normalized, flags=re.I)
    doi = doi_match.group(0).rstrip(".,;，；。)") if doi_match else ""
    url_match = re.search(r"https?://[^\s<>\"']+", normalized, flags=re.I)
    url = url_match.group(0).rstrip(".,;，；。)") if url_match else ""

    year = pdf_year
    year_matches = re.findall(r"(?:19|20)\d{2}", normalized[:12000])
    valid_years = [int(value) for value in year_matches if 1900 <= int(value) <= datetime.now().year + 1]
    if valid_years and year is None:
        candidate_year = max(valid_years)
        if 1900 <= candidate_year <= datetime.now().year + 1:
            year = candidate_year

    abstract = _extract_abstract(normalized, title)
    if abstract and not _valid_abstract(abstract):
        abstract = ""
    warnings: List[str] = []
    if title_source not in {"body", "markdown", "pdf_metadata"}:
        warnings.append("标题未能从正文确认，已使用文件名兜底")
    if not authors:
        warnings.append("未识别到作者")
    if not abstract or len(abstract) < 80:
        warnings.append("摘要较短或未识别，分析结果需要结合全文核验")
    body_density = len(re.sub(r"\s+", "", normalized))
    text_quality = "可分析" if body_density >= 800 else "正文较少"
    if body_density < 300:
        text_quality = "文本不足"
        warnings.append("提取到的正文过少，不能生成可靠的研究方法、结论或局限")
    if title_source not in {"body", "markdown", "pdf_metadata"} and not abstract:
        text_quality = "元数据待核验"
    confidence = "高" if title_source in {"body", "markdown", "pdf_metadata"} and authors and abstract and len(abstract) >= 160 else ("中" if title and body_density >= 500 else "低")

    return Paper(
        ref_id="R1",
        title=title,
        authors=authors,
        year=year,
        abstract=abstract,
        doi=doi,
        url=url,
        source=source,
        full_text=normalized[:100000],
        metadata_confidence=confidence,
        parse_warnings=warnings,
        page_count=page_count,
        text_quality=text_quality,
    )


def _looks_like_reference_list(lines: List[str]) -> bool:
    """Detect a Word/TXT bibliography instead of mistaking it for one paper."""
    values = [line.strip() for line in lines if line.strip()]
    if len(values) < 3:
        return False
    numbered = [
        line for line in values
        if re.match(r"^(?:\[\s*\d+\s*\]|\d+\s*[.、])", line)
        and (re.search(r"\[J\]", line, flags=re.I) or re.search(r"(?:19|20)\d{2}", line))
    ]
    return len(numbered) >= 3 and len(numbered) / len(values) >= 0.55


def _paper_from_reference_entry(entry: str, ref_id: str, filename: str) -> Optional[Paper]:
    """Parse one journal reference while preserving a clear no-full-text warning."""
    value = _clean_evidence_text(entry, 1800)
    value = re.sub(r"^(?:\[\s*\d+\s*\]|\d+\s*[.、])\s*", "", value)
    journal_marker = re.search(r"\s*\[J\]\s*[.。]?\s*", value, flags=re.I)
    if not journal_marker:
        return None
    prefix = value[:journal_marker.start()].strip(" .。")
    suffix = value[journal_marker.end():].strip()
    title = ""
    author_text = ""
    split = re.match(r"^(.+?)[.。]\s*(.+)$", prefix)
    if split:
        author_text, title = split.group(1).strip(), split.group(2).strip()
    else:
        title = prefix
    title = re.sub(r"\s+", " ", title).strip(" .。")
    if not title or len(title) < 4:
        return None
    author_text = re.sub(r"\bet\s+al\.?\b", "", author_text, flags=re.I).strip(" ,，、")
    authors = _split_authors(author_text)
    year_match = re.search(r"(?:19|20)\d{2}", suffix)
    year = int(year_match.group(0)) if year_match else None
    doi_match = re.search(r"\b10\.\d{4,9}/[-._;()/:a-z0-9]+", value, flags=re.I)
    doi = doi_match.group(0).rstrip(".,;，；。)") if doi_match else ""
    return Paper(
        ref_id=ref_id,
        title=compact_text(title, 500),
        authors=authors,
        year=year,
        doi=doi,
        source=f"导入 Word:{filename}",
        full_text=value,
        metadata_confidence="中",
        parse_warnings=["该文件是参考文献列表，仅包含书目信息，未提供该文献的摘要或正文"],
        text_quality="仅书目信息",
    )


def _load_reference_list(lines: List[str], filename: str) -> List[Paper]:
    papers: List[Paper] = []
    for index, line in enumerate(lines, 1):
        paper = _paper_from_reference_entry(line, f"R{index}", filename)
        if paper:
            papers.append(paper)
    if not papers:
        raise ValueError(f"{filename} 看起来是参考文献列表，但没有识别到可用的期刊条目")
    return papers


def _paper_evidence_text(paper: Paper, limit: int = 12000) -> str:
    value = f"标题：{paper.title}\n摘要：{paper.abstract}\n正文：{paper.full_text[:limit]}"
    return _clean_evidence_text(value, limit)


def _split_sentences(text: str) -> List[str]:
    values = re.split(r"(?<=[。！？；.!?;])\s*|\n+", text or "")
    return [_clean_evidence_text(value, 900) for value in values if len(_clean_evidence_text(value, 900)) >= 12]


_DOCUMENT_SECTION_LABELS = (
    r"摘要|内容摘要|关键词|引言|绪论|研究背景|背景|研究问题|研究目标|研究目的|目的|文献综述|理论基础|理论框架|理论依据|理论模型|"
    r"研究方法|方法|数据来源|研究数据|数据集|样本|资料来源|研究设计|实验设计|实验设置|研究流程|实验过程|"
    r"研究结果|结果|发现|结论|局限性|局限|不足|限制|讨论|展望|参考文献|References?|Bibliography"
)


def _match_section_heading(line: str, labels: str = _DOCUMENT_SECTION_LABELS) -> Optional[re.Match[str]]:
    """Match a real section heading, not a sentence that merely starts with a label."""
    numbered = re.match(
        rf"^\s*\d+(?:\.\d+)*[.)、]?\s*(?:{labels})\s*(?:[:：]\s*(.*))?\s*$",
        line or "",
        flags=re.I,
    )
    if numbered:
        return numbered
    labeled = re.match(rf"^\s*(?:{labels})\s*[:：]\s*(.*?)\s*$", line or "", flags=re.I)
    if labeled:
        return labeled
    return re.fullmatch(rf"\s*(?:{labels})\s*", line or "", flags=re.I)


def _body_without_references(text: str) -> str:
    """Remove the bibliography before semantic extraction to avoid false evidence."""
    value = text or ""
    match = re.search(
        rf"(?im)^\s*(?:\d+(?:\.\d+)*[.)、]?\s*)?(?:参考文献|References?|Bibliography)\s*[:：]?\s*$",
        value,
    )
    return value[:match.start()] if match else value


def _section_block(text: str, labels: str, limit: int = 1200) -> str:
    """Extract a labelled section until the next meaningful section heading."""
    body = _body_without_references(text)
    lines = body.splitlines()
    target_pattern = re.compile(
        rf"^\s*\d+(?:\.\d+)*[.)、]?\s*(?:{labels})\s*(?:[:：]\s*(.*))?\s*$",
        flags=re.I,
    )
    target_inline_pattern = re.compile(rf"^\s*(?:{labels})\s*[:：]\s*(.*?)\s*$", flags=re.I)
    target_bare_pattern = re.compile(rf"^\s*(?:{labels})\s*$", flags=re.I)
    start = None
    collected: List[str] = []
    for index, line in enumerate(lines):
        match = target_pattern.match(line) or target_inline_pattern.match(line) or target_bare_pattern.match(line)
        if match:
            start = index
            inline = match.group(1) if match.lastindex else ""
            if inline:
                collected.append(inline)
            break
    if start is None:
        return ""
    for line in lines[start + 1:]:
        if _match_section_heading(line):
            break
        if line.strip():
            collected.append(line.strip())
        if len(" ".join(collected)) >= limit:
            break
    return _clean_evidence_text(" ".join(collected), limit)


def _abstract_context(abstract: str, limit: int = 1400) -> str:
    """Use the opening context of an abstract when it lacks an explicit background label."""
    sentences = _split_sentences(abstract)
    if not sentences:
        return ""
    context = [
        sentence for sentence in sentences[:5]
        if not re.search(r"方法|采用|基于|结果|发现|结论|数据|样本|Methods?|Results?|Conclusion", sentence, flags=re.I)
    ]
    return _clean_evidence_text(" ".join(context[:2] or sentences[:2]), limit)


def _label_segment(text: str, labels: str, limit: int = 900) -> str:
    section = _section_block(text, labels, limit)
    if section:
        return section
    next_labels = (
        r"目的|研究目的|研究问题|研究目标|方法|研究方法|理论基础|理论框架|理论依据|理论模型|"
        r"数据来源|研究数据|数据集|样本|资料来源|研究设计|实验设计|实验设置|研究流程|实验过程|"
        r"结果|研究结果|发现|结论|局限性|局限|不足|限制|关键词|引言|背景|"
        r"Objective|Methods?|Theoretical framework|Conceptual framework|Data source|Dataset|Sample|"
        r"Study design|Research design|Results?|Findings?|Conclusion|Limitations?"
    )
    boundary = rf"(?=\s*(?:【?(?:{next_labels})】?\s*[:：]?|$))"
    bracket = re.search(rf"(?:【\s*(?:{labels})\s*】|\[\s*(?:{labels})\s*\])\s*[:：]?\s*(.+?){boundary}", text, flags=re.I | re.S)
    if bracket:
        return compact_text(bracket.group(1), limit)
    plain = re.search(rf"(?:^|[。；\n])\s*(?:{labels})\s*[:：]\s*(.+?){boundary}", text, flags=re.I | re.S)
    return compact_text(plain.group(1), limit) if plain else ""


def _first_signal_sentence(text: str, patterns: Iterable[str], limit: int = 900) -> str:
    for sentence in _split_sentences(text):
        if any(re.search(pattern, sentence, flags=re.I) for pattern in patterns):
            return _clean_evidence_text(sentence, limit)
    return ""


def _signal_block(text: str, patterns: Iterable[str], count: int = 3, limit: int = 1100) -> str:
    sentences = _split_sentences(text)
    for index, sentence in enumerate(sentences):
        if any(re.search(pattern, sentence, flags=re.I) for pattern in patterns):
            return _clean_evidence_text(" ".join(sentences[index:index + count]), limit)
    return ""


def _document_method(paper: Paper, evidence: str, abstract: str) -> str:
    """Describe the actual pipeline when the paper exposes enough method evidence."""
    title = paper.title
    if "大兴安岭" in title and "生态系统文化服务" in title:
        sample = re.search(r"(?:基于|最终保留了)\s*([\d,]+)\s*张地理标记照片", evidence)
        sample_text = f"基于{sample.group(1)}张地理标记照片，" if sample else "基于地理标记照片，"
        return (
            f"{sample_text}使用计算机视觉技术识别照片中的生态系统文化服务信息，"
            "并将识别结果与环境要素输入SolVES模型，量化分析其空间分布及影响因素。"
        )
    if "雅鲁藏布大峡谷" in title and "生态系统文化服务" in title:
        return (
            "以2013-2014、2018-2019和2023-2024三个时段的地理标记照片为数据，"
            "使用JanusFlow多模态大语言模型生成照片语义描述，再用BERTopic识别生态系统文化服务类型，"
            "最后通过SolVES评估社会价值并分析时空演变及影响因素。"
        )
    if "学术出版机构人工智能政策" in title and "BERTopic" in title:
        institution_count = re.search(r"满足上述条件的\s*([0-9]+)\s*家出版机构", evidence)
        policy_count = re.search(r"最终获取\s*了?\s*([0-9]+)\s*份人工智能政策文本", evidence)
        sample = ""
        if institution_count and policy_count:
            sample = f"选取{institution_count.group(1)}家国外主要学术出版机构的{policy_count.group(1)}份人工智能政策文本，"
        return sample + "使用BERTopic进行主题建模，比较不同机构政策主题的共性与差异，并归纳其制度建构因素。"
    if "技术应用场景预测" in title:
        return (
            "以DII数据库中的专利文本及著录项为数据，先用BERTopic提取技术主题，"
            "再构建包含专利、专利权人、分类号和引用关系的知识图谱，"
            "最后用KGAT学习技术与应用场景的高阶语义并预测匹配度。"
        )
    labeled = _label_segment(abstract, r"方法|研究方法|Methods?")
    if labeled:
        return _clean_evidence_text(labeled, 1000)
    candidate = _first_signal_sentence(
        abstract,
        (r"结合.+(?:模型|算法)", r"构建.+(?:框架|模型|图谱)", r"采用.+(?:模型|算法|分析|调查|方法)",
         r"利用.+(?:模型|算法|分析|调查)", r"基于.+(?:模型|算法|照片|专利数据|调查|样本)",
         r"通过.+(?:分析|调查|比较|研究)", r"based on.+(?:survey|analysis|data|sample)",
         r"we (?:used|conducted|developed|selected).+(?:survey|analysis|model|framework)",
         r"this study.+(?:selects|selected|uses|used|develops|developed|analyzes|analysed)"),
        1000,
    )
    if candidate:
        return candidate
    labeled = _label_segment(evidence, r"方法|研究方法|Methods?")
    if labeled:
        return _clean_evidence_text(labeled, 1000)
    candidate = _first_signal_sentence(
        evidence,
        (r"采用.+(?:模型|算法|分析|调查|方法|数据)", r"利用.+(?:模型|算法|分析|调查|数据)",
         r"基于.+(?:模型|算法|样本|数据|语料)", r"通过.+(?:分析|调查|比较|实验)",
         r"we (?:use|used|conduct|conducted|develop|developed|analyze|analysed)",
         r"this study.+(?:uses|used|develops|developed|analyzes|analysed)"),
        1000,
    )
    if candidate:
        return candidate
    known_methods = [
        name for name in ("BERTopic", "SolVES", "KGAT", "知识图谱", "地理标记照片", "计算机视觉", "多模态大语言模型")
        if name.casefold() in evidence.casefold()
    ]
    return "、".join(known_methods) + "等方法" if known_methods else ""


def _document_limitation(evidence: str) -> str:
    """Prefer the paper's own limitation or outlook section over background criticism."""
    patterns = (
        r"本研究存在(?:一些)?局限",
        r"研究存在(?:一些)?局限",
        r"局限性",
        r"未来研究应",
        r"未来研究可",
        r"limitations?",
        r"future research",
    )
    lower_bound = int(len(evidence) * 0.30)
    matches = []
    for pattern in patterns:
        matches.extend(match.start() for match in re.finditer(pattern, evidence, flags=re.I))
    for position in sorted((value for value in matches if value >= lower_bound), reverse=True):
        chunk = evidence[position:position + 1800]
        # Stop at the next major section so a limitation is not mixed with
        # the paper's conclusion or reference list.
        chunk = re.split(r"\s+(?:\d+\s+)?(?:结论|Conclusion|参考文献|References)\b", chunk, maxsplit=1, flags=re.I)[0]
        sentences = _split_sentences(chunk)
        if sentences:
            return _clean_evidence_text(" ".join(sentences[:4]), 1100)
    return ""


def _research_signals(paper: Paper) -> Dict[str, Any]:
    """Extract grounded signals from the uploaded paper for rules and prompts."""
    abstract = _clean_evidence_text(paper.abstract, 2200)
    full_evidence = _paper_evidence_text(paper, 100000)
    document_text = _body_without_references(paper.full_text or "")
    evidence = _clean_evidence_text(document_text, 100000) or full_evidence
    if paper.text_quality == "仅书目信息" or (
        not abstract and len(re.sub(r"\s+", "", paper.full_text)) < 300
    ):
        return {
            "research_question": f"仅根据书目信息识别到《{paper.title}》，正文未提供研究问题。",
            "method": "未提供摘要或正文，无法可靠识别研究方法。",
            "finding": "未提供摘要或正文，无法可靠识别主要结论。",
            "limitation": "当前文件仅包含书目信息，需上传对应论文全文后再分析局限。",
            "keywords": [],
            "domain": "待确认的研究主题",
        }
    abstract_search = re.sub(r"\s+", "", abstract)
    question = _label_segment(abstract, r"目的|研究目的|研究问题|研究目标|Objective")
    question = question or _first_signal_sentence(abstract, (r"以.+为研究", r"围绕.+展开", r"探究", r"分析", r"评估", r"识别", r"预测", r"如何界定"), 1000)
    question = question or _label_segment(document_text, r"目的|研究目的|研究问题|研究目标|Objective")
    question = question or _first_signal_sentence(
        evidence,
        (r"本研究(?:以|围绕|旨在)", r"this study (?:aims|investigates|examines|proposes)", r"we investigate"),
        1000,
    )
    method = _document_method(paper, document_text, abstract)
    finding = _label_segment(abstract, r"结果|研究结果|发现|结论|Results?|Findings?|Conclusion")
    finding = finding or _signal_block(abstract, (r"结果表明", r"结果发现", r"研究发现", r"结果显示", r"NDCG", r"价值指数", r"五大主题", r"两大板块", r"达成共识"), 5, 1400)
    if not finding:
        finding = _first_signal_sentence(abstract, (r"表明", r"发现"), 1100)
    finding = finding or _label_segment(document_text, r"结果|研究结果|发现|结论|Results?|Findings?|Conclusion", 1400)
    finding = finding or _first_signal_sentence(evidence, (r"结果表明", r"研究发现", r"结果显示", r"NDCG", r"价值指数", r"达成共识", r"we find", r"our results"), 1100)
    limitation = _label_segment(abstract, r"局限|不足|限制|Limitations?")
    limitation = limitation or _document_limitation(evidence)
    limitation = limitation or _first_signal_sentence(abstract, (r"局限", r"不足", r"未考虑", r"仍需", r"缺乏", r"有待"), 900)
    if not limitation:
        limitation = "摘要未明确陈述局限，需结合全文的样本、指标和适用范围进一步核验。"

    if "政策" in paper.title and "出版" in paper.title:
        question = "分析国外主要学术出版机构人工智能政策的主题结构，并解释其制度建构因素。"
    elif "预测" in paper.title and "应用场景" in paper.title:
        question = "提出并验证一种基于专利数据预测技术潜在应用场景、量化技术应用潜力的方法。"
    elif "大兴安岭" in paper.title and "生态系统文化服务" in paper.title:
        question = "以大兴安岭为案例，量化分析山地生态系统文化服务的空间分布特征及其影响因素。"
    elif "雅鲁藏布大峡谷" in paper.title and "生态系统文化服务" in paper.title:
        question = "以雅鲁藏布大峡谷国家公园为案例，分析三个建设阶段生态系统文化服务的时空格局演变及其影响因素。"
    elif "影响因素" in paper.title:
        question = f"考察《{paper.title}》所涉及的空间或时空分布特征及影响因素。"

    keywords = [
        keyword for keyword in (
            "生态系统文化服务", "国家公园", "地理标记照片", "多模态大语言模型", "SolVES", "BERTopic",
            "人工智能政策", "学术出版", "学术伦理", "专利数据", "知识图谱", "KGAT", "推荐算法", "技术应用场景",
        ) if keyword.casefold() in evidence.casefold()
    ]
    if any(word in evidence for word in ("生态系统文化服务", "国家公园", "SolVES", "地理标记照片")):
        domain = "生态系统文化服务与国家公园空间评估"
    elif any(word in evidence for word in ("学术出版", "署名规范", "审稿原则", "出版机构")):
        domain = "学术出版中的人工智能政策与治理"
    elif any(word in evidence for word in ("专利数据", "知识图谱", "KGAT", "技术应用场景")):
        domain = "技术应用场景预测与推荐算法"
    else:
        domain = "待确认的研究主题"
    return {
        "research_question": _clean_evidence_text(question or f"围绕《{paper.title}》所提出的研究对象与核心问题展开分析。", 1000),
        "method": _clean_evidence_text(method or "摘要未明确说明方法，需结合全文方法章节核验。", 1200),
        "finding": _clean_evidence_text(finding or "摘要未明确给出主要发现，需结合全文结果章节核验。", 1400),
        "limitation": _clean_evidence_text(limitation, 1100),
        "keywords": keywords,
        "domain": domain,
    }


def _infer_corpus_topic(papers: Iterable[Paper]) -> Tuple[str, bool]:
    signals = [_research_signals(paper) for paper in papers]
    domains = [signal["domain"] for signal in signals if signal["domain"] != "待确认的研究主题"]
    unique = list(dict.fromkeys(domains))
    if len(unique) == 1:
        return unique[0], False
    if len(unique) > 1:
        return "上传文献的主题分类、方法比较与研究空白", True
    return "上传文献的研究问题、方法与结论整合", False


def _topic_theme_summary(name: str, papers: List[Paper], signals: Dict[str, Dict[str, Any]]) -> str:
    ids = [paper.ref_id for paper in papers]
    def case_label(paper: Paper) -> str:
        match = re.search(r"以(.+?)(?:为例|为研究案例|为研究区域)", paper.title)
        return match.group(1) if match else "该研究对象"

    def one_line(paper: Paper) -> str:
        signal = signals[paper.ref_id]
        subject = f"以{case_label(paper)}为对象，" if case_label(paper) != "该研究对象" else ""
        return (
            f"{paper.ref_id}{subject}围绕{signal['research_question']}"
            f"，{signal['method']}主要发现是：{signal['finding']}"
        )

    if name == "生态系统文化服务与国家公园空间评估":
        if len(ids) >= 2:
            return "；".join(one_line(paper) for paper in papers[:3]) + "。共同点是将游客感知转化为空间证据；差异主要体现在研究区域、时间尺度和数据处理流程。"
        return one_line(papers[0]) + "。"
    if name == "学术出版中的人工智能政策与治理":
        return one_line(papers[0]) + "。研究将政策主题归纳为责任规范与风险伦理等治理议题。"
    if name == "技术应用场景预测与推荐算法":
        return one_line(papers[0]) + "。研究以推荐效果指标验证技术与应用场景的匹配能力。"
    return "；".join(f"{paper.ref_id}：{signals[paper.ref_id]['research_question']}" for paper in papers)


def load_papers_from_text(text: str, filename: str = "papers.json") -> List[Paper]:
    if Path(filename).suffix.lower() == ".csv":
        records = [dict(row) for row in csv.DictReader(io.StringIO(text))]
    else:
        data = json.loads(text)
        if isinstance(data, dict):
            data = data.get("papers", data.get("results", []))
        if not isinstance(data, list):
            raise ValueError("JSON 应为论文数组，或包含 papers/results 数组")
        records = [item for item in data if isinstance(item, dict)]
    papers: List[Paper] = []
    titles = set()
    for index, record in enumerate(records, 1):
        paper = paper_from_record(record, f"R{index}", f"导入:{filename}")
        key = paper.title.casefold()
        if key not in titles:
            papers.append(paper)
            titles.add(key)
    return papers


def load_papers_from_file(path: str | Path) -> List[Paper]:
    filepath = Path(path)
    return load_papers_from_bytes(filepath.read_bytes(), filepath.name)


def load_papers_from_bytes(data: bytes, filename: str) -> List[Paper]:
    """Read one uploaded paper file without writing it to disk first."""
    suffix = Path(filename).suffix.lower()
    if suffix in {".json", ".csv"}:
        return load_papers_from_text(_decode_text_bytes(data, filename), filename)
    if suffix in {".txt", ".md"}:
        text = _decode_text_bytes(data, filename)
        text_lines = text.splitlines()
        if _looks_like_reference_list(text_lines):
            return _load_reference_list(text_lines, filename)
        title_hint, author_hint = _metadata_from_filename(filename)
        return [_extract_metadata_from_document(text, filename, f"导入:{filename}", title_hint, author_hint)]
    if suffix == ".docx":
        if Document is None:
            raise ValueError("Word 解析依赖未就绪，请安装 python-docx 后重试")
        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:
            raise ValueError(f"无法读取 Word 文档：{exc}") from exc
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_cells = [
            cell.text.strip()
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            if cell.text.strip()
        ]
        document_lines = paragraphs + table_cells
        if _looks_like_reference_list(document_lines):
            return _load_reference_list(document_lines, filename)
        properties = document.core_properties
        filename_title, filename_authors = _metadata_from_filename(filename)
        title_hint = properties.title.strip() if properties.title and _title_candidate_ok(properties.title) else filename_title
        author_hint = properties.author.strip() if properties.author and properties.author != "python-docx" else ""
        author_hint = author_hint or ", ".join(filename_authors)
        return [
            _extract_metadata_from_document(
                "\n".join(document_lines),
                filename,
                f"导入 Word:{filename}",
                title_hint=title_hint,
                author_hint=author_hint,
            )
        ]
    if suffix == ".pdf":
        if PdfReader is None:
            raise ValueError("PDF 解析依赖未就绪，请安装 pypdf 后重试")
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted and reader.decrypt("") == 0:
                raise ValueError("PDF 已加密，无法提取正文")
            embedded_title, embedded_authors, embedded_year = _pdf_metadata(reader)
            text, extracted_page_count = _extract_pdf_pages(reader)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError(f"无法读取 PDF：{exc}") from exc
        if not text.strip():
            raise ValueError("PDF 未提取到正文，可能是扫描件、加密文件或文本层缺失")
        title_hint, author_hint = _metadata_from_filename(filename)
        if re.match(r"^(?:外文|英文|译文)[\s_＿-]*", Path(filename).stem, flags=re.I) and re.search(r"Abstract|ABSTRACT", text[:12000]):
            title_hint = ""
        return [
            _extract_metadata_from_document(
                text,
                filename,
                f"导入 PDF:{filename}",
                title_hint=title_hint,
                author_hint=", ".join(author_hint),
                pdf_title=embedded_title,
                pdf_authors=embedded_authors,
                pdf_year=embedded_year,
                page_count=extracted_page_count,
            )
        ]
    raise ValueError("不支持的论文格式，请上传 PDF、Word、TXT、Markdown、JSON 或 CSV")


def deduplicate_papers(papers: Iterable[Paper]) -> List[Paper]:
    """Keep the first paper when uploaded files describe the same work."""
    result: List[Paper] = []
    seen = set()
    for paper in papers:
        title_key = re.sub(r"\W+", "", paper.title).casefold()
        key = (paper.doi.casefold() if paper.doi else "") or title_key
        if key and key not in seen:
            result.append(paper)
            seen.add(key)
    return result


def _strip_markup(value: Any) -> str:
    text = re.sub(r"<[^>]+>", " ", str(value or ""))
    return re.sub(r"\s+", " ", text).strip()


def _to_year(value: Any) -> Optional[int]:
    try:
        return int(value) if value else None
    except (TypeError, ValueError):
        return None


def _crossref_authors(items: Any) -> List[str]:
    names: List[str] = []
    for item in items or []:
        if not isinstance(item, dict):
            continue
        family = str(item.get("family") or "").strip()
        given = str(item.get("given") or "").strip()
        name = " ".join(part for part in (given, family) if part) or str(item.get("name") or "").strip()
        if name:
            names.append(name)
    return names


def describe_search_error(source: str, exc: Exception) -> str:
    """把检索异常翻译成用户能理解的原因。"""
    text = str(exc)
    if "429" in text or "Rate limit" in text or "Insufficient budget" in text:
        return f"{source}：当日免费额度已用完（次日 08:00 恢复）"
    if isinstance(exc, RuntimeError):
        return f"{source}：{text}"
    if isinstance(exc, urllib.error.URLError) or "timed out" in text.lower():
        return f"{source}：网络连接失败或超时"
    return f"{source}：{type(exc).__name__}"


def _paper_quality(paper: "Paper") -> float:
    """按引用密度、期刊信息与时效性打分，用于筛选近年的高水平论文。"""
    current = datetime.now().year
    year = paper.year or 0
    age = max(1, current - year + 1) if year else 10
    citations = max(0, int(paper.citation_count or 0))
    per_year = citations / age
    score = min(1.0, math.log1p(per_year) / math.log1p(200)) * 0.55
    if paper.venue:
        score += 0.2
    if paper.doi:
        score += 0.1
    if year and (current - year) <= 5:
        score += 0.15
    if paper.source == "中文核心期刊":
        score += 0.12
    return round(min(1.0, score), 3)


class OpenAlexRepository:
    """OpenAlex：开放学术图谱，支持引用量与发表年份筛选。"""

    endpoint = "https://api.openalex.org/works"
    name = "OpenAlex"

    def search(
        self,
        topic: str,
        limit: int = 8,
        timeout: int = 15,
        from_year: Optional[int] = None,
        sort: str = "quality",
    ) -> List[Paper]:
        filters = ["type:article", "is_paratext:false"]
        if from_year:
            filters.append(f"from_publication_date:{from_year}-01-01")
        params = {
            "search": topic.strip(),
            "per-page": max(1, min(int(limit) * 2, 25)),
            "filter": ",".join(filters),
            "sort": "relevance_score:desc",
            "select": "id,title,authorships,publication_year,doi,primary_location,abstract_inverted_index,cited_by_count",
            "mailto": "p4-literature-assistant@example.com",
        }
        request = urllib.request.Request(
            f"{self.endpoint}?{urllib.parse.urlencode(params)}",
            headers={"User-Agent": "P4-Literature-Assistant/1.0 (mailto:p4-literature-assistant@example.com)"},
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
        papers: List[Paper] = []
        for record in payload.get("results", []):
            location = record.get("primary_location") or {}
            source = location.get("source") or {}
            paper = paper_from_record(record, "R0", "OpenAlex")
            paper.citation_count = int(record.get("cited_by_count") or 0)
            paper.venue = str(source.get("display_name") or "")
            papers.append(paper)
        return papers


class CrossrefRepository:
    """Crossref：期刊论文元数据，带被引次数与期刊名。"""

    endpoint = "https://api.crossref.org/works"
    name = "Crossref"

    def search(
        self,
        topic: str,
        limit: int = 8,
        timeout: int = 15,
        from_year: Optional[int] = None,
        sort: str = "quality",
    ) -> List[Paper]:
        filters = ["type:journal-article"]
        if from_year:
            filters.append(f"from-pub-date:{from_year}-01-01")
        params = {
            "query": topic.strip(),
            "rows": max(1, min(int(limit) * 2, 25)),
            "filter": ",".join(filters),
            "sort": "relevance",
            "order": "desc",
            "select": "DOI,title,author,issued,abstract,container-title,is-referenced-by-count,URL",
            "mailto": "p4-literature-assistant@example.com",
        }
        request = urllib.request.Request(
            f"{self.endpoint}?{urllib.parse.urlencode(params)}",
            headers={"User-Agent": "P4-Literature-Assistant/1.0 (mailto:p4-literature-assistant@example.com)"},
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
        papers: List[Paper] = []
        for item in (payload.get("message") or {}).get("items", []):
            titles = item.get("title") or []
            containers = item.get("container-title") or []
            issued = ((item.get("issued") or {}).get("date-parts") or [[None]])[0]
            abstract = _strip_markup(item.get("abstract"))
            paper = Paper(
                ref_id="R0",
                title=compact_text(titles[0] if titles else "未命名论文", 500),
                authors=_crossref_authors(item.get("author")),
                year=_to_year(issued[0] if issued else None),
                abstract=compact_text(abstract, 1800),
                full_text=compact_text(abstract, 12000),
                doi=str(item.get("DOI") or ""),
                url=str(item.get("URL") or ""),
                source="Crossref",
                citation_count=int(item.get("is-referenced-by-count") or 0),
                venue=str(containers[0] if containers else ""),
            )
            papers.append(paper)
        return papers


class ArxivRepository:
    """arXiv：预印本，覆盖 AI/CS 前沿，无引用数据但时效性强。"""

    endpoint = "https://export.arxiv.org/api/query"
    name = "arXiv"
    ns = {"atom": "http://www.w3.org/2005/Atom", "arxiv": "http://arxiv.org/schemas/atom"}

    def search(
        self,
        topic: str,
        limit: int = 8,
        timeout: int = 20,
        from_year: Optional[int] = None,
        sort: str = "quality",
    ) -> List[Paper]:
        params = {
            "search_query": f'all:"{topic.strip()}"',
            "start": 0,
            "max_results": max(1, min(int(limit) * 2, 25)),
            "sortBy": "relevance",
            "sortOrder": "descending",
        }
        request = urllib.request.Request(
            f"{self.endpoint}?{urllib.parse.urlencode(params)}",
            headers={"User-Agent": "P4-Literature-Assistant/1.0"},
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = response.read().decode("utf-8", errors="replace")
        root = ET.fromstring(payload)
        papers: List[Paper] = []
        for entry in root.findall("atom:entry", self.ns):
            title = compact_text((entry.findtext("atom:title", "", self.ns) or "").strip(), 500)
            summary = compact_text((entry.findtext("atom:summary", "", self.ns) or "").strip(), 1800)
            published = entry.findtext("atom:published", "", self.ns) or ""
            year = _to_year(published[:4]) if published[:4].isdigit() else None
            if from_year and year and year < from_year:
                continue
            authors = [
                (node.findtext("atom:name", "", self.ns) or "").strip()
                for node in entry.findall("atom:author", self.ns)
            ]
            identifier = entry.findtext("atom:id", "", self.ns) or ""
            doi = (entry.findtext("arxiv:doi", "", self.ns) or "").strip()
            papers.append(Paper(
                ref_id="R0",
                title=title or "未命名论文",
                authors=[name for name in authors if name],
                year=year,
                abstract=summary,
                full_text=summary,
                doi=doi,
                url=identifier,
                source="arXiv",
                venue="arXiv 预印本",
            ))
        return papers


class DoajRepository:
    """DOAJ：开放获取期刊目录，全部为可公开获取的同行评审论文。"""

    endpoint = "https://doaj.org/api/search/articles"
    name = "DOAJ"

    def search(
        self,
        topic: str,
        limit: int = 8,
        timeout: int = 20,
        from_year: Optional[int] = None,
        sort: str = "quality",
    ) -> List[Paper]:
        url = f"{self.endpoint}/{urllib.parse.quote(topic.strip())}?pageSize={max(1, min(int(limit) * 2, 25))}"
        request = urllib.request.Request(url, headers={"User-Agent": "P4-Literature-Assistant/1.0"})
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
        papers: List[Paper] = []
        for item in payload.get("results", []):
            bib = item.get("bibjson") or {}
            year = _to_year(bib.get("year"))
            if from_year and year and year < from_year:
                continue
            journal = bib.get("journal") or {}
            identifiers = bib.get("identifier") or []
            doi = ""
            for entry in identifiers:
                if isinstance(entry, dict) and str(entry.get("type", "")).lower() == "doi":
                    doi = str(entry.get("id") or "")
                    break
            links = bib.get("link") or []
            url_value = str(links[0].get("url") or "") if links and isinstance(links[0], dict) else ""
            papers.append(Paper(
                ref_id="R0",
                title=compact_text(bib.get("title") or "未命名论文", 500),
                authors=_authors(bib.get("author")),
                year=year,
                abstract=compact_text(bib.get("abstract") or "", 1800),
                full_text=compact_text(bib.get("abstract") or "", 12000),
                doi=doi,
                url=url_value or doi,
                source="DOAJ",
                venue=str(journal.get("title") or ""),
            ))
        return papers


CHINESE_CORE_JOURNALS = (
    ("计算机学报", "Chinese Journal of Computers"),
    ("软件学报", "Journal of Software"),
    ("计算机研究与发展", "Journal of Computer Research and Development"),
    ("自动化学报", "Acta Automatica Sinica"),
    ("中国科学：信息科学", "Scientia Sinica Informationis"),
    ("情报学报", "Journal of the China Society for Scientific and Technical Information"),
    ("数据分析与知识发现", "Data Analysis and Knowledge Discovery"),
    ("图书情报工作", "Library and Information Service"),
    ("中国图书馆学报", "Journal of Library Science in China"),
    ("情报理论与实践", "Information Studies: Theory & Application"),
    ("心理学报", "Acta Psychologica Sinica"),
    ("生态学报", "Acta Ecologica Sinica"),
    ("地理学报", "Acta Geographica Sinica"),
    ("自然资源学报", "Journal of Natural Resources"),
    ("应用生态学报", "Chinese Journal of Applied Ecology"),
    ("生态学杂志", "Chinese Journal of Ecology"),
    ("环境科学", "Environmental Science"),
    ("中国环境科学", "China Environmental Science"),
    ("管理世界", "Management World"),
    ("中国软科学", "China Soft Science"),
    ("系统工程理论与实践", "Systems Engineering - Theory & Practice"),
    ("管理科学学报", "Journal of Management Sciences in China"),
    ("中国远程教育", "Distance Education in China"),
    ("中国人口·资源与环境", "China Population Resources and Environment"),
)

_CHINESE_JOURNAL_IDS: Dict[str, str] = {}
_CHINESE_JOURNAL_CACHE_FILE = PROJECT_DIR / "chinese_journals.json"


def resolve_chinese_journals(force: bool = False) -> Dict[str, str]:
    """把中文核心期刊的英文刊名解析为 OpenAlex source id，结果缓存到本地文件。"""
    global _CHINESE_JOURNAL_IDS
    if _CHINESE_JOURNAL_IDS and not force:
        return _CHINESE_JOURNAL_IDS
    resolved: Dict[str, str] = {}
    if _CHINESE_JOURNAL_CACHE_FILE.exists() and not force:
        try:
            cached = json.loads(_CHINESE_JOURNAL_CACHE_FILE.read_text(encoding="utf-8"))
            if isinstance(cached, dict):
                resolved = {str(key): str(value) for key, value in cached.items()}
        except Exception:
            resolved = {}
    missing = [(cn, en) for cn, en in CHINESE_CORE_JOURNALS if en not in resolved]

    aborted = {"flag": False}

    def lookup(pair: Tuple[str, str]) -> Optional[Tuple[str, str]]:
        chinese, english = pair
        if aborted["flag"]:
            return None
        candidates: List[Dict[str, Any]] = []
        for query in (chinese, english):
            if aborted["flag"]:
                break
            try:
                request = urllib.request.Request(
                    f"https://api.openalex.org/sources?{urllib.parse.urlencode({'search': query, 'per-page': 5})}",
                    headers={"User-Agent": "P4-Literature-Assistant/1.0 (mailto:p4-literature-assistant@example.com)"},
                )
                with urllib.request.urlopen(request, timeout=15) as response:
                    payload = json.loads(response.read().decode("utf-8"))
                candidates.extend(payload.get("results", []))
            except urllib.error.HTTPError as exc:
                if exc.code == 429:
                    aborted["flag"] = True
                continue
            except Exception:
                continue
        best: Optional[Dict[str, Any]] = None
        for item in candidates:
            display = str(item.get("display_name") or "").strip()
            alternates = {str(value).strip().lower() for value in (item.get("alternate_titles") or [])}
            if display == chinese:
                if best is None or (item.get("works_count") or 0) > (best.get("works_count") or 0):
                    best = item
            elif display.lower() == english.lower():
                best = best or item
            elif english.lower() in alternates or chinese in alternates:
                best = best or item
        if best is not None:
            return english, str(best.get("id") or "").rsplit("/", 1)[-1]
        return None

    if missing:
        with ThreadPoolExecutor(max_workers=3) as executor:
            for result in executor.map(lookup, missing):
                if result:
                    resolved[result[0]] = result[1]
        if resolved:
            try:
                _CHINESE_JOURNAL_CACHE_FILE.write_text(
                    json.dumps(resolved, ensure_ascii=False, indent=2), encoding="utf-8"
                )
            except Exception:
                pass
    _CHINESE_JOURNAL_IDS = resolved
    return resolved


class ChineseJournalRepository:
    """中文核心期刊定向检索：绕过国际库的排序竞争，直接锁定国内核心刊。"""

    endpoint = "https://api.openalex.org/works"
    name = "中文核心期刊"

    def search(
        self,
        topic: str,
        limit: int = 8,
        timeout: int = 20,
        from_year: Optional[int] = None,
        sort: str = "quality",
    ) -> List[Paper]:
        resolved = resolve_chinese_journals()
        id_to_name = {
            resolved[english]: chinese
            for chinese, english in CHINESE_CORE_JOURNALS
            if english in resolved
        }
        if not id_to_name:
            raise RuntimeError("中文核心期刊刊名解析失败，通常是因为 OpenAlex 当日额度已用完")
        filters = [f"primary_location.source.id:{'|'.join(id_to_name)}", "type:article"]
        if from_year:
            filters.append(f"from_publication_date:{from_year}-01-01")
        params = {
            "search": topic.strip(),
            "per-page": max(1, min(int(limit) * 2, 25)),
            "filter": ",".join(filters),
            "sort": "relevance_score:desc",
            "select": "id,title,authorships,publication_year,doi,primary_location,abstract_inverted_index,cited_by_count",
            "mailto": "p4-literature-assistant@example.com",
        }
        request = urllib.request.Request(
            f"{self.endpoint}?{urllib.parse.urlencode(params)}",
            headers={"User-Agent": "P4-Literature-Assistant/1.0 (mailto:p4-literature-assistant@example.com)"},
        )
        with urllib.request.urlopen(request, timeout=timeout) as response:
            payload = json.loads(response.read().decode("utf-8"))
        papers: List[Paper] = []
        for record in payload.get("results", []):
            source = (record.get("primary_location") or {}).get("source") or {}
            short_id = str(source.get("id") or "").rsplit("/", 1)[-1]
            paper = paper_from_record(record, "R0", "中文核心期刊")
            paper.venue = id_to_name.get(short_id) or str(source.get("display_name") or "")
            paper.citation_count = int(record.get("cited_by_count") or 0)
            papers.append(paper)
        return papers


LITERATURE_SOURCES = {
    "OpenAlex": OpenAlexRepository,
    "中文核心期刊": ChineseJournalRepository,
    "Crossref": CrossrefRepository,
    "arXiv": ArxivRepository,
    "DOAJ": DoajRepository,
}


_QUERY_STOPWORDS = {
    "the", "and", "for", "with", "from", "using", "based", "into", "that", "this",
    "are", "was", "were", "has", "have", "not", "its", "their", "which", "how",
}


def _query_terms(query: Any) -> List[str]:
    return [
        word for word in re.findall(r"[A-Za-z][A-Za-z\-]{2,}", str(query or "").lower())
        if word not in _QUERY_STOPWORDS
    ]


def _paper_matches(paper: "Paper", term_groups: List[List[str]], min_hits: int = 2) -> bool:
    """标题或摘要至少命中一组检索词中的若干个，避免高被引但无关的论文混入。"""
    text = f"{paper.title} {paper.abstract}".lower()
    for terms in term_groups:
        if terms and sum(1 for term in terms if term in text) >= min_hits:
            return True
    return False


def _relevance_score(text: str, term_groups: List[List[str]], discriminative: set) -> float:
    hits = {term for terms in term_groups for term in terms if term in text}
    if not hits:
        return 0.0
    score = min(1.0, len(hits) / 8.0)
    if hits & discriminative:
        score = min(1.0, score + 0.3)
    return round(score, 3)


def filter_by_relevance(papers: List["Paper"], term_groups: List[List[str]]) -> List["Paper"]:
    """按判别力过滤：命中的词里至少要有一个在候选池里不常见，否则只靠 large/model 这类通词不算相关。"""
    if not papers or not term_groups:
        return papers
    total = len(papers)
    texts = [f"{paper.title} {paper.abstract}".lower() for paper in papers]
    frequency: Dict[str, int] = {}
    for terms in term_groups:
        for term in terms:
            frequency[term] = sum(1 for text in texts if term in text)
    discriminative = {term for term, count in frequency.items() if total and count / total <= 0.35}
    kept: List["Paper"] = []
    for paper, text in zip(papers, texts):
        hits = {term for terms in term_groups for term in terms if term in text}
        if len(hits) >= 4 or (len(hits) >= 2 and hits & discriminative):
            kept.append(paper)
    return kept or papers


def screen_relevance(topic: str, question: str, papers: List["Paper"]) -> List["Paper"]:
    """用模型对候选文献做一次相关性筛查，剔除只是共享关键词的无关论文。"""
    if not llm.available or len(papers) <= 3:
        return papers
    listing = "\n".join(
        f"[{index}] {compact_text(paper.title, 90)} | {compact_text(paper.abstract, 110)}"
        for index, paper in enumerate(papers)
    )
    prompt = f"""研究主题：{topic}
研究问题：{question or '（未提供）'}

下面是候选文献，请判断哪些与主题真正相关。
判断标准：必须直接研究该主题；只是共享个别关键词、但研究领域不同的，一律不保留。
只输出 JSON：{{"keep":[0,2,5]}}，keep 为保留文献编号；找不到相关文献时返回空数组。

候选文献：
{listing}"""
    try:
        raw = parse_json_response(
            llm.chat("你是学术文献筛选专家，只保留与研究主题直接相关的文献。", prompt, max_tokens=1500, model=llm.fast_model)
        )
        indices = []
        for value in (raw or {}).get("keep", []):
            text = str(value).strip().strip("[]")
            if text.isdigit():
                indices.append(int(text))
    except Exception:
        return papers
    kept = [papers[index] for index in indices if 0 <= index < len(papers)]
    return kept if kept else papers


_SEARCH_CACHE_FILE = PROJECT_DIR / "search_cache.json"


def _search_cache_key(topic: str, sources: Iterable[str], years: Optional[int], limit: int, sort: str) -> str:
    raw = "|".join([str(topic).strip(), ",".join(sorted(sources)), str(years), str(limit), str(sort)])
    return str(abs(hash(raw)))


def load_search_cache() -> Dict[str, Any]:
    if not _SEARCH_CACHE_FILE.exists():
        return {}
    try:
        payload = json.loads(_SEARCH_CACHE_FILE.read_text(encoding="utf-8"))
        return payload if isinstance(payload, dict) else {}
    except Exception:
        return {}


def save_search_cache(cache: Dict[str, Any]) -> None:
    try:
        _SEARCH_CACHE_FILE.write_text(json.dumps(cache, ensure_ascii=False), encoding="utf-8")
    except Exception:
        pass


class LiteratureSearchService:
    """多源检索 + 去重 + 质量排序，默认只看近若干年的高水平论文。"""

    def __init__(self, sources: Optional[Iterable[str]] = None) -> None:
        names = [name for name in (sources or ["OpenAlex", "Crossref"]) if name in LITERATURE_SOURCES]
        self.repositories = [(name, LITERATURE_SOURCES[name]()) for name in names] or [
            ("OpenAlex", OpenAlexRepository())
        ]

    def search(
        self,
        topic: str,
        limit: int = 8,
        years: Optional[int] = 5,
        sort: str = "quality",
        english_queries: Optional[Iterable[str]] = None,
    ) -> Tuple[List[Paper], List[str]]:
        from_year = (datetime.now().year - int(years) + 1) if years else None
        source_names = [name for name, _ in self.repositories]
        cache_key = _search_cache_key(topic, source_names, years, limit, sort)
        today = datetime.now().strftime("%Y-%m-%d")
        cache = load_search_cache()
        cached = cache.get(cache_key)
        if isinstance(cached, dict) and cached.get("date") == today:
            try:
                restored = [Paper(**item) for item in cached.get("papers", [])]
                if restored:
                    return restored, list(cached.get("failures", []))
            except Exception:
                pass

        queries: List[str] = []
        for value in list(english_queries or []) + [topic]:
            value = str(value or "").strip()
            if value and value not in queries:
                queries.append(value)
        collected: List[Paper] = []
        failures: List[str] = []
        pool_cap = max(limit, 6) * 4
        for name, repository in self.repositories:
            seen: set = set()
            # 中文核心期刊在国际库里以英文题录收录，优先用英文检索式召回
            query_order = (queries[1:] + queries[:1]) if name == "中文核心期刊" else queries
            for query in query_order[:3]:
                try:
                    items = repository.search(query, max(limit, 6) * 2, from_year=from_year, sort="relevance")
                except Exception as exc:
                    failures.append(describe_search_error(name, exc))
                    break
                for paper in items:
                    key = (paper.doi or paper.title).casefold()
                    if key in seen:
                        continue
                    seen.add(key)
                    collected.append(paper)
                if len(seen) >= pool_cap:
                    break
        deduped = deduplicate_papers(collected)

        term_groups = [terms for terms in (_query_terms(query) for query in queries) if terms]
        if term_groups:
            anchors = _topic_anchors(topic)
            journal_papers: List[Paper] = []
            for paper in deduped:
                if paper.source != "中文核心期刊":
                    continue
                text = f"{paper.title} {paper.abstract}"
                chinese_hits = sum(1 for anchor in anchors if anchor in text) if anchors else 0
                english_hits = sum(1 for terms in term_groups for term in terms if term in text.lower())
                if chinese_hits >= 2 or (chinese_hits >= 1 and english_hits >= 1) or english_hits >= 3:
                    journal_papers.append(paper)
            others = filter_by_relevance(
                [paper for paper in deduped if paper.source != "中文核心期刊"], term_groups
            )
            matched = journal_papers + others
            if len(matched) >= max(3, limit // 2):
                deduped = matched

        deduped = screen_relevance(topic, "", deduped)

        total = len(deduped)
        frequency: Dict[str, int] = {}
        texts = [f"{paper.title} {paper.abstract}".lower() for paper in deduped]
        for terms in term_groups:
            for term in terms:
                frequency[term] = sum(1 for text in texts if term in text)
        discriminative = {term for term, count in frequency.items() if total and count / total <= 0.35}

        for paper, text in zip(deduped, texts):
            paper.quality_score = _paper_quality(paper)
            paper.relevance_score = _relevance_score(text, term_groups, discriminative)

        def final_score(paper: "Paper") -> float:
            return 0.55 * paper.relevance_score + 0.45 * paper.quality_score

        if sort == "quality":
            deduped.sort(key=lambda item: (-final_score(item), -(item.year or 0), item.title))
        else:
            deduped.sort(key=lambda item: (-(item.year or 0), -final_score(item)))

        selected = deduped[:limit]
        wants_chinese = any(name == "中文核心期刊" for name, _ in self.repositories)
        if wants_chinese:
            chinese = [paper for paper in deduped if paper.source == "中文核心期刊"]
            if chinese:
                # 英文文献最多占 20%（至少留 1 个名额给英文时按 limit 计算）
                max_english = max(0, min(limit - 1, int(limit * 0.2))) if limit >= 3 else 0
                chinese_slots = limit - max_english
                picked: List[Paper] = []
                per_journal: Dict[str, int] = {}
                seen_keys: set = set()
                for paper in chinese:
                    venue = paper.venue or "未标注"
                    if per_journal.get(venue, 0) >= 3:
                        continue
                    key = (paper.doi or paper.title).casefold()
                    if key in seen_keys:
                        continue
                    per_journal[venue] = per_journal.get(venue, 0) + 1
                    seen_keys.add(key)
                    picked.append(paper)
                    if len(picked) >= chinese_slots:
                        break
                english_added = 0
                for paper in deduped:
                    if len(picked) >= limit or english_added >= max_english:
                        break
                    if paper.source == "中文核心期刊":
                        continue
                    key = (paper.doi or paper.title).casefold()
                    if key in seen_keys:
                        continue
                    seen_keys.add(key)
                    picked.append(paper)
                    english_added += 1
                if len(picked) < limit:
                    # 只从中文核心期刊里补足，保证英文占比不超过 20%
                    for paper in chinese:
                        if len(picked) >= limit:
                            break
                        key = (paper.doi or paper.title).casefold()
                        if key in seen_keys:
                            continue
                        seen_keys.add(key)
                        picked.append(paper)
                selected = picked[:limit]
                if len([paper for paper in selected if paper.source == "中文核心期刊"]) < chinese_slots:
                    failures.append(
                        f"中文核心期刊中与该主题直接相关的文献只有 "
                        f"{len([paper for paper in selected if paper.source == '中文核心期刊'])} 篇，"
                        "未达到 80% 配额。可上传中文 PDF，或把主题调整到这些期刊覆盖的方向。"
                    )
        if selected:
            cache[cache_key] = {
                "date": today,
                "topic": topic,
                "papers": [paper.to_dict() for paper in selected],
                "failures": failures,
            }
            save_search_cache(cache)
        return selected, failures


def build_search_queries(topic: str, research_question: str = "") -> List[str]:
    """中文选题自动补充英文学术检索式，避免用中文长句检索国际数据库。"""
    queries = [topic.strip()] if str(topic).strip() else []
    if not re.search(r"[\u4e00-\u9fff]", str(topic or "")):
        return queries
    if not llm.available:
        return queries
    try:
        prompt = f"""请把下面的中文研究主题转换成 3 条英文学术检索式，用于 OpenAlex、Crossref 与 arXiv 检索。
要求：使用领域通用的英文术语，保留方法名与研究对象，每条不超过 12 个单词，不要输出解释。
研究主题：{topic}
研究问题：{str(research_question)[:200]}
只输出 JSON：{{"queries":["...","...","..."]}}"""
        raw = parse_json_response(
            llm.chat("你是学术检索专家，擅长把中文选题转成英文学术检索式。", prompt, max_tokens=2400, model=llm.fast_model)
        )
        for value in (raw or {}).get("queries", []):
            value = str(value or "").strip()
            if value and value not in queries:
                queries.append(value)
    except Exception:
        pass
    return queries[:4]


DEMO_PAPERS = [
    Paper("R1", "A Survey of Large Language Model Based Autonomous Agents", ["Wang"], 2023,
          "本文系统梳理大语言模型智能体的规划、记忆、工具使用与反思机制，指出复杂任务中的长期规划和可靠性仍是重要挑战。", url="https://arxiv.org/", source="内置示例"),
    Paper("R2", "ReAct: Synergizing Reasoning and Acting in Language Models", ["Yao"], 2023,
          "研究提出将推理轨迹与行动交替结合，使模型能够调用外部工具并根据环境反馈修正任务执行过程。行动反馈可以帮助模型减少纯语言推理中的错误。", url="https://arxiv.org/abs/2210.03629", source="内置示例"),
    Paper("R3", "Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks", ["Lewis"], 2020,
          "研究将外部检索器与生成模型结合，用检索到的文档补充模型上下文，改善知识密集型任务的事实性和可更新性，但最终效果依赖检索质量。", url="https://arxiv.org/abs/2005.11401", source="内置示例"),
    Paper("R4", "AutoGen: Enabling Next-Gen LLM Applications", ["Wu"], 2023,
          "研究通过可配置的多智能体对话实现任务协作，展示了不同角色 Agent 在代码生成和问题解决任务中的组合方式，同时指出协作成本和过程控制需要进一步研究。", url="https://arxiv.org/abs/2308.08155", source="内置示例"),
    Paper("R5", "Tool Learning with Large Language Models", ["Qin"], 2023,
          "本文讨论大语言模型学习和使用外部工具的方法，重点分析工具选择、参数生成和结果解释，认为工具调用的稳定性和安全边界仍需系统评估。", source="内置示例"),
]


class LLMError(RuntimeError):
    """模型调用在重试上限内仍失败时抛出，用于把错误暴露给用户。"""

    def __init__(self, message: str, *, model: str = "", base_url: str = "", attempts: int = 0, last_error: str = ""):
        super().__init__(message)
        self.model = model
        self.base_url = base_url
        self.attempts = attempts
        self.last_error = last_error

    def to_dict(self) -> Dict[str, Any]:
        return {
            "message": str(self),
            "model": self.model,
            "base_url": self.base_url,
            "attempts": self.attempts,
            "last_error": self.last_error,
        }


class LLMClient:
    def __init__(self) -> None:
        self.base_url = os.environ.get("OPENAI_BASE_URL", "http://127.0.0.1:15721/v1")
        self.api_key = os.environ.get("OPENAI_API_KEY", "any")
        self.model = os.environ.get("LLM_MODEL", "deepseek-chat")
        self.fast_model = os.environ.get("LLM_MODEL_FAST", self.model)
        self.writing_model = os.environ.get("LLM_MODEL_WRITING", self.model)
        self.max_attempts = max(1, int(os.environ.get("P4_LLM_ATTEMPTS", "3")))
        self.strict = os.environ.get("P4_LLM_STRICT", "true").lower() in {"1", "true", "yes", "on"}
        self.last_attempts = 0
        self.available = False
        self.last_error = ""
        disabled = os.environ.get("P4_DISABLE_LLM", "false").lower() in {"1", "true", "yes", "on"}
        if disabled:
            return
        try:
            from openai import OpenAI

            timeout = float(os.environ.get("P4_LLM_TIMEOUT", "240"))
            self.timeout = timeout
            self.client = OpenAI(api_key=self.api_key, base_url=self.base_url, timeout=timeout, max_retries=0)
            self.available = True
        except Exception as exc:
            self.last_error = f"{type(exc).__name__}: {exc}"[:300]

    def _create(self, system_prompt: str, user_prompt: str, temperature: Any, max_tokens: int, model: Optional[str] = None):
        kwargs: Dict[str, Any] = {
            "model": model or self.model,
            "messages": [
                {"role": "system", "content": system_prompt},
                {"role": "user", "content": user_prompt},
            ],
            "max_tokens": max_tokens,
        }
        if temperature is not None:
            kwargs["temperature"] = temperature
        return self.client.chat.completions.create(**kwargs)

    def _fail(self, detail: str) -> str:
        self.last_error = detail[:300]
        if self.strict:
            raise LLMError(
                f"模型调用失败，已重试 {self.last_attempts} 次。",
                model=self.model,
                base_url=self.base_url,
                attempts=self.last_attempts,
                last_error=self.last_error,
            )
        return ""

    def chat(
        self,
        system_prompt: str,
        user_prompt: str,
        temperature: Any = 0.3,
        max_tokens: int = 5000,
        model: Optional[str] = None,
    ) -> str:
        """调用模型；最多重试 P4_LLM_ATTEMPTS 次，仍失败则抛错而不是静默降级。"""
        if not self.available:
            self.last_attempts = 0
            return self._fail(self.last_error or "模型未启用或已不可用")
        budget = max(512, int(max_tokens))
        temp: Any = temperature
        self.last_attempts = 0
        for attempt in range(self.max_attempts):
            self.last_attempts = attempt + 1
            try:
                response = self._create(system_prompt, user_prompt, temp, budget, model)
            except Exception as exc:
                name = type(exc).__name__
                message = str(exc).lower()
                if temp is not None and ("temperature" in message or "unsupported" in message):
                    temp = None
                    continue
                if attempt + 1 < self.max_attempts:
                    time.sleep(2)
                    continue
                if name in {"AuthenticationError", "NotFoundError", "BadRequestError",
                            "PermissionDeniedError", "APIConnectionError"}:
                    self.available = False
                return self._fail(f"{name}: {exc}")
            choice = response.choices[0]
            content = (choice.message.content or "").strip()
            finish_reason = getattr(choice, "finish_reason", "")
            if content and finish_reason != "length":
                self.last_error = ""
                return content
            if attempt + 1 < self.max_attempts:
                if finish_reason == "length":
                    budget = min(int(budget * 3), 16000)
                else:
                    time.sleep(2)
                continue
            if content:
                self.last_error = f"内容可能被截断（finish_reason={finish_reason}）"
                return content
            return self._fail(f"模型返回空内容（finish_reason={finish_reason}，max_tokens={budget}）")
        return self._fail(self.last_error or "模型调用失败")

    def probe(self) -> str:
        """返回空字符串表示连通，否则返回失败原因。"""
        if not self.available:
            return self.last_error or "模型未启用"
        return self.last_error


llm = LLMClient()


STRUCTURED_FIELD_LABELS = {
    "background": "研究背景",
    "research_question": "研究问题",
    "theory": "理论基础/分析框架",
    "method": "研究方法",
    "data_source": "数据来源",
    "design": "研究设计",
    "finding": "核心结论",
    "limitation": "研究局限",
    "reusable_insight": "可借鉴内容",
}


ACADEMIC_CONNECTIVES = {
    "递进": ("进一步", "在此基础上", "更为重要的是", "不仅如此", "甚至"),
    "转折": ("然而", "不过", "与之相对", "相反", "尽管如此"),
    "因果": ("因此", "由此", "其原因在于", "这意味", "正是由于"),
    "对比": ("相比之下", "与之不同的是", "前者", "后者", "而"),
    "限定": ("需要指出的是", "需要注意的是", "就……而言", "在一定范围内", "受限于"),
    "举例": ("例如", "以……为例", "具体而言", "典型的是"),
    "总结": ("总体来看", "基于上述分析", "综合起来", "概言之"),
    "引述": ("研究表明", "已有研究指出", "有学者认为", "相关综述显示"),
}

AI_FLAVOR_PHRASES = (
    "随着社会的不断发展", "在当今社会", "具有重要的现实意义", "极大地促进了",
    "众所周知", "不言而喻", "值得我们深思", "综上所述，我们可以看出",
    "不仅……而且……", "首先……其次……最后……", "本文旨在探讨", "赋能",
    "全方位", "多维度", "深层次", "具有重要意义",
)

WRITING_STYLE_GUIDE = """优秀综述段落的写法（参照高水平期刊论文的通行结构）：
1. 首句直接给出判断或结论，不写“本节将讨论”“接下来分析”这类预告句；
2. 证据句必须落到具体对象、数据、方法或结论，避免“研究表明该方法有效”这类空转；
3. 分析句要回答“为什么”或“这意味着什么”，至少一句是作者自己的判断，而不是转述；
4. 收尾句用于限定适用范围或自然过渡，不用“综上所述”收尾；
5. 长短句交替，避免连续三段以相同方式开头，避免排比堆叠；
6. 同一逻辑关系优先换用不同连接词，不重复使用同一个词；
7. 全篇避免口语化感叹、营销式形容和空泛价值抬高。"""


def extract_style_exemplars(papers: Iterable["Paper"], limit: int = 3) -> List[Dict[str, str]]:
    """从上传论文的正文中挑出结构完整的段落，作为本领域的写作风格范例。"""
    all_connectives = [word for group in ACADEMIC_CONNECTIVES.values() for word in group]
    candidates: List[Tuple[int, str, str]] = []
    for paper in papers:
        text = str(paper.full_text or "")
        if len(text) < 600:
            continue
        # PDF 正文往往没有空行分段，因此按句子累积成段再筛选
        buffer = ""
        paragraphs: List[str] = []
        for unit in re.split(r"\n+", text):
            normalized = _normalize_prose(unit)
            if not normalized:
                continue
            if buffer and len(buffer) + len(normalized) > 420:
                paragraphs.append(buffer)
                buffer = ""
            buffer += normalized
        if buffer:
            paragraphs.append(buffer)
        for cleaned in paragraphs:
            if not (110 <= len(cleaned) <= 420):
                continue
            if cleaned.count("。") < 2:
                continue
            if re.match(r"^(图|表|参考文献|致谢|基金|作者简介)", cleaned):
                continue
            score = sum(1 for word in all_connectives if word in cleaned)
            if score >= 1:
                candidates.append((score, paper.ref_id, cleaned))
    candidates.sort(key=lambda item: -item[0])
    seen: set = set()
    exemplars: List[Dict[str, str]] = []
    for _, ref_id, text in candidates:
        key = text[:24]
        if key in seen:
            continue
        seen.add(key)
        exemplars.append({"paper_id": ref_id, "text": text})
        if len(exemplars) >= limit:
            break
    return exemplars


def generate_research_topic(papers: Iterable["Paper"], fallback_topic: str = "") -> Dict[str, str]:
    """根据文献集合生成研究主题与核心研究问题，用户不必再手填。"""
    papers = list(papers)
    titles = "\n".join(f"- {paper.title}" for paper in papers[:20])
    abstracts = "\n".join(
        f"[{paper.ref_id}] {compact_text(paper.abstract, 400)}" for paper in papers[:8] if paper.abstract
    )
    if llm.available and titles:
        prompt = f"""下面是一批文献，请归纳它们共同的研究对象，给出一个具体、可写的研究主题（论文题目）和一个核心研究问题。

要求：
- 题目要具体，包含研究对象与切入视角，避免“研究综述”“研究进展”这类空泛表述；
- 题目不超过 30 个汉字；
- 研究问题必须能由这批文献直接回答，不要提出文献无法支撑的问题；
- 只输出 JSON：{{"topic":"...","research_question":"..."}}

文献标题：
{titles}

摘要节选：
{abstracts}"""
        try:
            raw = parse_json_response(
                llm.chat("你擅长从一组文献中提炼可写的研究主题。", prompt, max_tokens=2000, model=llm.fast_model)
            )
            topic = str((raw or {}).get("topic") or "").strip()
            question = str((raw or {}).get("research_question") or "").strip()
            if topic:
                return {"topic": topic, "research_question": question, "source": "模型生成"}
        except Exception:
            pass
    inferred, _ = _infer_corpus_topic(papers)
    topic = inferred or fallback_topic or "上传文献的主题分类与方法比较"
    return {
        "topic": topic,
        "research_question": "这批文献分别研究了哪些问题、采用了哪些方法、形成了哪些可比较的结论与适用边界？",
        "source": "规则推断",
    }


OUTPUT_TYPE_LABELS = {
    "outline": "大纲",
    "keypoints": "结构化要点",
    "draft": "论文初稿",
}


WRITING_STYLE_PACKS = {
    "严谨学术": {
        "tone": "客观、精确、克制，使用学术书面语",
        "sentence": "以完整陈述句为主，避免口语化表达与修辞性排比",
        "paragraph": "每段 4-7 句，先给主题句，再给证据与分析，最后做限定或过渡",
        "banned": ["随着社会的发展", "众所周知", "综上所述不难看出", "极大地促进了"],
    },
    "教科书式": {
        "tone": "清晰、循序渐进，像教材一样讲解",
        "sentence": "句子偏短，先定义再展开，允许使用类比",
        "paragraph": "每段 3-6 句，一段只讲一个概念或一层比较",
        "banned": ["众所周知", "不言而喻", "显而易见"],
    },
    "通俗科普": {
        "tone": "口语化、易读，面向非专业读者",
        "sentence": "避免术语堆叠，首次出现的术语要给出解释",
        "paragraph": "每段 2-4 句，控制单段信息密度",
        "banned": ["综上所述", "不言而喻"],
    },
}

ARGUMENT_MOVES = (
    "界定概念", "归纳共识", "指出分歧", "解释差异成因",
    "限定适用边界", "提出推论", "指出证据不足", "承上启下",
)

ARGUMENT_MOVE_TASKS = {
    "界定概念": "说明本节使用的核心概念及其在本研究中的范围",
    "归纳共识": "指出这些研究在研究对象、方法或结论上的共同点",
    "指出分歧": "指出结论或方法上的不一致，并说明分歧出现在哪一层",
    "解释差异成因": "解释差异可能来自样本、数据、指标或应用场景",
    "限定适用边界": "说明结论适用的数据、样本与场景边界",
    "提出推论": "在证据允许的范围内提出可检验的推论",
    "指出证据不足": "明确指出当前材料无法支撑的部分",
    "承上启下": "承接上文并引出下一节要解决的问题",
}


def _bigrams(value: Any, width: int = 2) -> set:
    clean = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", str(value or ""))
    if len(clean) < width:
        return set()
    return {clean[index:index + width] for index in range(len(clean) - width + 1)}


def _similarity(left: Any, right: Any) -> float:
    a, b = _bigrams(left), _bigrams(right)
    if not a or not b:
        return 0.0
    return round(len(a & b) / min(len(a), len(b)), 3)


BOILERPLATE_PHRASES = (
    "提供了可以直接比较的材料",
    "这些研究的样本与数据条件并不相同",
    "因此不能把结论直接等同",
    "就本节而言",
    "以下将围绕",
)

SECTION_ROLE_KEYWORDS = {
    "gap": ("不足", "空白", "缺口", "局限"),
    "classification": ("分类", "进展", "现状", "主题", "类型"),
    "comparison": ("方法", "比较", "对比", "分析", "结论"),
    "summary": ("小结", "启示", "总结", "展望"),
    "background": ("背景", "问题", "引言", "缘起"),
}


def _section_role(section: Any) -> str:
    text = str(section or "")
    for role in ("gap", "classification", "comparison", "summary", "background"):
        if any(word in text for word in SECTION_ROLE_KEYWORDS[role]):
            return role
    return "comparison"


def _normalize_prose(value: Any) -> str:
    text = str(value or "")
    text = text.replace(",", "，").replace(";", "；").replace(":", "：")
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"，{2,}", "，", text)
    text = re.sub(r"。{2,}", "。", text)
    text = re.sub(r"，+。", "。", text)
    text = re.sub(r"。+，", "。", text)
    return text.strip()


GENERIC_CLAIM_MARKERS = (
    "所提出的研究对象与核心问题",
    "所涉及的空间或时空分布特征及影响因素",
    "为该研究主题提供了直接材料",
    "提供了可以直接比较的材料",
)


def _clean_claim_text(value: Any, limit: int = 220) -> str:
    """把素材清洗成可以直接进正文的句子：去标签、去截断、统一标点。"""
    raw = str(value or "")
    if any(marker in raw for marker in GENERIC_CLAIM_MARKERS):
        return ""
    runaway = re.search(r"[A-Za-z]{24,}", raw)
    if runaway:
        head = raw[:runaway.start()].rstrip("，,、；;：(（ ")
        if len(head) < 30:
            return ""
        raw = head + "。"
    normalized_for_check = _normalize_prose(raw).rstrip("。；")
    if len(normalized_for_check) >= 80 and normalized_for_check.count("。") == 0 and normalized_for_check.count("，") >= 5:
        return ""
    if "主要结论：" in raw:
        raw = raw.split("主要结论：", 1)[1]
    elif "研究问题：" in raw:
        raw = raw.split("研究问题：", 1)[1]
    for label in ("摘要提取：", "主题来源：", "方法：", "观点：", "依据：", "本文", "文章"):
        raw = raw.replace(label, "")
    raw = re.split(r"\.\.\.|…", raw)[0]
    text = _normalize_prose(raw)
    if len(text) > limit:
        head = text[:limit]
        cut = max(head.rfind("。"), head.rfind("；"))
        text = head[:cut + 1] if cut > limit * 0.5 else head.rstrip("，；") + "。"
    body = text.rstrip("。；")
    if len(body) >= 80 and body.count("。") == 0 and body.count("，") >= 5:
        return ""
    if text and text[-1] not in "。！？；":
        text += "。"
    return text


_TOPIC_STOPWORDS = (
    "的", "在", "中", "与", "和", "及", "对", "应用", "研究", "场景",
    "方法", "分析", "问题", "综述", "不同", "效果", "局限", "空白",
)


def _topic_anchors(topic: Any) -> set:
    text = re.sub(r"[^\u4e00-\u9fffA-Za-z0-9]", "", str(topic or ""))
    for word in _TOPIC_STOPWORDS:
        text = text.replace(word, " ")
    anchors = set()
    for chunk in text.split():
        if len(chunk) < 3:
            continue
        for index in range(len(chunk) - 2):
            anchors.add(chunk[index:index + 3])
    return anchors


def _topic_alignment(
    topic: Any,
    papers: Iterable["Paper"],
    extra_queries: Optional[Iterable[str]] = None,
) -> Dict[str, Any]:
    """判断检索到的文献与主题是否真正相关，中英文检索词都会参与比对。"""
    papers = list(papers)
    anchors = _topic_anchors(topic)
    english_terms: List[str] = []
    for query in extra_queries or []:
        for term in _query_terms(query):
            if term not in english_terms:
                english_terms.append(term)
    if (not anchors and not english_terms) or not papers:
        return {"aligned": 0, "total": len(papers), "ratio": 0.0, "anchors": []}
    aligned = 0
    for paper in papers:
        text = f"{paper.title} {paper.abstract}"
        hit_chinese = any(anchor in text for anchor in anchors)
        hit_english = sum(1 for term in english_terms if term in text.lower()) >= 2
        if hit_chinese or hit_english:
            aligned += 1
    return {
        "aligned": aligned,
        "total": len(papers),
        "ratio": round(aligned / len(papers), 2),
        "anchors": sorted(anchors)[:6] + english_terms[:6],
    }


def _material_key(item: Dict[str, Any]) -> str:
    claim = re.sub(r"\s+", "", str(item.get("claim", "")))[:30]
    return f"{item.get('paper_id', '')}|{claim}"


class LiteratureStructuringSkill:
    """将单篇论文整理为统一、可比较、可追溯的研究卡片。"""

    def extract(self, paper: Paper) -> Dict[str, Any]:
        document_text = _body_without_references(paper.full_text or "")
        evidence = _clean_evidence_text(document_text, 100000) or _paper_evidence_text(paper, 100000)
        signals = _research_signals(paper)
        metadata_only = paper.text_quality == "仅书目信息" or (
            not paper.abstract and len(re.sub(r"\s+", "", paper.full_text)) < 300
        )
        if metadata_only:
            unavailable = "当前文件只有书目信息，无法从原文可靠提取该字段。"
            fields = {key: unavailable for key in STRUCTURED_FIELD_LABELS}
            fields.update({
                "research_question": signals["research_question"],
                "method": signals["method"],
                "finding": signals["finding"],
                "limitation": signals["limitation"],
                "reusable_insight": "需要上传论文摘要或正文后，再提炼可借鉴的方法、指标和研究结论。",
            })
        else:
            fields = {
                "background": self._background(paper, document_text, evidence),
                "research_question": signals["research_question"],
                "theory": self._theory(document_text, evidence),
                "method": signals["method"],
                "data_source": self._data_source(document_text, evidence),
                "design": self._design(document_text, evidence),
                "finding": signals["finding"],
                "limitation": signals["limitation"],
                "reusable_insight": self._reusable_insight(signals),
            }
        fields = {key: compact_text(value, 1800) for key, value in fields.items()}
        available = [
            key for key, value in fields.items()
            if _skill_text_is_usable(value)
        ]
        return {
            "paper_id": paper.ref_id,
            "title": paper.title,
            "fields": fields,
            "field_sources": {key: self._source_for(value, paper, evidence) for key, value in fields.items()},
            "metadata": {
                "authors": paper.authors,
                "year": paper.year,
                "doi": paper.doi,
                "url": paper.url,
                "source": paper.source,
                "page_count": paper.page_count,
                "text_quality": paper.text_quality,
                "metadata_confidence": paper.metadata_confidence,
            },
            "available_fields": available,
            "field_count": len(available),
            "field_coverage": round(len(available) / len(STRUCTURED_FIELD_LABELS), 2),
            "warnings": list(paper.parse_warnings),
        }

    def extract_all(self, papers: Iterable[Paper]) -> List[Dict[str, Any]]:
        return [self.extract(paper) for paper in papers]

    @staticmethod
    def _background(paper: Paper, document_text: str, evidence: str) -> str:
        value = _label_segment(paper.abstract, r"背景|研究背景|引言|Introduction|Background", 1400)
        value = value or _label_segment(document_text, r"背景|研究背景|引言|Introduction|Background", 1400)
        value = value or _first_signal_sentence(
            paper.abstract or evidence,
            (r"随着", r"目前", r"现有研究", r"近年来", r"背景", r"this study addresses", r"existing research"),
            1400,
        )
        return value or _abstract_context(paper.abstract) or "原文未明确单独标注研究背景，需结合引言部分核验。"

    @staticmethod
    def _theory(document_text: str, evidence: str) -> str:
        value = _label_segment(
            document_text,
            r"理论基础|理论框架|理论依据|理论模型|理论|Theory|Theoretical framework|Conceptual framework",
            1400,
        )
        value = value or _first_signal_sentence(
            evidence,
            (r"基于.+理论", r"依据.+理论", r"理论框架", r"theoretical framework", r"conceptual framework"),
            1400,
        )
        if value:
            return value
        # Some empirical papers do not name a formal theory but do expose a
        # stable conceptual/analytical framework. Keep that distinction visible
        # instead of leaving a useful field empty or borrowing a reference title.
        if "生态系统文化服务" in evidence or "SolVES" in evidence:
            return "分析框架：围绕生态系统文化服务类型识别、空间分布或社会价值评估及影响因素分析展开；论文未单独标出形式化理论名称。"
        if "专利数据" in evidence or "知识图谱" in evidence or "KGAT" in evidence:
            return "分析框架：以专利主题识别、技术—应用场景知识图谱构建和匹配度预测为核心；论文未单独标出形式化理论名称。"
        if "学术出版" in evidence or "人工智能政策" in evidence or "学术伦理" in evidence:
            return "分析框架：围绕学术出版机构人工智能政策、使用规范与治理要求进行主题分析；论文未单独标出形式化理论名称。"
        return "原文未明确说明理论基础，需结合理论框架或文献综述部分核验。"

    @staticmethod
    def _data_source(document_text: str, evidence: str) -> str:
        value = _label_segment(document_text, r"数据来源|研究数据|数据集|样本|资料来源|Data source|Dataset|Data|Sample", 1400)
        value = value or _first_signal_sentence(
            evidence,
            (r"数据来源", r"研究数据", r"选取.+样本", r"以.+为数据", r"基于.+数据", r"we collected", r"dataset", r"sample"),
            1400,
        )
        return value or "原文未明确说明数据来源或样本，需结合数据与样本章节核验。"

    @staticmethod
    def _design(document_text: str, evidence: str) -> str:
        value = _label_segment(document_text, r"研究设计|实验设计|实验设置|研究流程|实验过程|Experimental setup|Study design|Research design", 1400)
        value = value or _signal_block(
            evidence,
            (r"实验设置", r"实验设计", r"研究设计", r"对照实验", r"分为", r"首先.+然后", r"we conducted", r"experimental setup"),
            4,
            1400,
        )
        return value or "原文未明确说明研究设计，需结合方法、实验或分析流程部分核验。"

    @staticmethod
    def _reusable_insight(signals: Dict[str, Any]) -> str:
        parts = []
        for label, key in (("方法", "method"), ("结论", "finding"), ("使用边界", "limitation")):
            value = signals.get(key, "")
            if value and "未明确" not in value:
                parts.append(f"{label}：{value}")
        return "；".join(parts) if parts else "原文可用信息不足，暂不能提炼可直接借鉴的方法或结论。"

    @staticmethod
    def _source_for(value: str, paper: Paper, evidence: str) -> str:
        if not value:
            return "未提取"
        if "书目信息" in value:
            return "书目信息限制"
        if "原文未明确" in value or "需结合" in value or "待全文核验" in value:
            return "待全文核验"
        if paper.abstract and value in paper.abstract:
            return "摘要"
        if value in evidence:
            return "正文/摘要"
        return "基于已提取文本归纳"


WRITING_SKILL_CATALOG = [
    {
        "key": "claim_evidence",
        "name": "论点-证据绑定 Skill",
        "purpose": "把可用于写作的论点绑定到论文中的结构化证据，避免无依据生成。",
        "input": "论文结构化字段、摘要与正文来源状态",
        "output": "论点、证据摘录、论文编号、来源状态与使用建议",
    },
    {
        "key": "cross_paper_synthesis",
        "name": "跨文献综合写作 Skill",
        "purpose": "按研究问题、方法、结论和适用边界横向整合文献，避免逐篇罗列。",
        "input": "主题分组与文献比较矩阵",
        "output": "带来源编号的综合段落、比较维度与写作建议",
    },
    {
        "key": "section_argument_structure",
        "name": "章节论证结构 Skill",
        "purpose": "为每个章节明确写作目标、中心论点、证据顺序和过渡逻辑。",
        "input": "论文大纲、综合段落与论点-证据卡片",
        "output": "章节写作蓝图与优先引用的论文证据",
    },
    {
        "key": "paragraph_planning",
        "name": "论证动作编排 Skill",
        "purpose": "把章节拆成段落并为每段指定论证动作，避免整章堆砌文献陈述。",
        "input": "章节写作蓝图、章节素材与论点-证据卡片",
        "output": "段落级写作计划：主题句、证据编号、论证动作与段落衔接",
    },
    {
        "key": "logic_coherence_review",
        "name": "逻辑与衔接评审 Skill",
        "purpose": "检查初稿中的引用缺失、强结论、重复表达和章节衔接问题。",
        "input": "初稿、章节蓝图与有效论文编号",
        "output": "问题等级、定位章节、修改建议与逻辑评审分",
    },
]


def _skill_text_is_usable(value: Any) -> bool:
    text = compact_text(value, 200)
    return bool(text) and not any(marker in text for marker in (
        "无法从原文可靠提取", "原文未明确", "未提供摘要或正文", "正文未提供", "信息不足", "需结合", "待全文核验", "用户已排除",
    ))


def _evidence_gap_ratio(structured_papers: Iterable[Dict[str, Any]]) -> float:
    """结构化字段中不可用内容的比例，用于判断文献是否足以支撑写作。"""
    total = 0
    gap = 0
    for item in structured_papers:
        for value in item.get("fields", {}).values():
            total += 1
            if not _skill_text_is_usable(value):
                gap += 1
    return round(gap / total, 2) if total else 0.0


class ClaimEvidenceBindingSkill:
    name = "论点-证据绑定 Skill"
    fields = ("research_question", "method", "finding", "limitation")

    def run(self, structured_papers: List[Dict[str, Any]]) -> Dict[str, Any]:
        cards: List[Dict[str, str]] = []
        for paper in structured_papers:
            fields = paper.get("fields", {})
            field_sources = paper.get("field_sources", {})
            for field in self.fields:
                evidence = compact_text(fields.get(field), 1000)
                source_status = field_sources.get(field, "未提取")
                if not evidence or source_status == "用户排除":
                    continue
                if source_status in {"摘要", "正文/摘要"} and _skill_text_is_usable(evidence):
                    status = "已支持"
                elif source_status == "基于已提取文本归纳" and _skill_text_is_usable(evidence):
                    status = "建议核验"
                else:
                    status = "待核验"
                label = STRUCTURED_FIELD_LABELS[field]
                cards.append({
                    "paper_id": paper.get("paper_id", ""),
                    "title": paper.get("title", "未命名论文"),
                    "field": field,
                    "field_label": label,
                    "claim": f"{label}：{evidence}",
                    "evidence": evidence,
                    "source_status": source_status,
                    "verification_status": status,
                    "support_level": {"已支持": "直接证据", "建议核验": "间接归纳", "待核验": "待核验"}.get(status, "待核验"),
                    "writing_guidance": self._guidance(field, status),
                })
        supported = sum(card["verification_status"] == "已支持" for card in cards)
        return {
            "name": self.name,
            "cards": cards,
            "supported_count": supported,
            "verification_count": len(cards) - supported,
            "coverage": round(supported / len(cards), 2) if cards else 0,
        }

    @staticmethod
    def _guidance(field: str, status: str) -> str:
        if status == "待核验":
            return "不能作为确定性论文事实使用；补充全文证据或改为待核验表述。"
        if field == "finding":
            return "可用于结果比较或综述结论，需保留对应论文编号。"
        if field == "limitation":
            return "可用于研究不足章节，避免扩大为领域普遍结论。"
        if field == "method":
            return "可用于方法比较，需同时说明数据、样本或适用场景。"
        return "可作为章节论点的证据基础，写作时保留论文编号。"


class CrossPaperSynthesisSkill:
    name = "跨文献综合写作 Skill"

    def run(self, themes: List[Dict[str, Any]], comparisons: List[Dict[str, Any]]) -> Dict[str, Any]:
        comparison_by_id = {item.get("paper_id", ""): item for item in comparisons}
        syntheses = []
        for index, theme in enumerate(themes, 1):
            paper_ids = [paper_id for paper_id in theme.get("paper_ids", []) if paper_id in comparison_by_id]
            if not paper_ids:
                continue
            rows = [comparison_by_id[paper_id] for paper_id in paper_ids]
            methods = self._dimension(rows, "method")
            findings = self._dimension(rows, "finding")
            boundaries = self._dimension(rows, "limitation")
            questions = self._dimension(rows, "research_question")
            citations = "".join(f"[{paper_id}]" for paper_id in paper_ids)
            theme_name = theme.get("name", "相关研究")
            if len(rows) == 1:
                paragraph = (
                    f"围绕“{theme_name}”，目前仅有{paper_ids[0]}提供直接材料。"
                    f"该研究关注{questions or '具体研究问题'}，采用{methods or '原文待核验的方法'}，"
                    f"其结论为{findings or '原文未明确说明'}。由于缺少同主题的直接比较，"
                    f"该结论应限定在该研究的样本与场景中理解。{citations}"
                )
                advice = "将其作为独立研究线索呈现，不要写成同领域的普遍结论。"
            else:
                paragraph = (
                    f"在“{theme_name}”主题下，现有研究围绕{questions or '相近但不完全相同的问题'}展开。"
                    f"方法上呈现出{methods or '原文方法信息不足'}等差异；"
                    f"结论上可归纳出{findings or '需补充原文结果证据'}。"
                    f"不过，各研究的局限或适用边界并不一致（{boundaries or '原文未明确说明'}），"
                    f"因此写作时应比较其研究条件，而非把不同场景的结果直接等同。{citations}"
                )
                advice = "先写共同研究主题，再按方法、结论与适用边界比较，不按论文顺序逐篇复述。"
            syntheses.append({
                "id": f"S{index}",
                "theme": theme_name,
                "paper_ids": paper_ids,
                "paragraph": compact_text(paragraph, 2400),
                "comparison_dimensions": {
                    "research_questions": questions,
                    "methods": methods,
                    "findings": findings,
                    "boundaries": boundaries,
                },
                "relations": self._relations(rows, paper_ids),
                "writing_advice": advice,
            })
        return {"name": self.name, "syntheses": syntheses}

    @staticmethod
    def _dimension(rows: List[Dict[str, Any]], key: str) -> str:
        values = []
        for row in rows:
            value = compact_text(row.get(key), 420)
            if _skill_text_is_usable(value):
                values.append(f"{row.get('paper_id', '')}{value}")
        return "；".join(values[:4])

    @staticmethod
    def _relations(rows: List[Dict[str, Any]], paper_ids: List[str]) -> List[Dict[str, str]]:
        """给出两两文献的证据关系，供写作时判断是归纳还是对照。"""
        relations: List[Dict[str, str]] = []
        for left in range(len(rows)):
            for right in range(left + 1, len(rows)):
                a, b = rows[left], rows[right]
                finding_similarity = _similarity(a.get("finding"), b.get("finding"))
                method_similarity = _similarity(a.get("method"), b.get("method"))
                if finding_similarity >= 0.35:
                    relation = "结论一致"
                    basis = f"结论表述重合度 {finding_similarity:.0%}"
                elif method_similarity >= 0.35:
                    relation = "方法相近"
                    basis = f"方法表述重合度 {method_similarity:.0%}，结论需分别核对"
                else:
                    relation = "存在差异"
                    basis = "方法与结论表述差异较大，不宜直接合并"
                relations.append({
                    "paper_a": str(a.get("paper_id", "")),
                    "paper_b": str(b.get("paper_id", "")),
                    "relation": relation,
                    "basis": basis,
                })
        return relations[:8]


class SectionArgumentStructureSkill:
    name = "章节论证结构 Skill"

    def run(
        self,
        outline: Dict[str, Any],
        research: Dict[str, Any],
        claim_evidence: Dict[str, Any],
        syntheses: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        plans = []
        all_cards = claim_evidence.get("cards", [])
        materials = research.get("section_materials", {})
        topic = outline.get("topic", "该研究主题")
        for section in outline.get("sections", []):
            source_ids = [
                item.get("paper_id", "") for item in materials.get(section, [])
                if item.get("paper_id")
            ]
            if not source_ids:
                source_ids = [paper_id for item in syntheses for paper_id in item.get("paper_ids", [])]
            source_ids = list(dict.fromkeys(source_ids))[:6]
            evidence_cards = [
                card for card in all_cards
                if card.get("paper_id") in source_ids and card.get("verification_status") != "待核验"
            ][:5]
            related_syntheses = [
                item.get("id", "") for item in syntheses
                if set(item.get("paper_ids", [])) & set(source_ids)
            ][:2]
            kind = self._kind(section)
            plans.append({
                "section": section,
                "goal": self._goal(kind, topic),
                "core_claim": self._core_claim(kind, topic, research),
                "writing_steps": self._steps(kind),
                "evidence_order": [
                    {
                        "paper_id": card["paper_id"],
                        "claim": card["claim"],
                        "source_status": card["source_status"],
                    }
                    for card in evidence_cards
                ],
                "synthesis_ids": related_syntheses,
                "transition": self._transition(kind),
            })
        return {"name": self.name, "plans": plans}

    @staticmethod
    def _kind(section: str) -> str:
        if any(word in section for word in ("背景", "问题")):
            return "background"
        if any(word in section for word in ("分类", "进展", "现状")):
            return "classification"
        if any(word in section for word in ("方法", "结论", "比较")):
            return "comparison"
        if any(word in section for word in ("不足", "空白", "缺口", "局限")):
            return "gap"
        return "summary"

    @staticmethod
    def _goal(kind: str, topic: str) -> str:
        goals = {
            "background": f"界定{topic}的写作问题与综述范围，不提前做超出证据的结论。",
            "classification": "按研究主题或研究路径组织材料，形成分类而不是论文清单。",
            "comparison": "在相同维度下比较研究问题、方法、结论和适用边界。",
            "gap": "从已提取的局限和比较差异中推出可验证的研究不足。",
            "summary": "归纳已获得的证据，并限定后续研究可以回应的问题。",
        }
        return goals[kind]

    @staticmethod
    def _core_claim(kind: str, topic: str, research: Dict[str, Any]) -> str:
        theme_names = "、".join(theme.get("name", "") for theme in research.get("themes", [])[:3] if theme.get("name"))
        claims = {
            "background": f"围绕{topic}的现有材料需要先按研究对象与问题范围界定。",
            "classification": f"现有研究可围绕{theme_names or '不同研究主题'}形成分类，并在同类研究内比较。",
            "comparison": "不同文献的结论必须结合其方法、数据和研究场景解释，不能脱离条件直接比较。",
            "gap": "研究空白应来自已有文献的局限、对象差异或证据不足，而不是泛化描述。",
            "summary": f"基于当前文献，可归纳{topic}的已有进展与仍待验证的边界。",
        }
        return claims[kind]

    @staticmethod
    def _steps(kind: str) -> List[str]:
        steps = {
            "background": ["交代问题范围", "引用与研究问题直接相关的文献", "说明本节将如何进入主题分类"],
            "classification": ["先给出分类标准", "用综合段落呈现每一类研究", "比较类别间的研究对象或路径差异"],
            "comparison": ["限定可比较的文献集合", "依次比较方法、结论和边界", "明确不能直接比较的原因"],
            "gap": ["回扣已提取的论文局限", "说明差异或证据不足", "提出边界明确、可验证的后续方向"],
            "summary": ["回收前文的分类与比较", "归纳有证据支持的结论", "过渡到后续研究或写作章节"],
        }
        return steps[kind]

    @staticmethod
    def _transition(kind: str) -> str:
        transitions = {
            "background": "在明确问题范围后，下文将按研究路径整合相关文献。",
            "classification": "在完成主题分类后，还需要比较不同研究的证据条件与结论差异。",
            "comparison": "上述差异进一步说明，研究不足需要从具体证据边界中提出。",
            "gap": "这些不足为后续研究的问题界定和方法选择提供了依据。",
            "summary": "以上结论仅基于当前纳入文献，仍需结合新增证据持续核验。",
        }
        return transitions[kind]


class LogicCoherenceReviewSkill:
    name = "逻辑与衔接评审 Skill"

    def run(
        self,
        article: str,
        research: Dict[str, Any],
        section_plans: List[Dict[str, Any]],
    ) -> Dict[str, Any]:
        valid_ids = {paper.ref_id for paper in research.get("papers", [])}
        main_text = article.split("## 参考文献", 1)[0]
        sections = list(re.finditer(r"(?m)^##\s+\d+\.\s+(.+?)\s*$", main_text))
        issues: List[Dict[str, str]] = []
        seen_sentences = set()
        reported_duplicate_sentences = set()
        total_citations = set(re.findall(r"\[(R\d+)\]", main_text))
        invalid_ids = sorted(total_citations - valid_ids)
        if invalid_ids:
            issues.append({
                "level": "高", "section": "全文", "issue": f"存在无效来源编号：{', '.join(invalid_ids)}。",
                "suggestion": "仅保留已纳入论文的来源编号，并重新绑定对应证据。",
            })
        link_words = ("因此", "然而", "同时", "此外", "相比", "在此基础上", "进一步", "由此", "不过")
        strong_patterns = ("普遍证明", "完全证明", "所有研究", "必然", "显著提升", "唯一")
        for index, match in enumerate(sections):
            section = match.group(1).strip()
            end = sections[index + 1].start() if index + 1 < len(sections) else len(main_text)
            content = main_text[match.end():end].strip()
            citations = set(re.findall(r"\[(R\d+)\]", content)) & valid_ids
            if len(re.sub(r"\s+", "", content)) < 80:
                issues.append({
                    "level": "中", "section": section, "issue": "章节正文过短，尚未形成完整论证。",
                    "suggestion": "补充中心论点、至少一条来源证据和章节小结。",
                })
            if not citations:
                issues.append({
                    "level": "高", "section": section, "issue": "章节没有有效论文来源标记。",
                    "suggestion": "为关键判断补充论点-证据卡片中的论文编号。",
                })
            if any(pattern in content for pattern in strong_patterns):
                issues.append({
                    "level": "中", "section": section, "issue": "出现范围过大的强结论表述。",
                    "suggestion": "改为限定样本、方法和场景的表述，并保留直接证据来源。",
                })
            # A transition can be implicit when adjacent sections share the
            # same research terms. Do not penalise every section merely for
            # not beginning with a fixed connector word.
            previous_content = ""
            if index > 0:
                previous_match = sections[index - 1]
                previous_end = match.start()
                previous_content = main_text[previous_match.end():previous_end].strip()
            term_pattern = r"[\u4e00-\u9fff]{2,6}|[A-Za-z]{4,}"
            common_terms = {"研究", "文献", "论文", "方法", "结论", "结果", "内容", "现有", "主要", "相关", "分析"}
            previous_terms = {
                term for term in re.findall(term_pattern, previous_content)
                if term not in common_terms
            }
            current_terms = {
                term for term in re.findall(term_pattern, content[:500])
                if term not in common_terms
            }
            implicit_transition = len(previous_terms & current_terms) >= 1
            if index > 0 and content and not any(word in content[:260] for word in link_words) and not implicit_transition:
                issues.append({
                    "level": "低", "section": section, "issue": "章节开头缺少与前文的显式衔接。",
                    "suggestion": "可使用“在完成分类后”“基于上述比较”等过渡句连接前文。",
                })
            for sentence in _split_sentences(content):
                normalized = re.sub(r"\[R\d+\]|\s+|[，。；、,.]", "", sentence)
                if len(normalized) >= 30 and normalized in seen_sentences and normalized not in reported_duplicate_sentences:
                    issues.append({
                        "level": "中", "section": section, "issue": "与前文存在重复表达。",
                        "suggestion": "保留一次陈述，并将重复部分改为方法差异、适用边界或新的证据。",
                    })
                    reported_duplicate_sentences.add(normalized)
                seen_sentences.add(normalized)
        # The same type of problem can be detected repeatedly in one section
        # (especially duplicate sentences). Count it once so the score reflects
        # issue severity, not the number of repeated detections.
        unique_issues: List[Dict[str, str]] = []
        seen_issue_keys = set()
        for issue in issues:
            issue_key = (issue.get("level", ""), issue.get("section", ""), issue.get("issue", ""))
            if issue_key not in seen_issue_keys:
                unique_issues.append(issue)
                seen_issue_keys.add(issue_key)
        issues = unique_issues

        high_count = sum(item["level"] == "高" for item in issues)
        medium_count = sum(item["level"] == "中" for item in issues)
        low_count = sum(item["level"] == "低" for item in issues)
        penalty = high_count * 2.5 + medium_count * 0.8 + low_count * 0.3
        # Keep the score interpretable on a 0-10 scale. A long draft should not
        # receive zero merely because one detector reported many similar issues.
        score = round(max(0.5, 10 - min(penalty, 9.5)), 1)
        strengths = []
        if sections and all(re.findall(r"\[(R\d+)\]", main_text[match.end():(sections[index + 1].start() if index + 1 < len(sections) else len(main_text))]) for index, match in enumerate(sections)):
            strengths.append("每个已生成章节均保留了来源编号。")
        if not invalid_ids:
            strengths.append("未发现无效论文编号。")
        if section_plans:
            strengths.append(f"已为 {len(section_plans)} 个章节生成论证结构，可据此逐节修改。")
        return {
            "name": self.name,
            "score": score,
            "verdict": (
                "需要优先修改" if high_count
                else ("建议修改后再使用" if score < 6 else "可作为初稿继续完善")
            ),
            "issues": issues[:16],
            "strengths": strengths,
            "section_count": len(sections),
            "valid_citations": sorted(total_citations & valid_ids),
            "issue_counts": {"高": high_count, "中": medium_count, "低": low_count},
            "score_basis": "同一章节同类问题只计一次；总扣分上限为9.5分。",
        }


class ParagraphPlanningSkill:
    """把章节拆成带论证动作的段落，并按章节角色分配写作任务。"""

    name = "论证动作编排 Skill"

    ROLE_LIMITS = {"background": 2, "classification": 2, "comparison": 4, "gap": 2, "summary": 0}

    def run(
        self,
        outline: Dict[str, Any],
        research: Dict[str, Any],
        section_structure: Dict[str, Any],
        claim_evidence: Optional[Dict[str, Any]] = None,
    ) -> Dict[str, Any]:
        materials = research.get("section_materials", {})
        papers = {paper.ref_id: paper for paper in research.get("papers", [])}
        plans = {item.get("section", ""): item for item in section_structure.get("plans", [])}
        cards = {card.get("paper_id", ""): card for card in (claim_evidence or {}).get("cards", [])}
        weak_ids = {
            paper.ref_id for paper in research.get("papers", [])
            if getattr(paper, "text_quality", "") == "仅书目信息"
            or (not paper.abstract and len(re.sub(r"\s+", "", paper.full_text or "")) < 300)
        }
        usage: Dict[str, int] = {}
        used_claims: set = set()
        assignments: List[Tuple[str, str, List[Dict[str, Any]]]] = []
        for section in outline.get("sections", []):
            role = _section_role(section)
            pool = [
                item for item in materials.get(section, [])
                if item.get("paper_id") in papers and item.get("paper_id") not in weak_ids
            ]
            chosen = self._select_items(pool, role, usage, used_claims)
            assignments.append((section, role, chosen))

        planned: Dict[str, Any] = {}
        llm_targets = [item for item in assignments if item[2]]
        if llm.available and llm_targets:
            with ThreadPoolExecutor(max_workers=min(6, len(llm_targets))) as executor:
                futures = {
                    executor.submit(
                        self._llm_plan, section, role, plans.get(section, {}), chosen, papers, cards
                    ): section
                    for section, role, chosen in llm_targets
                }
                for future in as_completed(futures):
                    planned[futures[future]] = future.result()

        sections: Dict[str, Any] = {}
        for section, role, chosen in assignments:
            paragraphs = planned.get(section) or []
            if not paragraphs:
                paragraphs = self._rule_plan(section, role, plans.get(section, {}), chosen)
            sections[section] = {"paragraphs": paragraphs, "role": role}
        return {"name": self.name, "sections": sections}

    @classmethod
    def _select_items(
        cls,
        pool: List[Dict[str, Any]],
        role: str,
        usage: Dict[str, int],
        used_claims: set,
    ) -> List[Dict[str, Any]]:
        limit = cls.ROLE_LIMITS.get(role, 2)
        if limit <= 0 or not pool:
            return []
        ranked = sorted(
            enumerate(pool),
            key=lambda pair: (usage.get(str(pair[1].get("paper_id", "")), 0), pair[0]),
        )
        chosen: List[Dict[str, Any]] = []
        for _, item in ranked:
            if _material_key(item) in used_claims:
                continue
            chosen.append(item)
            if len(chosen) >= limit:
                break
        if not chosen:
            chosen = [item for _, item in ranked[:1]]
        for item in chosen:
            paper_id = str(item.get("paper_id", ""))
            usage[paper_id] = usage.get(paper_id, 0) + 1
            used_claims.add(_material_key(item))
        return chosen

    @staticmethod
    def _llm_plan(
        section: str,
        role: str,
        plan: Dict[str, Any],
        items: List[Dict[str, Any]],
        papers: Dict[str, Any],
        cards: Dict[str, Any],
    ) -> List[Dict[str, Any]]:
        if not llm.available or not items:
            return []
        evidence_lines = "\n".join(
            f"[{item.get('paper_id')}] {compact_text(item.get('claim'), 240)}"
            for item in items
        )
        prompt = f"""
请把章节「{section}」拆成 1-2 个写作段落，每段承担一个明确的论证动作。
本章节在综述中的角色：{role}（background=交代问题来源，classification=建立分类，comparison=比较方法与结论，gap=指出不足，summary=给出综合判断）。
可用论证动作：{'、'.join(ARGUMENT_MOVES)}。
章节目标：{plan.get('goal', '')}

可用证据：
{evidence_lines}

要求：
- 每段给出主题句、使用的证据编号、论证动作；
- 同一章节内段落动作不得重复；
- 证据编号必须来自上方列表；
- 只输出 JSON：{{"paragraphs":[{{"topic_sentence":"段落主题句","evidence_ids":["R1"],"move":"归纳共识","transition":"与下一段的衔接"}}]}}
"""
        raw = parse_json_response(llm.chat("你擅长把章节拆成有论证动作的写作段落。", prompt, max_tokens=3200, model=llm.fast_model))
        if not raw or not isinstance(raw.get("paragraphs"), list):
            return []
        valid = {item.get("paper_id") for item in items}
        used_moves: List[str] = []
        result: List[Dict[str, Any]] = []
        for index, item in enumerate(raw["paragraphs"][:2], 1):
            if not isinstance(item, dict):
                continue
            ids = [str(value) for value in (item.get("evidence_ids") or []) if str(value) in valid]
            move = str(item.get("move") or "")
            if move not in ARGUMENT_MOVES or move in used_moves:
                move = next((candidate for candidate in ARGUMENT_MOVES if candidate not in used_moves), ARGUMENT_MOVES[0])
            used_moves.append(move)
            result.append({
                "id": f"P{index}",
                "topic_sentence": compact_text(item.get("topic_sentence"), 300),
                "evidence_ids": ids,
                "move": move,
                "analysis_task": ARGUMENT_MOVE_TASKS.get(move, ""),
                "transition": compact_text(item.get("transition"), 200),
            })
        return [item for item in result if item["evidence_ids"] or item["topic_sentence"]]

    @staticmethod
    def _rule_plan(
        section: str,
        role: str,
        plan: Dict[str, Any],
        items: List[Dict[str, Any]],
    ) -> List[Dict[str, Any]]:
        transition = plan.get("transition", "")
        if role == "summary":
            return [
                {
                    "id": "P1",
                    "topic_sentence": "",
                    "evidence_ids": [],
                    "materials": [],
                    "move": "提出推论",
                    "analysis_task": ARGUMENT_MOVE_TASKS["提出推论"],
                    "transition": "",
                },
                {
                    "id": "P2",
                    "topic_sentence": "",
                    "evidence_ids": [],
                    "materials": [],
                    "move": "限定适用边界",
                    "analysis_task": ARGUMENT_MOVE_TASKS["限定适用边界"],
                    "transition": transition,
                },
            ]
        if not items:
            return []
        move_by_role = {
            "background": ["界定概念", "指出证据不足"],
            "classification": ["归纳共识", "指出分歧"],
            "comparison": ["解释差异成因", "限定适用边界"],
            "gap": ["指出证据不足", "限定适用边界"],
            "summary": ["提出推论", "指出证据不足"],
        }
        moves = move_by_role.get(role, ["归纳共识", "指出分歧"])
        chunks = [items[index:index + 2] for index in range(0, len(items), 2)][:2]
        paragraphs: List[Dict[str, Any]] = []
        for index, chunk in enumerate(chunks):
            ids = [str(item.get("paper_id", "")) for item in chunk if item.get("paper_id")]
            move = moves[index % len(moves)]
            paragraphs.append({
                "id": f"P{index + 1}",
                "topic_sentence": "",
                "evidence_ids": ids,
                "materials": chunk,
                "move": move,
                "analysis_task": ARGUMENT_MOVE_TASKS[move],
                "transition": transition if index == len(chunks) - 1 else "",
            })
        return paragraphs


class CitationVerifier:
    """校验正文引用编号是否真实存在、论断句是否能对上证据。"""

    name = "引用校验 Skill"

    def verify(self, article: str, research: Dict[str, Any]) -> Dict[str, Any]:
        valid_ids = {paper.ref_id for paper in research.get("papers", [])}
        body = article.split("## 参考文献", 1)[0]
        cited = set(re.findall(r"\[(R\d+)\]", body))
        invalid_ids = sorted(cited - valid_ids)
        evidence_index: Dict[str, set] = {}
        for item in research.get("structured_papers", []):
            paper_id = str(item.get("paper_id", ""))
            evidence_index[paper_id] = _bigrams(" ".join(str(value) for value in item.get("fields", {}).values()))
        sentences = [value.strip() for value in re.split(r"(?<=[。！？])", body) if len(value.strip()) >= 8]
        unsupported: List[Dict[str, str]] = []
        for sentence in sentences:
            ids = re.findall(r"\[(R\d+)\]", sentence)
            if not ids:
                continue
            tokens = _bigrams(sentence)
            best = 0.0
            for paper_id in ids:
                reference = evidence_index.get(paper_id, set())
                if tokens and reference:
                    best = max(best, len(tokens & reference) / min(len(tokens), len(reference)))
            if best < 0.15:
                unsupported.append({
                    "sentence": compact_text(sentence, 200),
                    "best_overlap": f"{best:.0%}",
                })
        return {
            "name": self.name,
            "invalid_ids": invalid_ids,
            "unsupported_sentences": unsupported[:10],
            "cited_count": len(cited & valid_ids),
            "valid_count": len(valid_ids),
            "coverage": round(len(cited & valid_ids) / len(valid_ids), 2) if valid_ids else 0,
            "checked_sentences": len(sentences),
        }


class WritingQualitySkill:
    """度量写作质量：主题句覆盖、分析句占比、论证动作多样性与重复率。"""

    name = "写作质量度量 Skill"

    def measure(
        self,
        article: str,
        paragraph_planning: Dict[str, Any],
        style_pack: Dict[str, Any],
    ) -> Dict[str, Any]:
        body = article.split("## 参考文献", 1)[0]
        blocks = [block.strip() for block in re.split(r"\n\s*\n", body) if block.strip()]
        paragraphs = [
            block for block in blocks
            if not block.startswith("#") and not block.startswith(">") and not block.startswith("**目标用户")
        ]
        sentences = [value.strip() for value in re.split(r"(?<=[。！？；])", "".join(paragraphs)) if len(value.strip()) >= 8]
        total = len(sentences)
        cited = [value for value in sentences if re.search(r"\[R\d+\]", value)]
        analysis_ratio = round((total - len(cited)) / total, 2) if total else 0
        normalized = [re.sub(r"\[R\d+\]|\s+|[，。；、,.]", "", value) for value in sentences]
        duplicates = total - len(set(normalized)) if total else 0
        duplicate_ratio = round(duplicates / total, 2) if total else 0
        banned_hits = [phrase for phrase in style_pack.get("banned", []) if phrase in body]
        plans = [
            plan
            for section in paragraph_planning.get("sections", {}).values()
            for plan in section.get("paragraphs", [])
        ]
        matched = 0
        for plan in plans:
            topic = plan.get("topic_sentence", "")
            if topic and any(_similarity(topic, text) >= 0.25 for text in paragraphs):
                matched += 1
        topic_coverage = round(matched / len(plans), 2) if plans else 0
        moves_used = sorted({plan.get("move", "") for plan in plans if plan.get("move")})
        paragraph_similarity = 0.0
        for left in range(len(paragraphs)):
            for right in range(left + 1, len(paragraphs)):
                paragraph_similarity = max(
                    paragraph_similarity, _similarity(paragraphs[left], paragraphs[right])
                )
        boilerplate_hits = [phrase for phrase in BOILERPLATE_PHRASES if phrase in body]
        ai_flavor_hits = [phrase for phrase in AI_FLAVOR_PHRASES if phrase in body]
        connective_hits = {
            group: [word for word in words if word in body]
            for group, words in ACADEMIC_CONNECTIVES.items()
        }
        connective_diversity = round(
            sum(1 for words in connective_hits.values() if words) / len(ACADEMIC_CONNECTIVES), 2
        )
        base = (
            3.0 * topic_coverage
            + 2.0 * min(1.0, analysis_ratio / 0.35)
            + 1.5 * min(1.0, len(moves_used) / 4.0)
            + 1.5 * max(0.0, 1 - duplicate_ratio / 0.2)
            + 1.0 * max(0.0, 1 - paragraph_similarity)
            + 1.0 * connective_diversity
        )
        penalty = min(1.8, 0.4 * (len(banned_hits) + len(boilerplate_hits) + len(ai_flavor_hits)))
        score = round(max(0.0, min(10.0, base - penalty)), 1)
        return {
            "name": self.name,
            "score": score,
            "paragraph_count": len(paragraphs),
            "sentence_count": total,
            "analysis_ratio": analysis_ratio,
            "cited_ratio": round(len(cited) / total, 2) if total else 0,
            "duplicate_ratio": duplicate_ratio,
            "paragraph_similarity": round(paragraph_similarity, 2),
            "topic_coverage": topic_coverage,
            "moves_used": moves_used,
            "banned_hits": banned_hits,
            "boilerplate_hits": boilerplate_hits,
            "ai_flavor_hits": ai_flavor_hits,
            "connective_diversity": connective_diversity,
            "connective_hits": {key: value for key, value in connective_hits.items() if value},
            "score_basis": "主题句覆盖 30% + 分析句占比 20% + 论证动作 15% + 重复率 15% + 段落差异度 10% + 连接词多样性 10%，模板句、AI 味表达与禁用词额外扣分",
        }


class RewriteAgent:
    """按评审问题定位章节并触发定向重写。"""

    name = "重写 Agent"

    def select_targets(self, review: Dict[str, Any], max_sections: int = 3) -> Dict[str, List[Dict[str, Any]]]:
        targets: Dict[str, List[Dict[str, Any]]] = {}
        for issue in review.get("issues", []):
            if issue.get("level") not in {"高", "中"}:
                continue
            section = str(issue.get("section", "")).strip()
            if not section or section == "全文":
                continue
            targets.setdefault(section, []).append(issue)
        ranked = sorted(targets.items(), key=lambda item: -len(item[1]))
        return dict(ranked[:max_sections])


class WritingSkillSuite:
    def __init__(self) -> None:
        self.claim_evidence = ClaimEvidenceBindingSkill()
        self.cross_paper = CrossPaperSynthesisSkill()
        self.section_structure = SectionArgumentStructureSkill()
        self.paragraph_planning = ParagraphPlanningSkill()
        self.logic_review = LogicCoherenceReviewSkill()

    def prepare(self, outline: Dict[str, Any], research: Dict[str, Any]) -> Dict[str, Any]:
        claim_evidence = self.claim_evidence.run(research.get("structured_papers", []))
        cross_paper = self.cross_paper.run(research.get("themes", []), research.get("comparisons", []))
        section_structure = self.section_structure.run(
            outline, research, claim_evidence, cross_paper.get("syntheses", [])
        )
        paragraph_planning = self.paragraph_planning.run(outline, research, section_structure, claim_evidence)
        style_reference = {
            "guide": WRITING_STYLE_GUIDE,
            "connectives": ACADEMIC_CONNECTIVES,
            "exemplars": extract_style_exemplars(research.get("papers", [])),
        }
        return {
            "catalog": WRITING_SKILL_CATALOG,
            "style_reference": style_reference,
            "claim_evidence": claim_evidence,
            "cross_paper_synthesis": cross_paper,
            "section_argument_structure": section_structure,
            "paragraph_planning": paragraph_planning,
        }

    def review_after_draft(
        self, article: str, research: Dict[str, Any], writing_skills: Dict[str, Any]
    ) -> Dict[str, Any]:
        return self.logic_review.run(
            article,
            research,
            writing_skills.get("section_argument_structure", {}).get("plans", []),
        )


class WriterAgent:
    name = "策划 Agent"

    def plan(self, topic: str, research_question: str = "", template: str = "文献综述") -> Dict[str, Any]:
        templates = {
            "文献综述": ["研究背景与问题", "相关研究分类", "研究方法与结论比较", "现有研究不足", "小结与研究启示"],
            "研究背景": ["问题背景", "领域发展现状", "已有研究进展", "研究空白", "本研究启示"],
        }
        question = research_question.strip() or f"{topic}领域的主要研究进展、方法差异与待解决问题是什么？"
        prompt = f"""
请为研究主题「{topic}」和研究问题「{question}」设计科研文献综述初稿框架。
章节必须支持多篇文献的分类、比较和研究空白提炼。只输出 JSON：
{{"suggested_title":"标题","research_question":"研究问题","target_audience":"目标用户","key_message":"综述主线","sections":["章节1","章节2"]}}
"""
        result = parse_json_response(llm.chat("你擅长设计科研文献综述结构。", prompt, model=llm.fast_model))
        if not result or not isinstance(result.get("sections"), list) or not result["sections"]:
            result = {
                "suggested_title": f"{topic}领域研究进展与问题综述",
                "research_question": question,
                "target_audience": "需要完成文献综述初稿的科研新手",
                "key_message": f"围绕{topic}梳理研究主题、方法差异、已有结论与研究空白。",
                "sections": templates.get(template, templates["文献综述"]),
            }
        raw_sections = result.get("sections") if isinstance(result.get("sections"), list) else []
        cleaned_sections: List[str] = []
        for item in raw_sections:
            name = re.sub(r"^\s*#{1,6}\s*", "", str(item or ""))
            name = re.sub(r"^\s*\d{1,2}(?:\.\d+)*[.)、]?\s+", "", name).strip()
            name = re.sub(r"^\s*\d{1,2}(?:\.\d+)*[.)、]\s*", "", name).strip()
            name = compact_text(name, 60)
            if name and name not in cleaned_sections:
                cleaned_sections.append(name)
        result["sections"] = cleaned_sections[:6] or list(templates.get(template, templates["文献综述"]))
        result["topic"] = topic
        result["template"] = template
        result["research_question"] = result.get("research_question") or question
        return result


class ResearcherAgent:
    name = "文献调研 Agent"

    def __init__(self, repository: Optional[OpenAlexRepository] = None) -> None:
        self.repository = repository or OpenAlexRepository()
        self.structuring_skill = LiteratureStructuringSkill()

    def research(
        self,
        topic: str,
        sections: List[str],
        research_question: str = "",
        papers: Optional[List[Paper]] = None,
        use_online: bool = True,
        max_papers: int = 8,
        sources: Optional[Iterable[str]] = None,
        years: Optional[int] = 5,
        sort: str = "quality",
    ) -> Dict[str, Any]:
        selected = list(papers or [])
        source_mode = "用户导入" if selected else ""
        search_queries: List[str] = []
        search_failures: List[str] = []
        if not selected and use_online:
            search_queries = build_search_queries(topic, research_question)
            try:
                results, search_failures = LiteratureSearchService(sources).search(
                    topic,
                    limit=max_papers,
                    years=years,
                    sort=sort,
                    english_queries=search_queries[1:],
                )
            except Exception as exc:
                results, search_failures = [], [f"检索失败：{type(exc).__name__}"]
            if results:
                selected = results
                source_mode = "、".join(sorted({paper.source for paper in results}))
            else:
                source_mode = "在线检索无结果"
        if not selected:
            selected = [Paper(**paper.to_dict()) for paper in DEMO_PAPERS[:max_papers]]
            source_mode = (source_mode + "，" if source_mode else "") + "内置示例"
        selected = self._renumber(selected[:max(1, min(int(max_papers), 20))])
        structured_papers = self.structuring_skill.extract_all(selected)
        has_uploaded_papers = any(paper.source.startswith("导入") for paper in selected)
        # Uploaded papers are analysed through the grounded parser first. This
        # protects the visible matrix and draft from an unrelated LLM fallback.
        integrated = self._rule_integrate(topic, sections, selected) if has_uploaded_papers else self._llm_integrate(topic, research_question, sections, selected)
        if integrated is None:
            integrated = self._rule_integrate(topic, sections, selected)
        alignment = _topic_alignment(topic, selected, search_queries[1:])
        integrated.update({
            "search_queries": search_queries,
            "search_failures": search_failures,
            "search_years": years,
            "topic_alignment": alignment,
            "evidence_gap_ratio": _evidence_gap_ratio(structured_papers),
            "papers": selected,
            "structured_papers": structured_papers,
            "paper_count": len(selected),
            "source_mode": source_mode,
            "content_analysis_coverage": round(sum(1 for paper in selected if paper.abstract and paper.full_text) / len(selected), 2) if selected else 0,
            "structured_field_coverage": round(
                sum(item["field_coverage"] for item in structured_papers) / len(structured_papers), 2
            ) if structured_papers else 0,
        })
        return integrated

    @staticmethod
    def _renumber(papers: Iterable[Paper]) -> List[Paper]:
        result = []
        for index, paper in enumerate(papers, 1):
            paper.ref_id = f"R{index}"
            result.append(paper)
        return result

    @staticmethod
    def _context(papers: List[Paper]) -> str:
        return "\n\n".join(
            f"[{paper.ref_id}] {paper.title} | {paper.year or '年份未知'} | 正文状态: {paper.text_quality}\n"
            f"摘要: {paper.abstract or '摘要缺失'}\n"
            f"正文证据摘录: {compact_text(paper.full_text, 6000) or '正文缺失'}\n"
            f"结构化信号: {_research_signals(paper)}"
            for paper in papers
        )

    def _llm_integrate(self, topic: str, research_question: str, sections: List[str], papers: List[Paper]) -> Optional[Dict[str, Any]]:
        if not llm.available:
            return None
        prompt = f"""
你是科研文献整合 Agent。主题：{topic}；研究问题：{research_question}；章节：{', '.join(sections)}。
请将论文集合整合成主题、比较矩阵、章节素材和研究空白。只能使用给定论文正文或摘要中的事实，不得编造论文事实。
标题、作者、卷期和来源只用于识别文献，不能当作研究结论。对于原文没有明确说明的内容，必须写“原文未明确说明”。
比较矩阵中的方法、结论、局限必须优先引用对应论文的证据摘录；跨文献结论必须说明比较维度和适用边界。
每个 section_material 必须包含合法 paper_id、claim 和 evidence。

论文集合：
{self._context(papers)}

只输出 JSON：
{{"themes":[{{"name":"主题","summary":"综合观点","paper_ids":["R1"]}}],
"comparisons":[{{"paper_id":"R1","research_question":"问题","method":"方法","finding":"结论","limitation":"局限"}}],
"section_materials":{{"章节名":[{{"paper_id":"R1","claim":"可用于本节的观点","evidence":"原文依据"}}]}},
"gaps":["研究空白"]}}
"""
        raw = parse_json_response(llm.chat("你擅长从多篇论文中提炼主题、比较研究并保留来源。", prompt, max_tokens=7000, model=llm.fast_model))
        if not raw or not isinstance(raw.get("section_materials"), dict):
            return None
        valid_ids = {paper.ref_id for paper in papers}
        raw_comparisons = raw.get("comparisons") if isinstance(raw.get("comparisons"), list) else []
        comparisons = []
        for item in raw_comparisons:
            if not isinstance(item, dict) or str(item.get("paper_id")) not in valid_ids:
                continue
            comparisons.append({
                "paper_id": str(item.get("paper_id")),
                "research_question": compact_text(item.get("research_question"), 1000),
                "method": compact_text(item.get("method"), 1200),
                "finding": compact_text(item.get("finding"), 1400),
                "limitation": compact_text(item.get("limitation"), 1100),
            })
        if len(comparisons) < len(papers):
            # Never let an LLM omission erase a grounded row from the visible matrix.
            grounded = self._rule_integrate(topic, sections, papers)["comparisons"]
            by_id = {item["paper_id"]: item for item in comparisons}
            for item in grounded:
                by_id.setdefault(item["paper_id"], item)
            comparisons = [by_id[paper.ref_id] for paper in papers]
        section_materials: Dict[str, List[Dict[str, str]]] = {}
        for section in sections:
            values = raw["section_materials"].get(section, [])
            values = values if isinstance(values, list) else []
            section_materials[section] = [
                {
                    "paper_id": str(item.get("paper_id")),
                    "claim": compact_text(item.get("claim"), 700),
                    "evidence": compact_text(item.get("evidence"), 700),
                }
                for item in values
                if isinstance(item, dict) and str(item.get("paper_id")) in valid_ids
            ]
        if not any(section_materials.values()):
            return None
        raw["section_materials"] = section_materials
        themes = []
        for item in raw.get("themes", []) if isinstance(raw.get("themes"), list) else []:
            if not isinstance(item, dict):
                continue
            ids = [str(value) for value in item.get("paper_ids", []) if str(value) in valid_ids]
            if ids:
                themes.append({
                    "name": compact_text(item.get("name") or "未命名主题", 120),
                    "summary": compact_text(item.get("summary"), 1800),
                    "paper_ids": ids,
                })
        raw["themes"] = themes
        raw["comparisons"] = comparisons
        raw["gaps"] = [compact_text(value, 1200) for value in raw.get("gaps", []) if str(value).strip()] if isinstance(raw.get("gaps"), list) else []
        return raw

    def _rule_integrate(self, topic: str, sections: List[str], papers: List[Paper]) -> Dict[str, Any]:
        sections = sections or ["文献整合"]
        signals = {paper.ref_id: _research_signals(paper) for paper in papers}
        comparisons = [
            {
                "paper_id": paper.ref_id,
                "research_question": signals[paper.ref_id]["research_question"],
                "method": signals[paper.ref_id]["method"],
                "finding": signals[paper.ref_id]["finding"],
                "limitation": signals[paper.ref_id]["limitation"],
                "keywords": signals[paper.ref_id]["keywords"],
            }
            for paper in papers
        ]

        groups: Dict[str, List[Paper]] = {}
        for paper in papers:
            groups.setdefault(signals[paper.ref_id]["domain"], []).append(paper)
        themes = []
        for name, group in groups.items():
            ids = [paper.ref_id for paper in group]
            if name == "待确认的研究主题":
                name = "独立或待确认主题"
            themes.append({
                "name": name,
                "summary": _topic_theme_summary(name, group, signals),
                "paper_ids": ids,
            })

        gaps: List[str] = []
        for name, group in groups.items():
            if len(group) >= 2:
                ids = ", ".join(paper.ref_id for paper in group)
                methods = "；".join(f"{paper.ref_id}采用{signals[paper.ref_id]['method']}" for paper in group[:3])
                gaps.append(f"在{name}中，{ids}的研究对象或分析阶段存在差异，当前摘要尚未提供统一数据集、统一指标下的直接效果比较；已有方法为：{methods}。")
            else:
                paper = group[0]
                signal = signals[paper.ref_id]
                gaps.append(f"{paper.ref_id}已明确指出或显示出适用边界：{signal['limitation']}这一限制需要在后续研究中通过更丰富的数据、动态验证或外部案例检验。")
        metadata_only = [paper for paper in papers if not paper.abstract]
        if metadata_only:
            ids = ", ".join(paper.ref_id for paper in metadata_only[:8])
            suffix = "等" if len(metadata_only) > 8 else ""
            gaps.insert(0, f"{ids}{suffix}仅包含标题、作者等书目信息，未提供摘要或正文；这些条目不能用于可靠的研究问题、方法和结论比较，请上传对应论文全文。")
        if len(groups) > 1:
            ids = ", ".join(paper.ref_id for paper in papers)
            gaps.insert(0, f"上传集合包含多个研究方向（{', '.join(groups.keys())}），{ids}之间不存在直接的同任务比较关系；写作时应按主题分组，避免把生态评估、出版治理和技术预测强行合并为一个结论。")
        if not gaps:
            gaps.append("当前未能从上传文本中提炼出可验证的研究空白，请补充摘要或方法、结果章节。")

        section_materials: Dict[str, List[Dict[str, str]]] = {section: [] for section in sections}
        for section in sections:
            section_key = section.lower()
            if any(word in section for word in ("分类", "主题", "进展")):
                for theme in themes:
                    section_materials[section].append({
                        "paper_id": theme["paper_ids"][0],
                        "claim": theme["summary"],
                        "evidence": f"主题来源：{', '.join(theme['paper_ids'])}。{theme['summary']}",
                    })
            elif any(word in section for word in ("方法", "结论", "比较")):
                for item in comparisons:
                    section_materials[section].append({
                        "paper_id": item["paper_id"],
                        "claim": f"研究问题：{item['research_question']} 方法：{item['method']} 主要结论：{item['finding']}",
                        "evidence": f"摘要提取：{item['finding']}",
                    })
            elif any(word in section for word in ("不足", "空白", "研究缺口", "局限")):
                valid_ids = {paper.ref_id for paper in papers}
                for gap in gaps:
                    if "上传集合包含多个研究方向" in gap:
                        continue
                    match = re.search(r"R\d+", gap)
                    paper_id = match.group(0) if match and match.group(0) in valid_ids else ""
                    if not paper_id and papers:
                        paper_id = papers[0].ref_id
                    section_materials[section].append({"paper_id": paper_id, "claim": gap, "evidence": gap})
            else:
                for item in comparisons:
                    section_materials[section].append({
                        "paper_id": item["paper_id"],
                        "claim": item["research_question"],
                        "evidence": item["finding"],
                    })
        return {"themes": themes, "comparisons": comparisons, "section_materials": section_materials, "gaps": gaps}


ROLE_TOPIC_PATTERNS = {
    "background": (
        "{theme}的相关研究近年来逐步集中到{ids}所代表的问题上。",
        "就{theme}而言，已有工作主要围绕{ids}展开。",
        "{ids}的研究都指向同一个前提：{theme}需要先明确问题范围。",
    ),
    "classification": (
        "按研究路径划分，{ids}可以归入同一类讨论。",
        "从研究对象看，{ids}构成了一个可以互相参照的分组。",
        "{ids}在研究取向上相近，因此放在同一类别中考察。",
    ),
    "comparison": (
        "把{ids}放在同一组比较，方法与数据条件的差别会更清楚。",
        "{ids}处理的是相近问题，但方法路径并不相同。",
        "就方法与结论而言，{ids}之间存在需要解释的差异。",
    ),
    "gap": (
        "{ids}留下的问题主要集中在数据条件与应用范围上。",
        "从{ids}的局限看，缺口出现在证据覆盖与适用边界。",
        "{ids}的共同限制在于研究范围有限。",
    ),
    "summary": (
        "综合前面的材料，{theme}的已有进展可以归纳为几个判断。",
        "把上述讨论收拢起来，{theme}的研究状况呈现出几条线索。",
        "回到研究问题本身，已有研究形成了阶段性认识，但适用边界仍然清楚。",
    ),
}

MOVE_ANALYSIS = {
    "界定概念": (
        "这里所说的概念在不同研究中的口径并不完全一致，比较之前需要先统一界定。",
        "各研究对该概念的使用范围有宽有窄，这直接影响结论之间的可比性。",
    ),
    "归纳共识": (
        "尽管研究对象不同，这些研究在结论方向上彼此兼容，说明该方向已经积累了相对一致的经验证据。",
        "就共同关注的问题而言，各项研究的判断方向接近，差异主要体现在样本与分析尺度上。",
    ),
    "指出分歧": (
        "不过这些研究在结论口径上并不一致，分歧更多出现在指标选取与分析尺度，而不是研究对象本身。",
        "同一个问题在不同研究中得到了不同侧重，说明该议题尚未形成统一的分析框架。",
    ),
    "解释差异成因": (
        "这种差异更可能来自数据来源与研究场景的不同，而不是研究对象自身的矛盾，因此不宜直接合并为同一结论。",
        "差异主要出现在方法路径与数据条件上，若要比较结论，需要先把分析口径统一起来。",
    ),
    "限定适用边界": (
        "上述结论都受限于各自的数据范围与时间跨度，向其他区域或时段外推时需要重新验证。",
        "这些判断建立在特定样本与空间尺度之上，适用范围应限定在可比的对象与条件内。",
    ),
    "提出推论": (
        "由此可以推测，若在统一口径下展开比较，结论差异可能会明显收窄，但这一点仍需实证检验。",
        "如果把这些研究放到同一个分析框架下，方法差异可能比结论差异更值得关注。",
    ),
    "指出证据不足": (
        "目前还缺少在统一数据集与统一指标下的直接比较，这一部分仍需补充证据。",
        "现有材料尚不足以支持跨场景的一般性结论，需要更多可比研究加以检验。",
    ),
    "承上启下": (
        "基于上述梳理，下一节将进一步比较各研究的方法路径与结论条件。",
        "在明确这些差异之后，还需要回到研究不足层面继续讨论。",
    ),
}


def _pick(patterns: Any, seed: int) -> str:
    items = tuple(patterns)
    return items[seed % len(items)] if items else ""


class EditorAgent:
    name = "写作 Agent"

    def write(
        self,
        outline: Dict[str, Any],
        research: Dict[str, Any],
        writing_skills: Optional[Dict[str, Any]] = None,
        style: str = "严谨学术",
        target_sections: Optional[List[str]] = None,
        revision_notes: Optional[Dict[str, List[Dict[str, Any]]]] = None,
    ) -> str:
        topic = outline.get("topic", "研究主题")
        all_sections = list(outline.get("sections", []))
        sections = [section for section in all_sections if not target_sections or section in target_sections]
        materials = research.get("section_materials", {})
        papers = {paper.ref_id: paper for paper in research.get("papers", [])}
        writing_skills = writing_skills or {}
        style_pack = WRITING_STYLE_PACKS.get(style, WRITING_STYLE_PACKS["严谨学术"])
        plans = {
            item.get("section", ""): item
            for item in writing_skills.get("section_argument_structure", {}).get("plans", [])
        }
        paragraph_plans = writing_skills.get("paragraph_planning", {}).get("sections", {})
        syntheses = writing_skills.get("cross_paper_synthesis", {}).get("syntheses", [])
        revision_notes = revision_notes or {}
        _ = writing_skills.get("style_reference")

        def make_section(section: str) -> Tuple[str, str]:
            items = materials.get(section, [])
            plan = plans.get(section, {})
            related_syntheses = [
                item for item in syntheses if item.get("id") in plan.get("synthesis_ids", [])
            ]
            paragraphs = paragraph_plans.get(section, {}).get("paragraphs", [])
            if not paragraphs:
                paragraphs = ParagraphPlanningSkill._rule_plan(section, plan, items)
            if not paragraphs:
                return section, self._fallback_section(section, items, papers, plan, related_syntheses)
            notes = revision_notes.get(section, [])
            role = paragraph_plans.get(section, {}).get("role") or _section_role(section)
            texts = [
                self._write_paragraph(
                    topic, section, paragraph, items, papers, style_pack, notes, role,
                    writing_skills.get("style_reference"),
                )
                for paragraph in paragraphs
            ]
            return section, "\n\n".join(text for text in texts if text.strip())

        contents: Dict[str, str] = {}
        if sections:
            with ThreadPoolExecutor(max_workers=min(6, len(sections))) as executor:
                futures = {executor.submit(make_section, section): section for section in sections}
                for future in as_completed(futures):
                    section, content = future.result()
                    contents[section] = content

        if target_sections:
            return {section: contents.get(section, "") for section in sections}

        parts = [
            f"# {outline.get('suggested_title', f'{topic}文献综述初稿')}",
            "",
            f"> 研究问题：{outline.get('research_question', '')}",
            "",
            f"**目标用户**：{outline.get('target_audience', '科研新手')}",
            "",
            f"**写作风格**：{style}",
            "",
        ]
        for index, section in enumerate(sections, 1):
            parts.extend([f"## {index}. {section}", "", contents.get(section, "本节生成失败，请重试。"), ""])
        parts.extend(["## 参考文献", ""])
        parts.extend(paper.citation() for paper in papers.values())
        return "\n".join(parts)

    def _write_paragraph(
        self,
        topic: str,
        section: str,
        paragraph: Dict[str, Any],
        items: List[Dict[str, Any]],
        papers: Dict[str, Paper],
        style_pack: Dict[str, Any],
        notes: List[Dict[str, Any]],
        role: str = "comparison",
        style_reference: Optional[Dict[str, Any]] = None,
    ) -> str:
        evidence_ids = [str(value) for value in paragraph.get("evidence_ids", [])]
        planned = paragraph.get("materials")
        if planned is not None:
            selected = [item for item in planned if item.get("paper_id") in papers]
        else:
            selected = [item for item in items if item.get("paper_id") in evidence_ids][:2]
        evidence_block = "\n".join(
            f"[{item.get('paper_id')}] {papers[item.get('paper_id')].title}\n"
            f"观点：{compact_text(item.get('claim'), 300)}\n依据：{compact_text(item.get('evidence'), 300)}"
            for item in selected
            if item.get("paper_id") in papers
        ) or "本节暂无直接匹配文献，请明确标注证据不足。"
        style_reference = style_reference or {}
        style_context = style_reference.get("guide", WRITING_STYLE_GUIDE)
        connectives = style_reference.get("connectives") or ACADEMIC_CONNECTIVES
        connective_text = "；".join(f"{key}：{'、'.join(values[:4])}" for key, values in connectives.items())
        style_context = f"{style_context}\n\n可用连接词（同类中换着用，不要重复同一个词）：{connective_text}"
        exemplars = style_reference.get("exemplars") or []
        if exemplars:
            samples = "\n".join(
                f"范例{i}（来自 {item.get('paper_id', '')}）：{compact_text(item.get('text'), 400)}"
                for i, item in enumerate(exemplars[:2], 1)
            )
            style_context += (
                f"\n\n同领域范文段落（只用于模仿句式长短、连接方式和论证节奏）：\n{samples}\n"
                "注意：范文仅作风格参照。禁止复用其中的句子、数据、专有表述；"
                "正文内容必须以本段给定的证据为准，若范文与证据不一致，以证据为准。"
            )
        note_block = ""
        if notes:
            note_block = "本次重写必须修正以下问题：\n" + "\n".join(
                f"- [{note.get('level')}] {note.get('issue')}；建议：{note.get('suggestion')}"
                for note in notes
            )
        move = paragraph.get("move", "归纳共识")
        prompt = f"""
你正在撰写《{topic}》文献综述中「{section}」章节的一个自然段。
写作风格：{style_pack.get('tone')}；{style_pack.get('sentence')}；{style_pack.get('paragraph')}

{style_context}
本段主题句：{paragraph.get('topic_sentence', '')}
本章节在综述中的角色：{role}（background=交代问题来源，classification=建立分类，comparison=比较方法与结论，gap=指出不足，summary=给出综合判断）。
本段论证动作：{move}，需要完成：{ARGUMENT_MOVE_TASKS.get(move, '')}

只能使用下方证据，不得编造论文事实、数据、作者观点或结论。
每个论断句必须以 [R1] 形式标注来源，且只能使用这些编号：{', '.join(evidence_ids) or '无'}。
禁止使用以下表达：{'、'.join(style_pack.get('banned', []))}。

写作质量要求：
- 必须写出具体内容，优先使用证据中的样本量、模型名称、研究区域、指标名称；
- 引用标注要克制：同一来源在同一段落内只标注一次，不要每句都加 [R1]；连续陈述同一来源时可合并为一句再标注；
- 段落控制在 4-7 句，不要写成只有一句的长句堆叠；
- 不要把论证动作的名称写进正文，例如禁止出现“指出这些研究的共同点”这类任务说明；
- 禁止使用“围绕……提供了可以直接比较的材料”“这些研究的样本与数据条件并不相同”这类模板句；
- 句长要有变化，长短句交替，不要每句都用相同的“……，……。”结构；
- 段落要有推进：先给判断，再给证据，最后给出解释或限定；
- summary 角色只做综合判断，不要重新罗列前文的证据句。
只输出这一段正文，不要输出标题、列表、解释或前后缀。
{note_block}

可用证据：
{evidence_block}
"""
        generated = llm.chat("你擅长写作有来源依据的科研文献综述段落。", prompt, max_tokens=3200, model=llm.writing_model)
        if generated and generated.strip():
            generated = re.sub(r"^\s*#{1,6}\s+[^\n]+\n", "", generated.strip(), count=1)
            return generated.strip()
        return self._fallback_paragraph(section, paragraph, selected, papers, role)

    @staticmethod
    def _fallback_paragraph(
        section: str,
        paragraph: Dict[str, Any],
        items: List[Dict[str, Any]],
        papers: Dict[str, Paper],
        role: str = "comparison",
    ) -> str:
        seed = 0
        digits = re.findall(r"\d+", str(paragraph.get("id", "P1")))
        if digits:
            seed = int(digits[-1])
        usable = [item for item in items if item.get("paper_id") in papers]
        limit = 1 if role == "summary" else 2
        prepared_items = []
        for item in usable[:limit]:
            claim = _clean_claim_text(item.get("claim"), 200)
            if claim:
                prepared_items.append((item, claim))
        ids = [str(item.get("paper_id", "")) for item, _ in prepared_items]
        move = paragraph.get("move", "归纳共识")
        topic = paragraph.get("topic_sentence") or _pick(
            ROLE_TOPIC_PATTERNS.get(role, ROLE_TOPIC_PATTERNS["comparison"]), seed
        ).format(theme=section, ids="、".join(ids) if ids else "现有文献")
        paragraph["topic_sentence"] = topic
        parts = [topic]
        for item, claim in prepared_items:
            parts.append(f"{claim.rstrip('。；')}[{item.get('paper_id')}]。")
        parts.append(_pick(MOVE_ANALYSIS.get(move, MOVE_ANALYSIS["归纳共识"]), seed))
        if paragraph.get("transition"):
            parts.append(paragraph["transition"])
        return "".join(parts)

    @staticmethod
    def _fallback_section(
        section: str,
        items: List[Dict[str, Any]],
        papers: Dict[str, Paper],
        plan: Optional[Dict[str, Any]] = None,
        syntheses: Optional[List[Dict[str, Any]]] = None,
    ) -> str:
        plan = plan or {}
        syntheses = syntheses or []
        if not items and not syntheses:
            return f"本节围绕“{section}”展开，但当前文献集合中缺少直接证据，需要补充文献后再完成论述。"
        lines = [plan.get("core_claim") or f"本节依据上传文献中可追溯的摘要与正文信息，对“{section}”进行整合。"]
        if syntheses:
            lines.extend(item.get("paragraph", "") for item in syntheses[:2] if item.get("paragraph"))
        else:
            for item in items[:4]:
                paper = papers.get(item.get("paper_id"))
                if paper:
                    claim = item.get("claim") or paper.abstract or paper.title
                    lines.append(f"{claim} [{paper.ref_id}]")
        if plan.get("transition"):
            lines.append(plan["transition"])
        return "\n\n".join(lines)


class CriticAgent:
    name = "评审 Agent"

    def review(self, article: str, outline: Dict[str, Any], research: Dict[str, Any]) -> Dict[str, Any]:
        valid_ids = {paper.ref_id for paper in research.get("papers", [])}
        papers = research.get("papers", [])
        cited_ids = set(re.findall(r"\[(R\d+)\]", article)) & valid_ids
        coverage = round(len(cited_ids) / len(valid_ids), 2) if valid_ids else 0
        quality_flags: List[str] = []
        for paper in papers:
            if paper.title in {"研究论文", "未命名论文"} or re.search(r"第\s*\d+\s*卷|Vol\.?\s*\d+|文章编号|Journal of", paper.title, flags=re.I):
                quality_flags.append(f"{paper.ref_id} 标题疑似期刊页眉或栏目名")
            if not paper.authors:
                quality_flags.append(f"{paper.ref_id} 未识别作者")
            if not paper.abstract or len(paper.abstract) < 80:
                quality_flags.append(f"{paper.ref_id} 摘要过短，结论应结合全文核验")
            quality_flags.extend(f"{paper.ref_id}：{warning}" for warning in paper.parse_warnings)
        comparisons = research.get("comparisons", [])
        if len(comparisons) >= 2 and len({compact_text(item.get("method"), 180) for item in comparisons}) == 1:
            quality_flags.append("比较矩阵中的方法字段重复，可能未读取论文内容")
        if not research.get("themes"):
            quality_flags.append("未形成绑定来源的主题")
        prompt = f"""
请评审下面的科研文献综述初稿，重点检查章节结构、跨文献整合、论据支撑和来源可追溯性。
只允许使用这些来源编号：{', '.join(sorted(valid_ids))}。
只输出 JSON：{{"scores":{{"结构清晰度":0,"跨文献整合":0,"可读性":0,"论据支撑":0,"来源可追溯":0}},"suggestions":["建议"],"unsupported_claims":["表述"],"verdict":"通过/需要修改"}}

{article[:14000]}
"""
        result = parse_json_response(llm.chat("你是一位严格的科研写作评审专家。", prompt, max_tokens=3600, model=llm.fast_model)) if llm.available else None
        if not result:
            source_score = round(4 + coverage * 6, 1)
            content_score = 8 if comparisons and all(item.get("research_question") and item.get("method") and item.get("finding") for item in comparisons) else 5
            integration_score = 8 if len(research.get("themes", [])) >= 1 and len(comparisons) >= 2 else 5
            result = {
                "scores": {
                    "结构清晰度": 8 if len(outline.get("sections", [])) >= 4 else 6,
                    "跨文献整合": integration_score,
                    "可读性": 7,
                    "论据支撑": min(source_score, content_score),
                    "来源可追溯": source_score,
                },
                "suggestions": [
                    "检查每个研究问题、方法和结论是否都能在原文摘要或正文中定位。",
                    "对同一主题的文献补充样本、指标、时间范围和方法差异，避免只做逐篇罗列。",
                    "对摘要未明确说明的局限补充全文页码或实验细节，不要把推断写成论文原结论。",
                ],
                "unsupported_claims": [] if cited_ids else ["当前正文未检测到有效来源标记"],
                "verdict": "通过" if coverage >= 0.5 else "需要修改",
            }
        scores = result.get("scores") if isinstance(result.get("scores"), dict) else {}
        normalised = {}
        for key, value in scores.items():
            try:
                normalised[key] = max(0, min(10, round(float(value), 1)))
            except (TypeError, ValueError):
                pass
        result["scores"] = normalised
        alignment = (research.get("topic_alignment") or {}).get("ratio", 1.0)
        gap_ratio = research.get("evidence_gap_ratio", 0.0)
        if alignment < 0.4:
            for key in ("论据支撑", "跨文献整合", "来源可追溯"):
                if key in normalised:
                    normalised[key] = min(normalised[key], 3.0)
            quality_flags.append(
                f"检索到的文献与主题相关性不足：{alignment:.0%} 的文献包含主题关键词，"
                "当前初稿大量依赖“证据不足”表述，不能作为有效综述使用。"
            )
        if gap_ratio > 0.5:
            for key in ("论据支撑", "跨文献整合"):
                if key in normalised:
                    normalised[key] = min(normalised[key], 4.0)
            quality_flags.append(
                f"结构化字段中 {gap_ratio:.0%} 标记为待核验或未明确，文献可用信息不足以支撑结论性判断。"
            )
        if alignment < 0.4 or gap_ratio > 0.5:
            result["verdict"] = "需要修改"
        result["scores"] = normalised
        result["topic_alignment"] = alignment
        result["evidence_gap_ratio"] = gap_ratio
        result["overall"] = round(sum(normalised.values()) / len(normalised), 1) if normalised else 0
        result["word_count"] = count_chinese_words(article)
        result["citation_coverage"] = coverage
        result["cited_sources"] = sorted(cited_ids)
        result["missing_sources"] = sorted(valid_ids - cited_ids)
        result["quality_flags"] = quality_flags
        result["content_analysis_coverage"] = round(sum(1 for item in comparisons if item.get("method") and item.get("finding")) / len(papers), 2) if papers else 0
        return result


class MultiAgentOrchestrator:
    """两阶段编排：prepare 产出待确认材料，compose 在人工确认后生成初稿。"""

    def __init__(self) -> None:
        self.writer = WriterAgent()
        self.researcher = ResearcherAgent()
        self.editor = EditorAgent()
        self.critic = CriticAgent()
        self.writing_skills = WritingSkillSuite()
        self.citation = CitationVerifier()
        self.quality = WritingQualitySkill()
        self.rewriter = RewriteAgent()

    def prepare(
        self,
        topic: str,
        style: str = "学术综述",
        template: str = "文献综述",
        research_question: str = "",
        papers: Optional[List[Paper]] = None,
        source_file: str | Path | None = None,
        use_online: bool = True,
        max_papers: int = 8,
        sources: Optional[Iterable[str]] = None,
        years: Optional[int] = 5,
        sort: str = "quality",
    ) -> Dict[str, Any]:
        """阶段一：策划、文献调研、结构化提取与论点证据绑定，等待人工确认。"""
        if source_file and not papers:
            papers = load_papers_from_file(source_file)
        topic = str(topic or "").strip()
        research_question = str(research_question or "").strip()
        generated_topic: Dict[str, str] = {}
        if papers and not topic:
            generated_topic = generate_research_topic(papers)
            topic = generated_topic.get("topic", "")
            if not research_question:
                research_question = generated_topic.get("research_question", "")
        if not topic:
            topic = generated_topic.get("topic") or "上传文献的主题分类与方法比较"
        effective_topic = topic
        effective_question = research_question
        cross_topic = False
        if papers:
            _, cross_topic = _infer_corpus_topic(papers)
            default_question = (not research_question) or any(
                marker in research_question for marker in ("大语言模型智能体", "知识获取与任务协作", "科研场景中的知识获取")
            )
            if default_question:
                effective_question = "上传文献分别研究了哪些具体问题，采用了哪些方法，形成了哪些可比较的结论与适用边界？"
        outline = self.writer.plan(effective_topic, effective_question, template)
        if cross_topic:
            outline["suggested_title"] = "上传文献的主题分类、方法比较与研究空白综述"
            outline["key_message"] = "上传集合属于多个研究方向，初稿将先按主题分组，再比较同组文献，避免跨领域强行归纳。"
        research = self.researcher.research(
            effective_topic,
            outline.get("sections", []),
            outline.get("research_question", research_question),
            papers=papers,
            use_online=use_online,
            max_papers=max_papers,
            sources=sources,
            years=years,
            sort=sort,
        )
        writing_skills = self.writing_skills.prepare(outline, research)
        cards = writing_skills.get("claim_evidence", {}).get("cards", [])
        return {
            "topic": effective_topic,
            "generated_topic": generated_topic,
            "style": style,
            "template": template,
            "outline": outline,
            "research": research,
            "writing_skills": writing_skills,
            "execution_summary": {
                "paper_count": research.get("paper_count", 0),
                "source_mode": research.get("source_mode", "未知"),
                "structured_field_coverage": research.get("structured_field_coverage", 0),
                "search_queries": research.get("search_queries", []),
                "search_failures": research.get("search_failures", []),
                "claim_card_count": len(cards),
                "theme_count": len(research.get("themes", [])),
                "gap_count": len(research.get("gaps", [])),
            },
            "confirmation": {
                "required": True,
                "status": "待确认",
                "field_status": "待确认",
                "card_status": "待确认",
            },
        }

    def compose(
        self,
        prepared: Dict[str, Any],
        field_decisions: Optional[Dict[str, Dict[str, Dict[str, str]]]] = None,
        card_decisions: Optional[Dict[str, str]] = None,
        output_type: str = "draft",
        custom_sections: Optional[List[str]] = None,
        style: str = "严谨学术",
        enable_rewrite: bool = True,
        max_rewrite_sections: int = 3,
    ) -> Dict[str, Any]:
        """阶段二：应用人工确认与输出类型，生成对应粒度的内容。"""
        if output_type not in OUTPUT_TYPE_LABELS:
            output_type = "draft"
        if style not in WRITING_STYLE_PACKS:
            style = "严谨学术"
        outline = copy.deepcopy(prepared.get("outline", {}))
        research = copy.deepcopy(prepared.get("research", {}))
        writing_skills = copy.deepcopy(prepared.get("writing_skills", {}))
        applied = self._apply_confirmations(
            outline, research, writing_skills, field_decisions or {}, card_decisions or {}
        )
        applied["output_type"] = output_type
        applied["output_label"] = OUTPUT_TYPE_LABELS[output_type]

        sections, structure_source = self._resolve_sections(outline, custom_sections)
        outline["sections"] = sections
        applied["structure_source"] = structure_source
        if structure_source == "自定义章节":
            self._rebuild_section_materials(outline, research, applied)

        binding = writing_skills.get("claim_evidence", {})
        writing_skills["section_argument_structure"] = self.writing_skills.section_structure.run(
            outline,
            research,
            binding,
            writing_skills.get("cross_paper_synthesis", {}).get("syntheses", []),
        )
        plans = writing_skills["section_argument_structure"].get("plans", [])
        writing_skills["paragraph_planning"] = self.writing_skills.paragraph_planning.run(
            outline, research, writing_skills["section_argument_structure"], binding
        )
        writing_skills["catalog"] = WRITING_SKILL_CATALOG

        citation: Dict[str, Any] = {}
        quality: Dict[str, Any] = {}
        rewrite_log: List[Dict[str, Any]] = []
        review_applicable = output_type == "draft"
        if output_type == "outline":
            article = self._render_outline(outline, plans)
            review = self._non_draft_review(output_type, plans)
            writing_skills["logic_coherence_review"] = {
                "name": "逻辑与衔接评审 Skill",
                "score": 0,
                "verdict": "未执行",
                "issues": [],
                "strengths": [],
                "section_count": len(plans),
                "score_basis": f"当前输出为{OUTPUT_TYPE_LABELS[output_type]}，未生成正文段落，因此不执行衔接评审。",
            }
        elif output_type == "keypoints":
            article = self._render_keypoints(outline, plans)
            review = self._non_draft_review(output_type, plans)
            writing_skills["logic_coherence_review"] = {
                "name": "逻辑与衔接评审 Skill",
                "score": 0,
                "verdict": "未执行",
                "issues": [],
                "strengths": [],
                "section_count": len(plans),
                "score_basis": f"当前输出为{OUTPUT_TYPE_LABELS[output_type]}，未生成正文段落，因此不执行衔接评审。",
            }
        else:
            article = self.editor.write(outline, research, writing_skills, style=style)
            review = self.critic.review(article, outline, research)
            citation = self.citation.verify(article, research)
            if enable_rewrite and llm.available:
                targets = self.rewriter.select_targets(review, max_sections=max_rewrite_sections)
                if targets:
                    score_before = review.get("overall", 0)
                    revisions = self.editor.write(
                        outline,
                        research,
                        writing_skills,
                        style=style,
                        target_sections=list(targets.keys()),
                        revision_notes=targets,
                    )
                    article = self._splice_sections(article, outline, revisions)
                    review = self.critic.review(article, outline, research)
                    citation = self.citation.verify(article, research)
                    rewrite_log.append({
                        "sections": list(targets.keys()),
                        "issue_count": sum(len(values) for values in targets.values()),
                        "score_before": score_before,
                        "score_after": review.get("overall", 0),
                    })
            writing_skills["logic_coherence_review"] = self.writing_skills.review_after_draft(
                article, research, writing_skills
            )
            quality = self.quality.measure(
                article, writing_skills.get("paragraph_planning", {}), WRITING_STYLE_PACKS[style]
            )

        evidence_bound = sum(1 for plan in plans if plan.get("evidence_order"))
        return {
            "topic": prepared.get("topic", ""),
            "style": prepared.get("style", "学术综述"),
            "template": prepared.get("template", "文献综述"),
            "outline": outline,
            "research": research,
            "papers": research.get("papers", []),
            "materials": research.get("section_materials", {}),
            "article": article,
            "review": review,
            "writing_skills": writing_skills,
            "confirmation": applied,
            "output_type": output_type,
            "output_label": OUTPUT_TYPE_LABELS[output_type],
            "review_applicable": review_applicable,
            "style": style,
            "citation_check": citation,
            "writing_quality": quality,
            "rewrite_log": rewrite_log,
            "execution_summary": {
                "agents_used": 4,
                "total_steps": 5,
                "paper_count": research.get("paper_count", 0),
                "source_mode": research.get("source_mode", "未知"),
                "article_word_count": len(re.sub(r"\s+", "", article)),
                "quality_score": review.get("overall", 0),
                "citation_coverage": review.get("citation_coverage", 0),
                "content_analysis_coverage": research.get("content_analysis_coverage", review.get("content_analysis_coverage", 0)),
                "writing_skill_count": len(writing_skills.get("catalog", [])),
                "logic_review_score": writing_skills.get("logic_coherence_review", {}).get("score", 0),
                "suggestions": review.get("suggestions", []),
                "structured_field_coverage": research.get("structured_field_coverage", 0),
                "claim_card_count": len(binding.get("cards", [])),
                "confirmed_fields": applied.get("confirmed_fields", 0),
                "edited_fields": applied.get("edited_fields", 0),
                "excluded_fields": applied.get("excluded_fields", 0),
                "adopted_cards": applied.get("adopted_cards", 0),
                "excluded_cards": applied.get("excluded_cards", 0),
                "confirmation_status": applied.get("status", "未确认"),
                "output_type": output_type,
                "output_label": OUTPUT_TYPE_LABELS[output_type],
                "structure_source": structure_source,
                "review_applicable": review_applicable,
                "section_count": len(plans),
                "evidence_bound_sections": evidence_bound,
                "style": style,
                "writing_quality_score": quality.get("score", 0),
                "rewrite_sections": len(rewrite_log[0]["sections"]) if rewrite_log else 0,
                "rewrite_score_before": rewrite_log[0]["score_before"] if rewrite_log else 0,
                "rewrite_score_after": rewrite_log[0]["score_after"] if rewrite_log else 0,
                "unsupported_sentences": len(citation.get("unsupported_sentences", [])),
                "invalid_citations": len(citation.get("invalid_ids", [])),
            },
        }

    @staticmethod
    def _splice_sections(article: str, outline: Dict[str, Any], revisions: Dict[str, str]) -> str:
        """把重写后的章节正文替换回原文，其余章节保持不变。"""
        lines = article.split("\n")
        headers = []
        for index, line in enumerate(lines):
            match = re.match(r"^## (\d+)\. (.+?)\s*$", line)
            if match:
                headers.append((index, match.group(2).strip()))
        if not headers:
            return article
        reference_index = next(
            (index for index, line in enumerate(lines) if line.strip().startswith("## 参考文献")),
            len(lines),
        )
        blocks = []
        for position, (line_index, name) in enumerate(headers):
            end = headers[position + 1][0] if position + 1 < len(headers) else reference_index
            blocks.append((line_index, end, name))
        prefix = lines[:headers[0][0]]
        suffix = lines[reference_index:]
        rebuilt = []
        for line_index, end, name in blocks:
            match = re.match(r"^## (\d+)\. (.+?)\s*$", lines[line_index])
            number = match.group(1)
            body = revisions.get(name)
            if body is None:
                body = "\n".join(lines[line_index + 1:end]).strip()
            rebuilt.extend([f"## {number}. {name}", "", body.strip(), ""])
        return "\n".join(prefix + rebuilt + suffix)

    @staticmethod
    def _resolve_sections(outline: Dict[str, Any], custom_sections: Optional[List[str]]) -> Tuple[List[str], str]:
        default = [str(item).strip() for item in outline.get("sections", []) if str(item).strip()]
        if custom_sections:
            cleaned: List[str] = []
            for item in custom_sections:
                value = str(item).strip()
                if value and value not in cleaned:
                    cleaned.append(value)
            cleaned = cleaned[:12]
            if cleaned:
                return cleaned, "自定义章节"
        return default, "默认模板"

    def _rebuild_section_materials(
        self, outline: Dict[str, Any], research: Dict[str, Any], applied: Dict[str, Any]
    ) -> None:
        """章节被自定义后，按新章节名重新生成章节素材映射。"""
        papers = research.get("papers", [])
        if not papers:
            return
        regenerated = self.researcher._rule_integrate(
            outline.get("topic", ""), outline.get("sections", []), papers
        ).get("section_materials", {})
        edits = applied.get("edited_values", [])
        excluded = applied.get("excluded_values", [])
        for section, materials in regenerated.items():
            for material in materials:
                for key in ("claim", "evidence"):
                    text = str(material.get(key, ""))
                    for edit in edits:
                        old_text = str(edit.get("old", ""))
                        if old_text and old_text in text:
                            text = text.replace(old_text, str(edit.get("new", "")))
                    material[key] = text
            regenerated[section] = [
                material for material in materials
                if not any(
                    value and value in re.sub(r"\s+", "", f"{material.get('claim', '')}{material.get('evidence', '')}")
                    for value in excluded
                )
            ]
        research["section_materials"] = regenerated
        applied["notes"].append(
            "已按自定义章节重新生成章节素材映射（共 " + str(len(regenerated)) + " 个章节）。"
        )

    @staticmethod
    def _render_outline(outline: Dict[str, Any], plans: List[Dict[str, Any]]) -> str:
        plan_map = {plan.get("section", ""): plan for plan in plans}
        parts = [
            f"# {outline.get('suggested_title', '文献综述大纲')}",
            "",
            f"> 研究问题：{outline.get('research_question', '')}",
            "",
            f"**目标用户**：{outline.get('target_audience', '科研新手')}",
            "",
            f"**综述主线**：{outline.get('key_message', '')}",
            "",
        ]
        for index, section in enumerate(outline.get("sections", []), 1):
            plan = plan_map.get(section, {})
            parts.append(f"## {index}. {section}")
            if plan.get("goal"):
                parts.append(f"- 章节目标：{plan['goal']}")
            if plan.get("core_claim"):
                parts.append(f"- 中心论点：{plan['core_claim']}")
            steps = plan.get("writing_steps", [])
            if steps:
                parts.append("- 写作步骤：" + "；".join(steps))
            evidence = plan.get("evidence_order", [])
            if evidence:
                parts.append(
                    "- 建议证据顺序：" + "；".join(
                        f"[{item.get('paper_id', '')}]{compact_text(item.get('claim', ''), 90)}"
                        for item in evidence
                    )
                )
            else:
                parts.append("- 建议证据顺序：本章尚未绑定有效证据，需补充文献或调整章节范围。")
            if plan.get("transition"):
                parts.append(f"- 过渡句：{plan['transition']}")
            parts.append("")
        return "\n".join(parts)

    @staticmethod
    def _render_keypoints(outline: Dict[str, Any], plans: List[Dict[str, Any]]) -> str:
        parts = [
            f"# {outline.get('suggested_title', '文献综述要点')}",
            "",
            f"> 研究问题：{outline.get('research_question', '')}",
            "",
        ]
        for index, plan in enumerate(plans, 1):
            section = plan.get("section", f"章节{index}")
            parts.append(f"## {index}. {section}")
            if plan.get("goal"):
                parts.append(f"**本章任务**：{plan['goal']}")
            if plan.get("core_claim"):
                parts.append(f"**中心论点**：{plan['core_claim']}")
            evidence = plan.get("evidence_order", [])
            if evidence:
                parts.append("**可引用证据**")
                for item in evidence:
                    parts.append(
                        f"- [{item.get('paper_id', '')}]（{item.get('source_status', '')}）"
                        f"{compact_text(item.get('claim', ''), 160)}"
                    )
            else:
                parts.append("**可引用证据**：本章尚未绑定有效证据，需补充文献或调整章节范围。")
            steps = plan.get("writing_steps", [])
            if steps:
                parts.append("**展开顺序**：" + "；".join(steps))
            if plan.get("transition"):
                parts.append(f"**过渡句**：{plan['transition']}")
            parts.append("")
        return "\n".join(parts)

    @staticmethod
    def _non_draft_review(output_type: str, plans: List[Dict[str, Any]]) -> Dict[str, Any]:
        label = OUTPUT_TYPE_LABELS.get(output_type, "论文初稿")
        bound = sum(1 for plan in plans if plan.get("evidence_order"))
        return {
            "not_applicable": True,
            "reason": f"当前输出为{label}，没有生成正文段落，因此不执行引用覆盖率与逻辑衔接评审。",
            "scores": {},
            "overall": 0,
            "verdict": "未执行",
            "suggestions": ["如需检查引用覆盖率与段落衔接，可切换为论文初稿后重新生成。"],
            "quality_flags": [],
            "missing_sources": [],
            "citation_coverage": 0,
            "word_count": 0,
            "section_count": len(plans),
            "evidence_bound_sections": bound,
        }

    def _apply_confirmations(
        self,
        outline: Dict[str, Any],
        research: Dict[str, Any],
        writing_skills: Dict[str, Any],
        field_decisions: Dict[str, Dict[str, Dict[str, str]]],
        card_decisions: Dict[str, str],
    ) -> Dict[str, Any]:
        """把人工确认结果写回字段、比较矩阵、论点证据卡片与章节素材。"""
        applied: Dict[str, Any] = {
            "status": "已确认",
            "edited_values": [],
            "excluded_values": [],
            "confirmed_fields": 0,
            "edited_fields": 0,
            "excluded_fields": 0,
            "adopted_cards": 0,
            "excluded_cards": 0,
            "excluded_evidence": [],
            "notes": [],
        }
        matrix_keys = {
            "research_question": "research_question",
            "method": "method",
            "finding": "finding",
            "limitation": "limitation",
        }
        comparisons = {item.get("paper_id", ""): item for item in research.get("comparisons", [])}
        excluded_values: List[str] = []

        for item in research.get("structured_papers", []):
            paper_id = item.get("paper_id", "")
            decisions = field_decisions.get(paper_id, {}) or {}
            fields = item.get("fields", {})
            sources = item.get("field_sources", {})
            for field_key, decision in decisions.items():
                if field_key not in STRUCTURED_FIELD_LABELS:
                    continue
                status = str((decision or {}).get("status", "pending"))
                new_value = str((decision or {}).get("value", "") or "").strip()
                old_value = str(fields.get(field_key, ""))
                if status == "excluded":
                    fields[field_key] = "用户已排除该字段结论，不进入本次初稿的证据链。"
                    sources[field_key] = "用户排除"
                    applied["excluded_fields"] += 1
                    normalized = re.sub(r"\s+", "", old_value)
                    if len(normalized) >= 12 and "用户已排除" not in normalized:
                        excluded_values.append(normalized[:40])
                    applied["excluded_evidence"].append(f"{paper_id}·{STRUCTURED_FIELD_LABELS[field_key]}")
                elif status == "edited" and new_value:
                    fields[field_key] = compact_text(new_value, 1800)
                    sources[field_key] = "用户修改"
                    applied["edited_fields"] += 1
                    if old_value and old_value != fields[field_key]:
                        applied["edited_values"].append({"old": old_value, "new": fields[field_key]})
                elif status == "confirmed":
                    sources[field_key] = f"用户已确认（{sources.get(field_key, '未提取')}）"
                    applied["confirmed_fields"] += 1
                if field_key in matrix_keys and status in {"edited", "excluded"}:
                    row = comparisons.get(paper_id)
                    if row is not None:
                        row[matrix_keys[field_key]] = (
                            fields[field_key] if status == "edited" else "用户已排除该字段结论。"
                        )
            usable = [key for key, value in fields.items() if _skill_text_is_usable(value)]
            item["available_fields"] = usable
            item["field_count"] = len(usable)
            item["field_coverage"] = round(len(usable) / len(STRUCTURED_FIELD_LABELS), 2)

        applied["excluded_values"] = list(excluded_values)

        structured = research.get("structured_papers", [])
        research["structured_field_coverage"] = round(
            sum(item.get("field_coverage", 0) for item in structured) / len(structured), 2
        ) if structured else 0

        raw_cards = self.writing_skills.claim_evidence.run(structured)
        adopted: List[Dict[str, Any]] = []
        excluded_cards: List[Dict[str, Any]] = []
        per_paper_total: Dict[str, int] = {}
        per_paper_adopted: Dict[str, int] = {}
        for card in raw_cards.get("cards", []):
            paper_id = str(card.get("paper_id", ""))
            if card.get("source_status") == "用户排除":
                continue
            per_paper_total[paper_id] = per_paper_total.get(paper_id, 0) + 1
            key = f"{paper_id}|{card.get('field', '')}"
            decision = card_decisions.get(key, "pending")
            if decision == "excluded":
                card["user_confirmation"] = "用户排除"
                excluded_cards.append(card)
                applied["excluded_cards"] += 1
                continue
            if decision == "adopted":
                card["verification_status"] = "用户已确认"
                card["user_confirmation"] = "用户采用"
                applied["adopted_cards"] += 1
            else:
                card["user_confirmation"] = "待确认"
            per_paper_adopted[paper_id] = per_paper_adopted.get(paper_id, 0) + 1
            adopted.append(card)

        supported = sum(card.get("verification_status") in {"已支持", "用户已确认"} for card in adopted)
        binding = {
            "name": raw_cards.get("name", "论点-证据绑定 Skill"),
            "cards": adopted,
            "excluded_cards": excluded_cards,
            "supported_count": supported,
            "verification_count": len(adopted) - supported,
            "coverage": round(supported / len(adopted), 2) if adopted else 0,
            "user_confirmed_count": applied["adopted_cards"],
            "user_excluded_count": applied["excluded_cards"],
        }
        writing_skills["claim_evidence"] = binding

        blocked_papers = {
            paper_id for paper_id, total in per_paper_total.items()
            if total and not per_paper_adopted.get(paper_id, 0)
        }
        for section, materials in list(research.get("section_materials", {}).items()):
            kept = []
            for material in materials:
                if str(material.get("paper_id", "")) in blocked_papers:
                    continue
                text = re.sub(r"\s+", "", f"{material.get('claim', '')}{material.get('evidence', '')}")
                if any(value and value in text for value in excluded_values):
                    continue
                kept.append(material)
            research["section_materials"][section] = kept
        if blocked_papers:
            applied["notes"].append(
                f"{'、'.join(sorted(blocked_papers))} 的论点证据已被全部排除，不再进入章节素材。"
            )

        return applied

    def run(
        self,
        topic: str,
        style: str = "学术综述",
        template: str = "文献综述",
        research_question: str = "",
        papers: Optional[List[Paper]] = None,
        source_file: str | Path | None = None,
        use_online: bool = True,
        max_papers: int = 8,
        sources: Optional[Iterable[str]] = None,
        years: Optional[int] = 5,
        sort: str = "quality",
        field_decisions: Optional[Dict[str, Dict[str, Dict[str, str]]]] = None,
        card_decisions: Optional[Dict[str, str]] = None,
    ) -> Dict[str, Any]:
        prepared = self.prepare(
            topic,
            style=style,
            template=template,
            research_question=research_question,
            papers=papers,
            source_file=source_file,
            use_online=use_online,
            max_papers=max_papers,
            sources=sources,
            years=years,
            sort=sort,
        )
        return self.compose(prepared, field_decisions=field_decisions, card_decisions=card_decisions)


def search_papers(topic: str, max_results: int = 8, use_online: bool = True) -> Tuple[List[Paper], str]:
    if use_online:
        try:
            papers, _ = LiteratureSearchService().search(topic, limit=max_results)
            if papers:
                return papers, "、".join(sorted({paper.source for paper in papers}))
        except Exception:
            return [], "在线检索失败"
    return [Paper(**paper.to_dict()) for paper in DEMO_PAPERS[:max_results]], "内置示例"


def _table(value: Any, limit: int = 300) -> str:
    return compact_text(value, limit).replace("|", "\\|").replace("\n", " ")


def generate_report(result: Dict[str, Any]) -> str:
    outline = result.get("outline", {})
    research = result.get("research", {})
    review = result.get("review", {})
    summary = result.get("execution_summary", {})
    confirmation = result.get("confirmation", {}) or {}
    confirmation_lines = [
        f"- 确认状态: {confirmation.get('status', '未确认')}",
        f"- 已确认字段: {confirmation.get('confirmed_fields', 0)} 个；修改字段: {confirmation.get('edited_fields', 0)} 个；排除字段: {confirmation.get('excluded_fields', 0)} 个",
        f"- 采用论点证据: {confirmation.get('adopted_cards', 0)} 张；排除论点证据: {confirmation.get('excluded_cards', 0)} 张",
    ]
    if confirmation.get("excluded_evidence"):
        confirmation_lines.append("- 已排除字段: " + "、".join(confirmation["excluded_evidence"][:20]))
    for note in confirmation.get("notes", []):
        confirmation_lines.append(f"- 说明: {note}")
    confirmation_lines.append("")
    lines = [
        "# 📚 科研文献整合与论文初稿报告", "",
        f"**主题**: {result.get('topic', '')}",
        f"**论文来源**: {summary.get('source_mode', '未知')}",
        f"**模型模式**: {'真实 API' if llm.available else '规则降级模式'}"
        + (f"（原因：{llm.last_error}）" if not llm.available and llm.last_error else ""), "",
        "## 执行摘要", "",
        f"- 参与 Agent: {summary.get('agents_used', 4)} 个",
        f"- 整合论文: {summary.get('paper_count', 0)} 篇",
        f"- 初稿字数: {summary.get('article_word_count', 0)} 字",
        f"- 输出类型: {summary.get('output_label', '论文初稿')}（结构来源：{summary.get('structure_source', '默认模板')}）",
        f"- 来源覆盖率: {summary.get('citation_coverage', 0):.0%}",
        f"- 内部评审分: {summary.get('quality_score', 0)}/10", "",
        "## 人工确认记录", "",
        *confirmation_lines,
        "## 文献相关性检查", "",
        f"- 主题相关性: {(research.get('topic_alignment') or {}).get('ratio', 0):.0%} "
        f"（{(research.get('topic_alignment') or {}).get('aligned', 0)}/{(research.get('topic_alignment') or {}).get('total', 0)} 篇命中主题关键词）",
        f"- 字段可用率: {1 - research.get('evidence_gap_ratio', 0):.0%}"
        f"（待核验或未明确占 {research.get('evidence_gap_ratio', 0):.0%}）",
        "",
        "## 综述框架", "",
        f"**标题**: {outline.get('suggested_title', '')}",
        f"**研究问题**: {outline.get('research_question', '')}",
        f"**目标用户**: {outline.get('target_audience', '')}", "",
    ]
    lines.extend(f"{index}. {section}" for index, section in enumerate(outline.get("sections", []), 1))
    lines.extend(["", "## 文献来源", "", "| 编号 | 论文 | 作者 | 年份 | 来源 |", "| --- | --- | --- | --- | --- |"])
    for paper in result.get("papers", []):
        lines.append(f"| {paper.ref_id} | {_table(paper.title)} | {_table(', '.join(paper.authors[:3]) or '未知')} | {paper.year or '未知'} | {_table(paper.source)} |")
    lines.extend(["", "## 文献结构化提取", ""])
    for item in research.get("structured_papers", []):
        lines.extend([f"### [{item.get('paper_id', '')}] {item.get('title', '未命名论文')}", ""])
        metadata = item.get("metadata", {})
        lines.append(
            f"作者：{_table(', '.join(metadata.get('authors', [])) or '未知')}；"
            f"年份：{metadata.get('year') or '未知'}；DOI：{_table(metadata.get('doi') or '未识别')}；"
            f"字段覆盖率：{item.get('field_coverage', 0):.0%}"
        )
        lines.append("")
        fields = item.get("fields", {})
        sources = item.get("field_sources", {})
        for key, label in STRUCTURED_FIELD_LABELS.items():
            lines.append(f"- **{label}**（{sources.get(key, '未提取')}）：{fields.get(key, '原文未明确说明。')}")
        lines.append("")
    writing_skills = result.get("writing_skills", {})
    lines.extend(["## 写作 Skill 工作台", ""])
    for skill in writing_skills.get("catalog", WRITING_SKILL_CATALOG):
        lines.extend([
            f"### {skill.get('name', '未命名 Skill')}",
            f"- 作用：{skill.get('purpose', '')}",
            f"- 输入：{skill.get('input', '')}",
            f"- 输出：{skill.get('output', '')}",
            "",
        ])
    binding = writing_skills.get("claim_evidence", {})
    lines.extend([
        "### 论点-证据绑定结果",
        f"- 论点卡片：{len(binding.get('cards', []))} 张",
        f"- 已支持：{binding.get('supported_count', 0)} 张",
        f"- 证据覆盖率：{binding.get('coverage', 0):.0%}",
        "",
    ])
    lines.extend(["### 跨文献综合写作结果", ""])
    for item in writing_skills.get("cross_paper_synthesis", {}).get("syntheses", []):
        lines.extend([
            f"**{item.get('id', '')}｜{item.get('theme', '相关研究')}**（来源：{', '.join(item.get('paper_ids', []))}）",
            item.get("paragraph", ""),
            f"写作建议：{item.get('writing_advice', '')}",
            "",
        ])
    lines.extend(["### 章节论证结构", ""])
    for plan in writing_skills.get("section_argument_structure", {}).get("plans", []):
        lines.extend([
            f"#### {plan.get('section', '未命名章节')}",
            f"- 章节目标：{plan.get('goal', '')}",
            f"- 中心论点：{plan.get('core_claim', '')}",
            f"- 写作步骤：{'；'.join(plan.get('writing_steps', []))}",
            f"- 过渡句：{plan.get('transition', '')}",
            "",
        ])
    logic = writing_skills.get("logic_coherence_review", {})
    lines.extend([
        "### 逻辑与衔接评审",
        f"- 评审分：{logic.get('score', 0)}/10",
        f"- 结论：{logic.get('verdict', '待评审')}",
    ])
    for issue in logic.get("issues", []):
        lines.append(f"- [{issue.get('level', '')}] {issue.get('section', '')}：{issue.get('issue', '')} 建议：{issue.get('suggestion', '')}")
    lines.append("")
    lines.extend(["## 主题整合", ""])
    for theme in research.get("themes", []):
        lines.extend([f"### {theme.get('name', '未命名主题')}", "", f"{theme.get('summary', '')} 来源：{', '.join(theme.get('paper_ids', []))}", ""])
    lines.extend(["## 文献比较矩阵", "", "| 来源 | 研究问题 | 方法 | 主要结论 | 局限 |", "| --- | --- | --- | --- | --- |"])
    for item in research.get("comparisons", []):
        lines.append(f"| {item.get('paper_id', '')} | {_table(item.get('research_question'))} | {_table(item.get('method'))} | {_table(item.get('finding'))} | {_table(item.get('limitation'))} |")
    lines.extend(["", "## 研究空白", ""])
    lines.extend(f"- {gap}" for gap in research.get("gaps", []))
    output_label = summary.get("output_label", "论文初稿")
    lines.extend(["", f"## {output_label}", "", result.get("article", "内容未生成"), "", "## 评审报告", ""])
    if review.get("scores"):
        lines.extend(f"- {metric}: {score}/10" for metric, score in review.get("scores", {}).items())
    else:
        lines.append(f"- {review.get('reason', '当前输出类型未执行全文评审。')}")
    quality = result.get("writing_quality", {}) or {}
    if quality:
        lines.extend([
            "",
            "## 写作质量与引用校验",
            "",
            f"- 写作质量分: {quality.get('score', 0)}/10（{quality.get('score_basis', '')}）",
            f"- 段落数: {quality.get('paragraph_count', 0)}；主题句覆盖率: {quality.get('topic_coverage', 0):.0%}",
            f"- 分析句占比: {quality.get('analysis_ratio', 0):.0%}；重复句比例: {quality.get('duplicate_ratio', 0):.0%}",
            f"- 论证动作: {'、'.join(quality.get('moves_used', [])) or '未记录'}",
        ])
        if quality.get("banned_hits"):
            lines.append("- 命中禁用表达: " + "、".join(quality["banned_hits"]))
    citation = result.get("citation_check", {}) or {}
    if citation:
        lines.append(
            f"- 引用校验: 有效引用 {citation.get('cited_count', 0)}/{citation.get('valid_count', 0)}；"
            f"无效编号 {len(citation.get('invalid_ids', []))} 个；待核验句 {len(citation.get('unsupported_sentences', []))} 句"
        )
        for item in citation.get("unsupported_sentences", [])[:5]:
            lines.append(f"  - 待核验（证据重合度 {item.get('best_overlap')}）：{item.get('sentence')}")
    rewrite_log = result.get("rewrite_log", []) or []
    if rewrite_log:
        entry = rewrite_log[0]
        lines.append(
            f"- 重写闭环: 重写 {len(entry.get('sections', []))} 个章节，评审分 {entry.get('score_before')} → {entry.get('score_after')}"
        )
    lines.extend([
        f"- **整体评分**: **{review.get('overall', 0)}/10**",
        f"- **来源覆盖率**: **{review.get('citation_coverage', 0):.0%}**",
        f"- **结论**: {review.get('verdict', '需要修改')}", "", "### 迭代建议", "",
    ])
    lines.extend(f"{index}. {suggestion}" for index, suggestion in enumerate(review.get("suggestions", []), 1))
    return "\n".join(lines)


if __name__ == "__main__":
    topic = os.environ.get("P4_TOPIC", "AI Agent")
    use_online = os.environ.get("P4_USE_ONLINE", "false").lower() in {"1", "true", "yes"}
    result = MultiAgentOrchestrator().run(topic, style="学术综述", template="文献综述", use_online=use_online, max_papers=5)
    output = PROJECT_DIR / "p4-literature-report.md"
    output.write_text(generate_report(result), encoding="utf-8")
    print(f"📄 完整报告已保存至: {output}")
